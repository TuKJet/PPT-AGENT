from __future__ import annotations

import base64
import io
import json
import math
import re
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from PIL import Image, ImageChops, ImageFilter, ImageStat
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from html_pipeline.html_builder import render_html_screenshot, write_slide_status
from vendor_presentation_core.export.powerpoint_preview_renderer import (
    detect_powerpoint_render_support,
    export_powerpoint_slide_previews,
)


_STATUS_PRIORITY = {"pass": 0, "warn": 1, "fail": 2}
_RESAMPLE_LANCZOS = getattr(getattr(Image, "Resampling", Image), "LANCZOS")


def _utc_now_iso() -> str:
    return datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")


class _HtmlTextSignalParser(HTMLParser):
    _IGNORED_TAGS = {"script", "style", "title", "head", "meta", "link", "noscript"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._tag_stack: list[str] = []
        self.segments: list[str] = []

    def handle_starttag(self, tag: str, attrs) -> None:
        self._tag_stack.append(tag.lower())

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        for index in range(len(self._tag_stack) - 1, -1, -1):
            if self._tag_stack[index] == tag:
                del self._tag_stack[index:]
                break

    def handle_data(self, data: str) -> None:
        if any(tag in self._IGNORED_TAGS for tag in self._tag_stack):
            return
        text = re.sub(r"\s+", " ", data).strip()
        if not text:
            return
        if not re.search(r"[0-9A-Za-z\u4e00-\u9fff]", text):
            return
        self.segments.append(text[:80])


def _estimate_html_text_signal(html: str) -> dict[str, Any]:
    parser = _HtmlTextSignalParser()
    parser.feed(html)
    parser.close()
    return {
        "text_node_count": len(parser.segments),
        "samples": parser.segments[:5],
    }


def _iter_shapes(shapes) -> Any:
    for shape in shapes:
        yield shape
        if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)


def _shape_text(shape) -> str:
    raw_text = getattr(shape, "text", "") or ""
    return re.sub(r"\s+", " ", raw_text).strip()


def _merge_status(*statuses: str | None) -> str:
    merged = "pass"
    for status in statuses:
        if status is None:
            continue
        if _STATUS_PRIORITY.get(status, 0) > _STATUS_PRIORITY.get(merged, 0):
            merged = status
    return merged


def _compute_dhash(image: Image.Image, size: int = 8) -> int:
    grayscale = image.convert("L").resize((size + 1, size), _RESAMPLE_LANCZOS)
    pixels = grayscale.load()
    digest = 0
    for row in range(size):
        for col in range(size):
            digest <<= 1
            digest |= 1 if pixels[col, row] > pixels[col + 1, row] else 0
    return digest


def _hamming_distance(left: int, right: int) -> int:
    return (left ^ right).bit_count()


def _compare_preview_images(source_png: Path, preview_png: Path) -> dict[str, Any]:
    with Image.open(source_png) as source_image, Image.open(preview_png) as preview_image:
        source_rgb = source_image.convert("RGB")
        preview_rgb = preview_image.convert("RGB")
        if preview_rgb.size != source_rgb.size:
            preview_rgb = preview_rgb.resize(source_rgb.size, _RESAMPLE_LANCZOS)

        compare_size = (320, 180)
        source_compare = source_rgb.resize(compare_size, _RESAMPLE_LANCZOS).filter(
            ImageFilter.GaussianBlur(1.2)
        )
        preview_compare = preview_rgb.resize(compare_size, _RESAMPLE_LANCZOS).filter(
            ImageFilter.GaussianBlur(1.2)
        )
        diff = ImageChops.difference(source_compare, preview_compare)
        stat = ImageStat.Stat(diff)
        mean_pixel_delta = sum(stat.mean) / (len(stat.mean) * 255.0)
        dhash_distance = _hamming_distance(
            _compute_dhash(source_compare),
            _compute_dhash(preview_compare),
        )

    reasons: list[str] = []
    if mean_pixel_delta >= 0.14 or (dhash_distance >= 20 and mean_pixel_delta >= 0.05):
        reasons.append(
            "powerpoint readback diverges visually "
            f"(dhash={dhash_distance}, mean_delta={mean_pixel_delta:.4f})"
        )
        status = "fail"
    elif dhash_distance >= 15 or mean_pixel_delta >= 0.09:
        reasons.append(
            "powerpoint readback shows noticeable drift "
            f"(dhash={dhash_distance}, mean_delta={mean_pixel_delta:.4f})"
        )
        status = "warn"
    else:
        status = "pass"

    return {
        "status": status,
        "reasons": reasons,
        "mean_pixel_delta": round(mean_pixel_delta, 4),
        "dhash_distance": dhash_distance,
    }


def _ensure_dom_preview_fallback(slides: list[dict[str, Any]]) -> None:
    for slide in slides:
        source_path = Path(slide["html_preview_png_path"])
        preview_path = Path(slide["preview_png_path"])
        if preview_path.exists():
            continue
        preview_path.write_bytes(source_path.read_bytes())


def _render_powerpoint_readback_previews(
    ppt_path: Path,
    preview_dir: Path,
    slides: list[dict[str, Any]],
) -> dict[str, Any]:
    support = detect_powerpoint_render_support()
    if not support["available"]:
        return {
            "available": False,
            "engine": None,
            "reason": support["reason"],
            "binary_path": support["binary_path"],
            "used_ascii_workspace": False,
            "slides": [],
            "summary": None,
            "failures": [],
        }

    try:
        render_result = export_powerpoint_slide_previews(
            ppt_path=ppt_path,
            out_dir=preview_dir,
            slide_count=len(slides),
            prefix="editable-preview",
        )
    except Exception as exc:
        return {
            "available": False,
            "engine": support["engine"],
            "reason": str(exc),
            "binary_path": support["binary_path"],
            "used_ascii_workspace": False,
            "slides": [],
            "summary": None,
            "failures": [],
        }

    results: list[dict[str, Any]] = []
    failures: list[str] = []
    for index, slide in enumerate(slides, start=1):
        preview_path = Path(slide["preview_png_path"])
        if not preview_path.exists():
            result = {
                "index": index,
                "status": "fail",
                "reasons": ["PowerPoint readback preview missing"],
                "mean_pixel_delta": None,
                "dhash_distance": None,
                "preview_png_path": str(preview_path),
            }
        else:
            comparison = _compare_preview_images(
                Path(slide["html_preview_png_path"]),
                preview_path,
            )
            result = {
                "index": index,
                "status": comparison["status"],
                "reasons": comparison["reasons"],
                "mean_pixel_delta": comparison["mean_pixel_delta"],
                "dhash_distance": comparison["dhash_distance"],
                "preview_png_path": str(preview_path),
            }
        results.append(result)
        if result["status"] == "fail":
            failures.append(
                f"slide {index:02d}: " + "; ".join(result["reasons"] or ["PowerPoint readback audit failed"])
            )

    summary = {
        "pass": sum(1 for item in results if item["status"] == "pass"),
        "warn": sum(1 for item in results if item["status"] == "warn"),
        "fail": sum(1 for item in results if item["status"] == "fail"),
    }
    return {
        "available": True,
        "engine": render_result["engine"],
        "reason": None,
        "binary_path": render_result["binary_path"],
        "used_ascii_workspace": render_result["used_ascii_workspace"],
        "slides": results,
        "summary": summary,
        "failures": failures,
    }


def _merge_audit_results(
    structure_audit: dict[str, Any],
    readback_audit: dict[str, Any],
) -> dict[str, Any]:
    readback_by_index = {
        int(item["index"]): item for item in readback_audit.get("slides", []) if "index" in item
    }
    results: list[dict[str, Any]] = []
    failures = list(structure_audit.get("failures", []))
    readback_available = bool(readback_audit.get("available"))
    audit_basis = "pptx-structure+powerpoint-readback" if readback_available else "pptx-structure"
    preview_basis = "powerpoint-readback" if readback_available else "dom-source"

    for slide_result in structure_audit["slides"]:
        index = int(slide_result["index"])
        readback_result = readback_by_index.get(index)
        combined_reasons = list(slide_result["reasons"])
        combined_status = slide_result["status"]
        readback_status = "skipped"
        readback_reasons: list[str] = []
        readback_mean_pixel_delta = None
        readback_dhash_distance = None
        preview_png_path = None

        if readback_result:
            combined_status = _merge_status(combined_status, readback_result["status"])
            combined_reasons.extend(readback_result["reasons"])
            readback_status = readback_result["status"]
            readback_reasons = list(readback_result["reasons"])
            readback_mean_pixel_delta = readback_result.get("mean_pixel_delta")
            readback_dhash_distance = readback_result.get("dhash_distance")
            preview_png_path = readback_result.get("preview_png_path")

        merged = {
            **slide_result,
            "status": combined_status,
            "reasons": combined_reasons,
            "structure_status": slide_result["status"],
            "structure_reasons": list(slide_result["reasons"]),
            "readback_status": readback_status,
            "readback_reasons": readback_reasons,
            "readback_mean_pixel_delta": readback_mean_pixel_delta,
            "readback_dhash_distance": readback_dhash_distance,
            "preview_png_path": preview_png_path,
            "preview_basis": preview_basis,
            "audit_basis": audit_basis,
            "readback_available": readback_available,
            "readback_error": readback_audit.get("reason"),
            "preview_engine": readback_audit.get("engine"),
            "powerpoint_binary_path": readback_audit.get("binary_path"),
            "powerpoint_used_ascii_workspace": readback_audit.get("used_ascii_workspace", False),
        }
        results.append(merged)
        if merged["status"] == "fail" and not any(
            failure.startswith(f"slide {index:02d}") for failure in failures
        ):
            failures.append(
                f"slide {index:02d} {merged['title']}: " + "; ".join(merged["reasons"] or ["audit failed"])
            )

    summary = {
        "pass": sum(1 for item in results if item["status"] == "pass"),
        "warn": sum(1 for item in results if item["status"] == "warn"),
        "fail": sum(1 for item in results if item["status"] == "fail"),
    }
    return {
        "slides": results,
        "failures": failures,
        "summary": summary,
        "preview_basis": preview_basis,
        "audit_basis": audit_basis,
        "readback_available": readback_available,
        "readback_reason": readback_audit.get("reason"),
        "readback_summary": readback_audit.get("summary"),
        "preview_engine": readback_audit.get("engine"),
        "powerpoint_binary_path": readback_audit.get("binary_path"),
        "powerpoint_used_ascii_workspace": readback_audit.get("used_ascii_workspace", False),
    }


def _audit_exported_pptx(ppt_path: Path, slides: list[dict[str, Any]]) -> dict[str, Any]:
    deck = Presentation(str(ppt_path))
    results: list[dict[str, Any]] = []
    failures: list[str] = []

    if len(deck.slides) != len(slides):
        failures.append(
            f"slide count mismatch: pptx={len(deck.slides)} html={len(slides)}"
        )

    for idx, slide_payload in enumerate(slides, start=1):
        ppt_slide = deck.slides[idx - 1] if idx - 1 < len(deck.slides) else None
        text_signal = _estimate_html_text_signal(slide_payload["html"])
        html_text_nodes = text_signal["text_node_count"]
        shape_count = 0
        text_shape_count = 0
        picture_shape_count = 0
        reasons: list[str] = []

        if ppt_slide is None:
            reasons.append("missing slide in exported pptx")
            status = "fail"
        else:
            all_shapes = list(_iter_shapes(ppt_slide.shapes))
            shape_count = len(all_shapes)
            for shape in all_shapes:
                if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE:
                    picture_shape_count += 1
                if getattr(shape, "has_text_frame", False) and _shape_text(shape):
                    text_shape_count += 1

            min_text_shapes = 0
            min_total_shapes = 0
            if html_text_nodes >= 10:
                min_text_shapes = max(3, math.ceil(html_text_nodes * 0.12))
            elif html_text_nodes >= 4:
                min_text_shapes = 1
            if html_text_nodes >= 12:
                min_total_shapes = max(4, math.ceil(html_text_nodes * 0.18))
            elif html_text_nodes >= 4:
                min_total_shapes = 2

            screenshot_like = (
                html_text_nodes >= 3
                and shape_count <= 2
                and picture_shape_count >= 1
                and text_shape_count <= 1
            )
            if screenshot_like:
                reasons.append("pptx structure looks rasterized")
            if min_text_shapes and text_shape_count < min_text_shapes:
                reasons.append(
                    f"text shapes too low ({text_shape_count} < {min_text_shapes})"
                )
            if min_total_shapes and shape_count < min_total_shapes:
                reasons.append(
                    f"shape count too low ({shape_count} < {min_total_shapes})"
                )

            if screenshot_like or (len(reasons) >= 2 and html_text_nodes >= 10):
                status = "fail"
            elif reasons:
                status = "warn"
            else:
                status = "pass"

        result = {
            "index": idx,
            "title": slide_payload.get("title", f"slide-{idx:02d}"),
            "status": status,
            "shape_count": shape_count,
            "text_shape_count": text_shape_count,
            "picture_shape_count": picture_shape_count,
            "html_text_node_count": html_text_nodes,
            "html_text_samples": text_signal["samples"],
            "reasons": reasons,
        }
        results.append(result)
        if status == "fail":
            failures.append(
                f"slide {idx:02d} {result['title']}: " + "; ".join(reasons or ["audit failed"])
            )

    summary = {
        "pass": sum(1 for item in results if item["status"] == "pass"),
        "warn": sum(1 for item in results if item["status"] == "warn"),
        "fail": sum(1 for item in results if item["status"] == "fail"),
    }
    return {"slides": results, "failures": failures, "summary": summary}


def _render_html_background_only_screenshot(html_path: Path) -> bytes:
    html = html_path.read_text(encoding="utf-8")
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch()
        page = browser.new_page(viewport={"width": 1280, "height": 720})
        try:
            page.set_content(html, wait_until="networkidle")
            page.add_style_tag(
                content="""
                .slide > * {
                  visibility: hidden !important;
                }
                """
            )
            page.wait_for_timeout(60)
            locator = page.locator(".slide").first
            if locator.count() == 0:
                return page.screenshot(type="png")
            return locator.screenshot(type="png")
        finally:
            browser.close()


def _replace_full_slide_background_picture(ppt_path: Path, slides: list[dict[str, Any]]) -> None:
    deck = Presentation(str(ppt_path))
    slide_width = int(deck.slide_width)
    slide_height = int(deck.slide_height)
    patched = False

    for idx, slide_payload in enumerate(slides, start=1):
        html = str(slide_payload.get("html") or "")
        if "radial-gradient(" not in html:
            continue

        ppt_slide = deck.slides[idx - 1] if idx - 1 < len(deck.slides) else None
        if ppt_slide is None:
            continue

        target_picture = None
        for shape in ppt_slide.shapes:
            if getattr(shape, "shape_type", None) != MSO_SHAPE_TYPE.PICTURE:
                continue
            if abs(int(shape.left)) > 2000 or abs(int(shape.top)) > 2000:
                continue
            if int(shape.width) < int(slide_width * 0.98) or int(shape.height) < int(slide_height * 0.98):
                continue
            target_picture = shape
            break

        if target_picture is None:
            continue

        background_png = _render_html_background_only_screenshot(Path(slide_payload["html_path"]))
        sp_tree = ppt_slide.shapes._spTree
        target_index = list(sp_tree).index(target_picture._element)
        sp_tree.remove(target_picture._element)

        replacement = ppt_slide.shapes.add_picture(
            io.BytesIO(background_png),
            left=0,
            top=0,
            width=deck.slide_width,
            height=deck.slide_height,
        )
        sp_tree.remove(replacement._element)
        sp_tree.insert(target_index, replacement._element)
        patched = True

    if patched:
        deck.save(str(ppt_path))


def _write_dom_review(out_dir: Path, slide_result: dict[str, Any], html_path: Path) -> Path:
    review_path = out_dir / f"review-{slide_result['index']:02d}.md"
    status_label = {
        "pass": "PASS",
        "warn": "WARN",
        "fail": "FAIL",
    }[slide_result["status"]]
    lines = [
        f"# {html_path.name}",
        "",
        f"RESULT: {status_label}",
        "",
        "## Reasons",
    ]
    lines.extend([f"- {item}" for item in slide_result["reasons"]] or ["- DOM export passed all enabled audits"])
    lines.extend(
        [
            "",
            "## Audit Layers",
            f"- structure_status: {slide_result['structure_status']}",
            f"- readback_status: {slide_result['readback_status']}",
            f"- preview_basis: {slide_result['preview_basis']}",
            f"- audit_basis: {slide_result['audit_basis']}",
            "",
            "## Metrics",
            f"- shape_count: {slide_result['shape_count']}",
            f"- text_shape_count: {slide_result['text_shape_count']}",
            f"- picture_shape_count: {slide_result['picture_shape_count']}",
            f"- html_text_node_count: {slide_result['html_text_node_count']}",
            f"- readback_mean_pixel_delta: {slide_result.get('readback_mean_pixel_delta')}",
            f"- readback_dhash_distance: {slide_result.get('readback_dhash_distance')}",
            "",
            "## HTML Text Samples",
        ]
    )
    lines.extend([f"- {item}" for item in slide_result["html_text_samples"]] or ["- no sampled text nodes"])
    if slide_result.get("readback_reasons"):
        lines.extend(["", "## Readback Notes"])
        lines.extend([f"- {item}" for item in slide_result["readback_reasons"]])
    elif slide_result.get("readback_error"):
        lines.extend(["", "## Readback Notes", f"- {slide_result['readback_error']}"])
    review_path.write_text("\n".join(lines), encoding="utf-8")
    return review_path


class DomPptxExporter:
    """Run the browser-side DOM exporter inside Playwright."""

    def __init__(self, bundle_path: Path | None = None):
        self.bundle_path = bundle_path or (Path(__file__).resolve().parent / "dom-to-pptx.bundle.js")

    def export_slides(self, slides: list[dict[str, Any]], output_path: Path) -> Path:
        if not slides:
            raise ValueError("No slides provided for DOM export")

        shell_html = """<!DOCTYPE html><html><head><meta charset='utf-8'><title>dom export</title></head><body></body></html>"""
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch()
            page = browser.new_page(viewport={"width": 1400, "height": 900})
            page.set_content(shell_html, wait_until="domcontentloaded")
            page.add_script_tag(path=str(self.bundle_path))
            payload = {"slides": slides, "fileName": output_path.name}
            b64 = page.evaluate(
                """
                async (payload) => {
                  const host = document.createElement('div');
                  host.style.cssText = 'position:fixed;left:-9999px;top:-9999px;opacity:0;pointer-events:none;';
                  document.body.appendChild(host);

                  function stabilizeLayoutForPpt(doc) {
                    const win = doc.defaultView;
                    if (!win) return;

                    const clampPx = (value, fallback = 0) => {
                      const parsed = Number.parseFloat(String(value ?? ''));
                      return Number.isFinite(parsed) ? parsed : fallback;
                    };

                    const extractLastLinearGradient = (value) => {
                      const input = String(value || '');
                      const lower = input.toLowerCase();
                      const token = 'linear-gradient(';
                      const start = lower.lastIndexOf(token);
                      if (start < 0) return '';

                      let depth = 1;
                      let quote = '';
                      for (let index = start + token.length; index < input.length; index++) {
                        const ch = input[index];
                        if (quote) {
                          if (ch === quote && input[index - 1] !== '\\\\') {
                            quote = '';
                          }
                          continue;
                        }
                        if (ch === '"' || ch === "'") {
                          quote = ch;
                          continue;
                        }
                        if (ch === '(') {
                          depth += 1;
                          continue;
                        }
                        if (ch === ')') {
                          depth -= 1;
                          if (depth === 0) {
                            return input.slice(start, index + 1);
                          }
                        }
                      }
                      return '';
                    };

                    const extractRadialGlowSpecs = (value) => {
                      const input = String(value || '');
                      const specs = [];
                      const pattern = /radial-gradient\\(\\s*circle at\\s*([^,]+),\\s*(rgba?\\([^\\)]+\\)|#[0-9a-fA-F]+)\\s*,\\s*transparent\\s+([0-9.]+)%\\s*\\)/gi;
                      let match = null;
                      while ((match = pattern.exec(input))) {
                        specs.push({
                          position: String(match[1] || '').trim().toLowerCase(),
                          color: String(match[2] || '').trim(),
                          stop: clampPx(match[3], 30),
                        });
                      }
                      return specs.slice(0, 3);
                    };

                    const parseCssColorAlpha = (value) => {
                      const raw = String(value || '').trim();
                      const rgbaMatch = raw.match(/^rgba?\\((\\d+),\\s*(\\d+),\\s*(\\d+)(?:,\\s*([\\d.]+))?\\)$/i);
                      if (rgbaMatch) {
                        return {
                          color: `rgb(${rgbaMatch[1]}, ${rgbaMatch[2]}, ${rgbaMatch[3]})`,
                          alpha: rgbaMatch[4] != null ? Math.max(0, Math.min(Number.parseFloat(rgbaMatch[4]), 1)) : 1,
                        };
                      }
                      return { color: raw || '#FFFFFF', alpha: 1 };
                    };

                    const buildGlowOverlayDataUrl = (specs) => {
                      if (!Array.isArray(specs) || !specs.length) {
                        return '';
                      }

                      const defs = [];
                      const bodies = [];
                      specs.forEach((spec, index) => {
                        const baseSize = Math.max(1280, 720);
                        const size = Math.max(
                          260,
                          Math.round(baseSize * Math.max(0.28, Math.min((spec.stop || 30) / 100 * 1.45, 0.64)))
                        );
                        const blur = Math.max(60, Math.round(size * 0.2));
                        const innerSize = Math.max(180, Math.round(size * 0.62));
                        const pos = String(spec.position || '').toLowerCase();
                        let outerX = -size * 0.16;
                        let outerY = -size * 0.16;
                        if (pos.includes('100% 100%') || pos.includes('bottom right') || pos.includes('right bottom')) {
                          outerX = 1280 - size * 0.84;
                          outerY = 720 - size * 0.84;
                        } else if (pos.includes('100% 0%') || pos.includes('top right') || pos.includes('right top')) {
                          outerX = 1280 - size * 0.84;
                          outerY = -size * 0.16;
                        } else if (pos.includes('0% 100%') || pos.includes('bottom left') || pos.includes('left bottom')) {
                          outerX = -size * 0.16;
                          outerY = 720 - size * 0.84;
                        }
                        const innerX = outerX + Math.round((size - innerSize) * 0.5);
                        const innerY = outerY + Math.round((size - innerSize) * 0.5);
                        const colorInfo = parseCssColorAlpha(spec.color);
                        const outerOpacity = Math.max(0.04, Math.min(colorInfo.alpha * 0.95, 0.10));
                        const innerOpacity = Math.max(0.06, Math.min(colorInfo.alpha * 1.35, 0.15));
                        const filterId = `pptGlow${index}`;
                        defs.push(
                          `<filter id="${filterId}" x="-50%" y="-50%" width="200%" height="200%">` +
                          `<feGaussianBlur stdDeviation="${Math.max(16, Math.round(blur * 0.25))}" />` +
                          `</filter>`
                        );
                        bodies.push(
                          `<ellipse cx="${Math.round(outerX + size / 2)}" cy="${Math.round(outerY + size / 2)}" rx="${Math.round(size / 2)}" ry="${Math.round(size / 2)}" fill="${colorInfo.color}" fill-opacity="${outerOpacity.toFixed(3)}" filter="url(#${filterId})"/>`
                        );
                        bodies.push(
                          `<ellipse cx="${Math.round(innerX + innerSize / 2)}" cy="${Math.round(innerY + innerSize / 2)}" rx="${Math.round(innerSize / 2)}" ry="${Math.round(innerSize / 2)}" fill="${colorInfo.color}" fill-opacity="${innerOpacity.toFixed(3)}"/>`
                        );
                      });

                      const svg =
                        `<svg xmlns="http://www.w3.org/2000/svg" width="1280" height="720" viewBox="0 0 1280 720">` +
                        `<defs>${defs.join('')}</defs>` +
                        bodies.join('') +
                        `</svg>`;
                      return `data:image/svg+xml;charset=utf-8,${encodeURIComponent(svg)}`;
                    };

                    const ensureRootBackdrop = () => {
                      const htmlEl = doc.documentElement;
                      const body = doc.body;
                      if (!(htmlEl instanceof win.HTMLElement) || !(body instanceof win.HTMLElement)) {
                        return;
                      }

                      htmlEl.style.width = '1280px';
                      htmlEl.style.height = '720px';
                      htmlEl.style.margin = '0';
                      htmlEl.style.overflow = 'hidden';
                      body.style.width = '1280px';
                      body.style.height = '720px';
                      body.style.minHeight = '720px';
                      body.style.margin = '0';
                      body.style.overflow = 'hidden';
                      if (!body.style.position || body.style.position === 'static') {
                        body.style.position = 'relative';
                      }

                      const htmlStyle = win.getComputedStyle(htmlEl);
                      const bodyStyle = win.getComputedStyle(body);
                      const transparentColor = 'rgba(0, 0, 0, 0)';
                      const bgSourceImage =
                        bodyStyle.backgroundImage && bodyStyle.backgroundImage !== 'none'
                          ? bodyStyle.backgroundImage
                          : htmlStyle.backgroundImage;
                      const bgSourceColor =
                        bodyStyle.backgroundColor && bodyStyle.backgroundColor !== transparentColor
                          ? bodyStyle.backgroundColor
                          : htmlStyle.backgroundColor;
                      const linearGradient = extractLastLinearGradient(bgSourceImage);
                      const radialSpecs = extractRadialGlowSpecs(bgSourceImage);
                      const hasRenderableBackdrop =
                        !!linearGradient ||
                        (!!bgSourceColor && bgSourceColor !== transparentColor) ||
                        radialSpecs.length > 0;
                      if (!hasRenderableBackdrop) {
                        return;
                      }

                      let backdrop = Array.from(body.children).find(
                        (child) =>
                          child instanceof win.HTMLElement &&
                          child.getAttribute('data-ppt-root-backdrop') === 'true'
                      );
                      if (!(backdrop instanceof win.HTMLElement)) {
                        backdrop = doc.createElement('div');
                        backdrop.setAttribute('data-ppt-root-backdrop', 'true');
                        body.insertBefore(backdrop, body.firstChild);
                      }

                      backdrop.style.position = 'absolute';
                      backdrop.style.left = '0';
                      backdrop.style.top = '0';
                      backdrop.style.width = '1280px';
                      backdrop.style.height = '720px';
                      backdrop.style.pointerEvents = 'none';
                      backdrop.style.overflow = 'hidden';
                      backdrop.style.zIndex = '0';
                      backdrop.style.borderRadius = '0';
                      backdrop.style.backgroundColor =
                        bgSourceColor && bgSourceColor !== transparentColor ? bgSourceColor : 'transparent';
                      backdrop.style.backgroundImage = linearGradient || 'none';
                      backdrop.style.backgroundRepeat = 'no-repeat';
                      backdrop.style.backgroundSize = '100% 100%';
                      backdrop.style.backgroundPosition = 'center';

                      const glowOverlayUrl = buildGlowOverlayDataUrl(radialSpecs);
                      Array.from(body.children).forEach((child) => {
                        if (
                          child instanceof win.HTMLElement &&
                          child.getAttribute('data-ppt-root-glow-image') === 'true'
                        ) {
                          child.remove();
                        }
                      });
                      if (glowOverlayUrl) {
                        const glowImage = doc.createElement('img');
                        glowImage.setAttribute('data-ppt-root-glow-image', 'true');
                        glowImage.src = glowOverlayUrl;
                        glowImage.alt = '';
                        glowImage.style.position = 'absolute';
                        glowImage.style.left = '0';
                        glowImage.style.top = '0';
                        glowImage.style.width = '1280px';
                        glowImage.style.height = '720px';
                        glowImage.style.pointerEvents = 'none';
                        glowImage.style.zIndex = '0';
                        body.insertBefore(glowImage, backdrop.nextSibling);
                      }

                      Array.from(body.children).forEach((child) => {
                        if (
                          !(child instanceof win.HTMLElement) ||
                          child === backdrop ||
                          child.getAttribute('data-ppt-root-glow-image') === 'true'
                        ) {
                          return;
                        }
                        const childStyle = win.getComputedStyle(child);
                        if (childStyle.position === 'static') {
                          child.style.position = 'relative';
                        }
                        const childZIndex = Number.parseInt(childStyle.zIndex, 10);
                        if (!Number.isFinite(childZIndex) || childZIndex < 1) {
                          child.style.zIndex = '1';
                        }
                      });

                      body.style.backgroundImage = 'none';
                      body.style.backgroundColor = 'transparent';
                      htmlEl.style.backgroundImage = 'none';
                      htmlEl.style.backgroundColor = 'transparent';
                    };

                    const ensureSlideBackdrop = () => {
                      const slide = doc.querySelector('.slide');
                      if (!(slide instanceof win.HTMLElement)) {
                        return;
                      }

                      const slideStyle = win.getComputedStyle(slide);
                      const slideBgImage = slideStyle.backgroundImage;
                      if (!slideBgImage || slideBgImage === 'none') {
                        return;
                      }

                      const linearGradient = extractLastLinearGradient(slideBgImage);
                      const radialSpecs = extractRadialGlowSpecs(slideBgImage);

                      Array.from(slide.children).forEach((child) => {
                        if (
                          child instanceof win.HTMLElement &&
                          child.getAttribute('data-ppt-slide-glow-image') === 'true'
                        ) {
                          child.remove();
                        }
                      });

                      if (!slide.style.position || slide.style.position === 'static') {
                        slide.style.position = 'relative';
                      }
                      slide.style.backgroundImage = linearGradient || 'none';
                      slide.style.backgroundRepeat = 'no-repeat';
                      slide.style.backgroundSize = '100% 100%';
                      slide.style.backgroundPosition = 'center';

                      const glowOverlayUrl = buildGlowOverlayDataUrl(radialSpecs);
                      if (glowOverlayUrl) {
                        const glowImage = doc.createElement('img');
                        glowImage.setAttribute('data-ppt-slide-glow-image', 'true');
                        glowImage.src = glowOverlayUrl;
                        glowImage.alt = '';
                        glowImage.style.position = 'absolute';
                        glowImage.style.left = '0';
                        glowImage.style.top = '0';
                        glowImage.style.width = '1280px';
                        glowImage.style.height = '720px';
                        glowImage.style.pointerEvents = 'none';
                        glowImage.style.zIndex = '0';
                        slide.insertBefore(glowImage, slide.firstChild);
                      }

                      Array.from(slide.children).forEach((child) => {
                        if (
                          !(child instanceof win.HTMLElement) ||
                          child.getAttribute('data-ppt-slide-glow-image') === 'true'
                        ) {
                          return;
                        }
                        const childStyle = win.getComputedStyle(child);
                        if (childStyle.position === 'static') {
                          child.style.position = 'relative';
                        }
                        const childZIndex = Number.parseInt(childStyle.zIndex, 10);
                        if (!Number.isFinite(childZIndex) || childZIndex < 1) {
                          child.style.zIndex = '1';
                        }
                      });
                    };

                    const stabilizeHeader = (header) => {
                      const headerStyle = win.getComputedStyle(header);
                      if (headerStyle.display !== 'flex') return;

                      const children = Array.from(header.children).filter(
                        (node) => node instanceof win.HTMLElement
                      );
                      if (children.length < 2) return;

                      const lead = children[0];
                      const tail = children[children.length - 1];
                      const headerRect = header.getBoundingClientRect();
                      const tailRect = tail.getBoundingClientRect();
                      if (!headerRect.width) return;

                      // Flex headers with `space-between` often keep the text column at max-content width.
                      // That is stable in Chromium, but PowerPoint may reflow with different font metrics.
                      const reserve = Math.max(110, tailRect.width + 36);
                      const available = Math.max(360, Math.floor(headerRect.width - reserve));

                      lead.style.flex = `0 0 ${available}px`;
                      lead.style.width = `${available}px`;
                      lead.style.maxWidth = `${available}px`;
                      lead.style.minWidth = '0';
                      if (!lead.style.paddingRight) {
                        lead.style.paddingRight = '24px';
                      }

                      tail.style.flex = '0 0 auto';
                      tail.style.marginLeft = '24px';
                      tail.style.alignSelf = 'flex-start';
                      tail.style.maxWidth = `${Math.max(132, Math.floor(headerRect.width * 0.34))}px`;
                      tail.style.minWidth = tail.style.minWidth || '92px';

                      if (tail.classList.contains('corner-cluster') || tail.querySelector('.page-tag, .meta-chip, [class*="chip"], [class*="tag"]')) {
                        tail.style.display = 'grid';
                        tail.style.justifyItems = 'end';
                        tail.style.rowGap = tail.style.rowGap || '10px';
                      }

                      const textNodes = Array.from(
                        lead.querySelectorAll('h1, h2, h3, p, .title, .subtitle')
                      ).filter((node) => node instanceof win.HTMLElement);

                      const tailNodes = Array.from(
                        tail.querySelectorAll('.page-tag, .meta-chip, [class*="chip"], [class*="tag"], [class*="badge"]')
                      ).filter((node) => node instanceof win.HTMLElement);

                      textNodes.forEach((node) => {
                        node.style.width = '100%';
                        node.style.maxWidth = '100%';
                        node.style.minWidth = '0';

                        const computed = win.getComputedStyle(node);
                        const fontSize = clampPx(computed.fontSize);
                        const lineHeight = clampPx(computed.lineHeight, fontSize * 1.08);

                        if (node.matches('h1, .title')) {
                          node.style.whiteSpace = 'normal';
                          node.style.wordBreak = 'break-word';
                          node.style.lineHeight = `${Math.max(lineHeight, fontSize * 1.1)}px`;

                          const letterSpacing = clampPx(computed.letterSpacing);
                          if (letterSpacing < -0.6) {
                            node.style.letterSpacing = '-0.4px';
                          }
                        } else if (node.matches('p, .subtitle')) {
                          node.style.lineHeight = `${Math.max(lineHeight, fontSize * 1.2)}px`;
                        }
                      });

                      tailNodes.forEach((node) => {
                        node.style.whiteSpace = 'nowrap';
                        node.style.minWidth = node.classList.contains('page-tag') ? '68px' : (node.style.minWidth || '0');
                        node.style.paddingLeft = node.style.paddingLeft || '10px';
                        node.style.paddingRight = node.style.paddingRight || '10px';
                        if (node.classList.contains('page-tag')) {
                          node.style.display = 'inline-flex';
                          node.style.alignItems = 'center';
                          node.style.justifyContent = 'center';
                        }
                      });
                    };

                    Array.from(doc.querySelectorAll('header')).forEach((header) => {
                      if (header instanceof win.HTMLElement) {
                        stabilizeHeader(header);
                      }
                    });

                    ensureRootBackdrop();
                    ensureSlideBackdrop();
                  }

                  function canParseCssColor(value) {
                    if (!value) return false;
                    try {
                      const cvs = document.createElement('canvas');
                      cvs.width = cvs.height = 1;
                      const ctx = cvs.getContext('2d');
                      if (!ctx) return false;
                      ctx.fillStyle = '#000';
                      ctx.fillStyle = value;
                      return !!ctx.fillStyle;
                    } catch (_) {
                      return false;
                    }
                  }

                  function toHexFallback(colorValue) {
                    if (!canParseCssColor(colorValue)) return colorValue;
                    try {
                      const cvs = document.createElement('canvas');
                      cvs.width = cvs.height = 1;
                      const ctx = cvs.getContext('2d');
                      if (!ctx) return colorValue;
                      ctx.fillStyle = '#000';
                      ctx.fillStyle = colorValue;
                      return ctx.fillStyle || colorValue;
                    } catch (_) {
                      return colorValue;
                    }
                  }

                  function convertModernColors(rootEl) {
                    if (!rootEl) return;
                    const modernColorRe = /\\b(oklch|oklab|lch|lab|color)\\s*\\(/i;
                    const colorCssProps = [
                      'color', 'background-color', 'border-color',
                      'border-top-color', 'border-right-color', 'border-bottom-color', 'border-left-color',
                      'outline-color', 'text-decoration-color', 'caret-color', 'column-rule-color',
                      'fill', 'stroke', 'stop-color', 'flood-color', 'lighting-color'
                    ];
                    const colorJsProps = [
                      'color', 'backgroundColor', 'borderColor',
                      'borderTopColor', 'borderRightColor', 'borderBottomColor', 'borderLeftColor',
                      'outlineColor', 'textDecorationColor', 'caretColor', 'columnRuleColor',
                      'fill', 'stroke', 'stopColor', 'floodColor', 'lightingColor'
                    ];

                    rootEl.querySelectorAll('style').forEach((styleEl) => {
                      if (!modernColorRe.test(styleEl.textContent || '')) return;
                      styleEl.textContent = styleEl.textContent.replace(
                        /(oklch|oklab|lch|lab|color)\\([^)]*\\)/gi,
                        (match) => toHexFallback(match)
                      );
                    });

                    const walk = [rootEl, ...Array.from(rootEl.querySelectorAll('*'))];
                    for (const el of walk) {
                      if (!el || !el.style || el.tagName === 'STYLE' || el.tagName === 'SCRIPT') continue;

                      let cs;
                      try {
                        cs = window.getComputedStyle(el);
                      } catch (_) {
                        continue;
                      }

                      for (let i = 0; i < colorCssProps.length; i++) {
                        const val = cs.getPropertyValue(colorCssProps[i]);
                        if (val && modernColorRe.test(val)) {
                          el.style[colorJsProps[i]] = toHexFallback(val);
                        }
                      }

                      const bgImg = cs.getPropertyValue('background-image');
                      if (bgImg && modernColorRe.test(bgImg)) {
                        el.style.backgroundImage = bgImg.replace(
                          /(oklch|oklab|lch|lab|color)\\([^)]*\\)/gi,
                          (match) => toHexFallback(match)
                        );
                      }

                      for (const shadowProp of ['box-shadow', 'text-shadow']) {
                        const shadowValue = cs.getPropertyValue(shadowProp);
                        if (!shadowValue || !modernColorRe.test(shadowValue)) continue;
                        const normalized = shadowValue.replace(
                          /(oklch|oklab|lch|lab|color)\\([^)]*\\)/gi,
                          (match) => toHexFallback(match)
                        );
                        el.style[shadowProp === 'box-shadow' ? 'boxShadow' : 'textShadow'] = normalized;
                      }
                    }
                  }

                  function extractFirstUrlFromBackgroundImage(bgValue, sourceWindow) {
                    if (!bgValue || !/url\\s*\\(/i.test(bgValue)) return null;
                    const match = /url\\s*\\(\\s*(['"]?)(.*?)\\1\\s*\\)/i.exec(bgValue);
                    if (!match || !match[2]) return null;
                    const rawUrl = match[2].trim();
                    if (!rawUrl) return null;
                    try {
                      return new URL(rawUrl, sourceWindow.location.href).href;
                    } catch (_) {
                      return rawUrl;
                    }
                  }

                  function inferObjectFitFromBackgroundSize(bgSize) {
                    const value = String(bgSize || '').toLowerCase();
                    if (value.includes('contain')) return 'contain';
                    if (value.includes('cover')) return 'cover';
                    if (value.includes('100% 100%') || value.includes('100%')) return 'fill';
                    return 'cover';
                  }

                  function materializeBackgroundImagesForExport(sourceRoot, clonedRoot, sourceWindow) {
                    if (!sourceRoot || !clonedRoot || !sourceWindow) return;
                    const sourceNodes = [sourceRoot, ...Array.from(sourceRoot.querySelectorAll('*'))];
                    const clonedNodes = [clonedRoot, ...Array.from(clonedRoot.querySelectorAll('*'))];
                    const pairCount = Math.min(sourceNodes.length, clonedNodes.length);

                    for (let i = 0; i < pairCount; i++) {
                      const src = sourceNodes[i];
                      const dst = clonedNodes[i];
                      if (!src || !dst || !dst.style) continue;

                      let srcStyle;
                      try {
                        srcStyle = sourceWindow.getComputedStyle(src);
                      } catch (_) {
                        continue;
                      }
                      const bgImage = srcStyle.getPropertyValue('background-image');
                      if (!bgImage || !/url\\s*\\(/i.test(bgImage)) continue;
                      if (bgImage.includes('gradient(') && bgImage.includes('url(')) continue;

                      const hasChildren = dst.children && dst.children.length > 0;
                      const hasText = !!(dst.textContent && dst.textContent.trim());
                      if (hasChildren || hasText) continue;

                      const imageUrl = extractFirstUrlFromBackgroundImage(bgImage, sourceWindow);
                      if (!imageUrl) continue;

                      const img = document.createElement('img');
                      img.src = imageUrl;
                      img.alt = '';
                      img.setAttribute('data-export-bg-image', 'true');
                      img.style.width = '100%';
                      img.style.height = '100%';
                      img.style.display = 'block';
                      img.style.objectFit = inferObjectFitFromBackgroundSize(srcStyle.getPropertyValue('background-size'));
                      img.style.objectPosition = srcStyle.getPropertyValue('background-position') || '50% 50%';

                      dst.style.setProperty('background-image', 'none');
                      dst.style.setProperty('background', 'none');
                      dst.appendChild(img);
                    }
                  }

                  function replaceClonedCanvasesWithImages(sourceDoc, clonedRoot) {
                    if (!sourceDoc || !clonedRoot) return;

                    const sourceCanvases = Array.from(sourceDoc.querySelectorAll('canvas'));
                    const clonedCanvases = Array.from(clonedRoot.querySelectorAll('canvas'));
                    const pairCount = Math.min(sourceCanvases.length, clonedCanvases.length);

                    for (let i = 0; i < pairCount; i++) {
                      const sourceCanvas = sourceCanvases[i];
                      const clonedCanvas = clonedCanvases[i];
                      if (!sourceCanvas || !clonedCanvas || !clonedCanvas.parentNode) continue;

                      try {
                        const dataUrl = sourceCanvas.toDataURL('image/png');
                        if (!dataUrl || dataUrl.length < 128) continue;

                        const img = document.createElement('img');
                        img.src = dataUrl;
                        img.alt = '';
                        img.className = clonedCanvas.className || '';
                        img.style.cssText = clonedCanvas.getAttribute('style') || '';
                        if (!img.style.width) img.style.width = (sourceCanvas.style.width || sourceCanvas.width + 'px');
                        if (!img.style.height) img.style.height = (sourceCanvas.style.height || sourceCanvas.height + 'px');
                        if (!img.style.display) img.style.display = 'block';

                        clonedCanvas.parentNode.replaceChild(img, clonedCanvas);
                      } catch (_) {
                        // Ignore tainted/unsupported canvas and keep original clone.
                      }
                    }
                  }

                  const EXPORT_COMPUTED_STYLE_PROPS = [
                    'display', 'position', 'top', 'right', 'bottom', 'left', 'z-index',
                    'width', 'height', 'min-width', 'min-height', 'max-width', 'max-height',
                    'margin', 'margin-top', 'margin-right', 'margin-bottom', 'margin-left',
                    'padding', 'padding-top', 'padding-right', 'padding-bottom', 'padding-left',
                    'box-sizing', 'overflow', 'overflow-x', 'overflow-y',
                    'transform', 'transform-origin', 'opacity',
                    'font-family', 'font-size', 'font-weight', 'font-style', 'line-height', 'letter-spacing',
                    'text-align', 'text-transform', 'text-decoration', 'white-space', 'word-break',
                    'color', 'background', 'background-color', 'background-image', 'background-size', 'background-position', 'background-repeat',
                    'border', 'border-top', 'border-right', 'border-bottom', 'border-left', 'border-radius',
                    'box-shadow', 'filter', 'backdrop-filter',
                    'align-items', 'align-content', 'justify-content', 'justify-items',
                    'flex', 'flex-direction', 'flex-wrap', 'flex-grow', 'flex-shrink', 'flex-basis', 'gap',
                    'grid-template-columns', 'grid-template-rows', 'grid-column', 'grid-row',
                    'object-fit', 'object-position'
                  ];

                  function copyComputedStylesForExport(sourceRoot, clonedRoot, sourceWindow) {
                    if (!sourceRoot || !clonedRoot || !sourceWindow) return;

                    const sourceNodes = [sourceRoot, ...Array.from(sourceRoot.querySelectorAll('*'))];
                    const clonedNodes = [clonedRoot, ...Array.from(clonedRoot.querySelectorAll('*'))];
                    const pairCount = Math.min(sourceNodes.length, clonedNodes.length);

                    for (let i = 0; i < pairCount; i++) {
                      const sourceNode = sourceNodes[i];
                      const clonedNode = clonedNodes[i];
                      if (!sourceNode || !clonedNode || !clonedNode.style) continue;
                      if (clonedNode.tagName === 'SCRIPT' || clonedNode.tagName === 'STYLE') continue;

                      try {
                        const computed = sourceWindow.getComputedStyle(sourceNode);
                        for (const prop of EXPORT_COMPUTED_STYLE_PROPS) {
                          const value = computed.getPropertyValue(prop);
                          if (value) clonedNode.style.setProperty(prop, value);
                        }
                      } catch (_) {}
                    }
                  }

                  function parseUniformScaleFromTransform(transformValue) {
                    const raw = String(transformValue || '').trim();
                    if (!raw || raw === 'none') return null;

                    const matrixMatch = raw.match(/^matrix\\(([^)]+)\\)$/i);
                    if (matrixMatch) {
                      const vals = matrixMatch[1].split(',').map((value) => parseFloat(value.trim()));
                      if (vals.length >= 6 && vals.every((value) => Number.isFinite(value))) {
                        const [a, b, c, d, e, f] = vals;
                        if (Math.abs(b) < 1e-4 && Math.abs(c) < 1e-4 && Math.abs(a - d) < 1e-3 && Math.abs(e) < 0.5 && Math.abs(f) < 0.5) {
                          return a;
                        }
                      }
                    }

                    const matrix3dMatch = raw.match(/^matrix3d\\(([^)]+)\\)$/i);
                    if (matrix3dMatch) {
                      const vals = matrix3dMatch[1].split(',').map((value) => parseFloat(value.trim()));
                      if (vals.length >= 16 && vals.every((value) => Number.isFinite(value))) {
                        const sx = vals[0];
                        const sy = vals[5];
                        const tx = vals[12];
                        const ty = vals[13];
                        if (Math.abs(sx - sy) < 1e-3 && Math.abs(tx) < 0.5 && Math.abs(ty) < 0.5) {
                          return sx;
                        }
                      }
                    }

                    const scaleMatch = raw.match(/^scale\\(\\s*([-\\d.]+)(?:\\s*,\\s*([-\\d.]+))?\\s*\\)$/i);
                    if (scaleMatch) {
                      const sx = parseFloat(scaleMatch[1]);
                      const sy = scaleMatch[2] ? parseFloat(scaleMatch[2]) : sx;
                      if (Number.isFinite(sx) && Number.isFinite(sy) && Math.abs(sx - sy) < 1e-3) {
                        return sx;
                      }
                    }

                    return null;
                  }

                  function neutralizeViewportFitScaleForExport(sourceRoot, clonedRoot, sourceWindow) {
                    if (!sourceRoot || !clonedRoot || !sourceWindow) return;

                    const sourceNodes = [sourceRoot, ...Array.from(sourceRoot.querySelectorAll('*'))];
                    const clonedNodes = [clonedRoot, ...Array.from(clonedRoot.querySelectorAll('*'))];
                    const pairCount = Math.min(sourceNodes.length, clonedNodes.length);

                    for (let i = 0; i < pairCount; i++) {
                      const src = sourceNodes[i];
                      const dst = clonedNodes[i];
                      if (!src || !dst || !dst.style) continue;

                      let cs;
                      try {
                        cs = sourceWindow.getComputedStyle(src);
                      } catch (_) {
                        continue;
                      }
                      if (!cs) continue;

                      const uniformScale = parseUniformScaleFromTransform(cs.transform);
                      if (!Number.isFinite(uniformScale) || Math.abs(uniformScale - 1) < 0.01) continue;

                      const widthPx = parseFloat(cs.width) || src.getBoundingClientRect().width;
                      const heightPx = parseFloat(cs.height) || src.getBoundingClientRect().height;
                      if (!(widthPx > 900 && heightPx > 500)) continue;
                      const ratio = widthPx / Math.max(1, heightPx);
                      if (!(ratio > 1.7 && ratio < 1.8)) continue;

                      let parentLooksViewportFitter = false;
                      const parent = src.parentElement;
                      if (parent) {
                        try {
                          const ps = sourceWindow.getComputedStyle(parent);
                          parentLooksViewportFitter = String(ps.display || '').includes('flex') &&
                            String(ps.justifyContent || '').includes('center') &&
                            String(ps.alignItems || '').includes('center');
                        } catch (_) {}
                      }

                      const idLikeSlideRoot = !!(src.id && /^(slide|ppt|page|slide-container|ppt-page|page-root)$/i.test(src.id));
                      if (!parentLooksViewportFitter && !idLikeSlideRoot) continue;

                      dst.style.setProperty('transform', 'none', 'important');
                      dst.style.setProperty('transform-origin', 'center center', 'important');
                    }
                  }

                  async function waitForStylesheetsReadyInContainer(container, timeoutMs = 2200) {
                    if (!container || !container.querySelectorAll) return;
                    const links = Array.from(container.querySelectorAll('link[rel="stylesheet"]'));
                    if (links.length === 0) return;

                    const waits = links.map((link) => new Promise((resolve) => {
                      if (link.sheet) return resolve(true);
                      const done = () => {
                        link.removeEventListener('load', done);
                        link.removeEventListener('error', done);
                        resolve(true);
                      };
                      link.addEventListener('load', done, { once: true });
                      link.addEventListener('error', done, { once: true });
                      setTimeout(done, timeoutMs);
                    }));
                    await Promise.allSettled(waits);
                  }

                  async function waitForImagesReadyInContainer(container, timeoutMs = 2600) {
                    if (!container || !container.querySelectorAll) return;
                    const images = Array.from(container.querySelectorAll('img'));
                    const pending = images.filter((img) => !img.complete);
                    if (pending.length === 0) return;

                    await Promise.race([
                      Promise.allSettled(
                        pending.map((img) => new Promise((resolve) => {
                          const done = () => {
                            img.removeEventListener('load', done);
                            img.removeEventListener('error', done);
                            resolve(true);
                          };
                          img.addEventListener('load', done, { once: true });
                          img.addEventListener('error', done, { once: true });
                        }))
                      ),
                      new Promise((resolve) => setTimeout(resolve, timeoutMs))
                    ]);
                  }

                  function buildExportContainerFromDocument(sourceDoc, container) {
                    try {
                      if (!sourceDoc || !sourceDoc.body || !container) return false;

                      const sourceBody = sourceDoc.body;
                      const sourceWindow = sourceDoc.defaultView || window;
                      const styleClone = document.createElement('div');
                      styleClone.style.cssText = 'width:1280px;height:720px;overflow:hidden;position:relative;';
                      styleClone.setAttribute('data-export-root', 'true');

                      const freezeMotionStyle = document.createElement('style');
                      freezeMotionStyle.textContent = '[data-export-root="true"] *, [data-export-root="true"] *::before, [data-export-root="true"] *::after { animation: none !important; transition: none !important; }';
                      styleClone.appendChild(freezeMotionStyle);

                      sourceDoc.querySelectorAll('style').forEach((styleNode) => {
                        styleClone.appendChild(styleNode.cloneNode(true));
                      });

                      sourceDoc.querySelectorAll('link[rel="stylesheet"]').forEach((linkNode) => {
                        styleClone.appendChild(linkNode.cloneNode(true));
                      });

                      const bodyStyle = sourceWindow.getComputedStyle(sourceBody);
                      const bodyBg = bodyStyle.backgroundColor;
                      const bodyBgImage = bodyStyle.backgroundImage;
                      if (bodyBg && bodyBg !== 'rgba(0, 0, 0, 0)') {
                        styleClone.style.backgroundColor = bodyBg;
                      }
                      if (bodyBgImage && bodyBgImage !== 'none') {
                        styleClone.style.backgroundImage = bodyBgImage;
                      }

                      const bodyContent = sourceBody.cloneNode(true);
                      copyComputedStylesForExport(sourceBody, bodyContent, sourceWindow);
                      neutralizeViewportFitScaleForExport(sourceBody, bodyContent, sourceWindow);
                      materializeBackgroundImagesForExport(sourceBody, bodyContent, sourceWindow);
                      replaceClonedCanvasesWithImages(sourceDoc, bodyContent);
                      bodyContent.style.margin = '0';
                      bodyContent.style.width = '1280px';
                      bodyContent.style.height = '720px';
                      bodyContent.style.overflow = 'hidden';
                      styleClone.appendChild(bodyContent);

                      container.innerHTML = '';
                      container.appendChild(styleClone);
                      return true;
                    } catch (_) {
                      return false;
                    }
                  }

                  async function waitForIframeReady(iframe) {
                    await new Promise((resolve, reject) => {
                      const timer = setTimeout(() => resolve(), 6000);
                      iframe.onload = () => {
                        clearTimeout(timer);
                        resolve();
                      };
                      iframe.onerror = (err) => {
                        clearTimeout(timer);
                        reject(err || new Error('iframe load failed'));
                      };
                    });
                    const doc = iframe.contentDocument || iframe.contentWindow.document;
                    if (!doc || !doc.body) {
                      throw new Error('Unable to access iframe content');
                    }
                    if (doc.fonts && doc.fonts.ready) {
                      try { await Promise.race([doc.fonts.ready, new Promise(r => setTimeout(r, 1800))]); } catch (_) {}
                    }
                    const images = Array.from(doc.images || []).filter((img) => !img.complete);
                    if (images.length) {
                      await Promise.race([
                        Promise.all(images.map((img) => new Promise((resolve) => {
                          img.onload = img.onerror = () => resolve();
                        }))),
                        new Promise((resolve) => setTimeout(resolve, 2400))
                      ]);
                    }
                    stabilizeLayoutForPpt(doc);
                    await new Promise((resolve) => requestAnimationFrame(() => requestAnimationFrame(resolve)));
                    return doc.body;
                  }

                  const slideElementStream = (async function* () {
                    for (const slide of payload.slides) {
                      const iframe = document.createElement('iframe');
                      iframe.style.cssText = 'width:1280px;height:720px;border:none;position:absolute;left:0;top:0;';
                      iframe.setAttribute('sandbox', 'allow-scripts allow-same-origin');
                      host.appendChild(iframe);
                      iframe.srcdoc = slide.html;
                      const root = await waitForIframeReady(iframe);
                      const sourceDoc = iframe.contentDocument || iframe.contentWindow.document;
                      const container = document.createElement('div');
                      container.style.cssText = 'width:1280px;height:720px;overflow:hidden;position:relative;background:white;';
                      host.appendChild(container);
                      let exportRoot = root;
                      try {
                        const built = buildExportContainerFromDocument(sourceDoc, container);
                        if (built) {
                          await waitForStylesheetsReadyInContainer(container, 2600);
                          await waitForImagesReadyInContainer(container, 2800);
                          convertModernColors(container);
                          exportRoot = container;
                        }
                        yield exportRoot;
                      } finally {
                        container.remove();
                        iframe.remove();
                      }
                    }
                  })();

                  const blob = await window.domToPptx.exportToPptx(slideElementStream, {
                    fileName: payload.fileName,
                    skipDownload: true,
                    svgAsVector: false,
                  });

                  const bytes = new Uint8Array(await blob.arrayBuffer());
                  let binary = '';
                  const chunkSize = 0x8000;
                  for (let i = 0; i < bytes.length; i += chunkSize) {
                    binary += String.fromCharCode.apply(null, bytes.subarray(i, i + chunkSize));
                  }
                  host.remove();
                  return btoa(binary);
                }
                """,
                payload,
            )
            browser.close()

        output_path.write_bytes(base64.b64decode(b64))
        return output_path


def build_dom_editable_deck_from_html(
    html_dir: Path,
    out_dir: Path,
    slide_meta: list[dict] | None = None,
    deck_name: str | None = None,
) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    preview_dir = out_dir / "previews"
    preview_dir.mkdir(exist_ok=True)
    html_files = sorted(html_dir.glob("*.html"))
    if not html_files:
        raise ValueError(f"HTML directory is empty: {html_dir}")

    meta_by_index = {}
    if slide_meta:
        meta_by_index = {item["index"]: item for item in slide_meta if "index" in item}

    slides = []
    html_paths_by_index: dict[int, Path] = {}
    for idx, html_path in enumerate(html_files, start=1):
        meta = meta_by_index.get(idx, {})
        title = meta.get("title", html_path.stem)
        page_role = meta.get("page_role", "summary")
        html = html_path.read_text(encoding="utf-8")
        source_png = preview_dir / f"html-source-{idx:02d}.png"
        source_png.write_bytes(render_html_screenshot(html_path))
        dom_preview_png = preview_dir / f"editable-preview-{idx:02d}.png"
        slides.append(
            {
                "index": idx,
                "title": title,
                "page_role": page_role,
                "html": html,
                "html_path": str(html_path),
                "preview_png_path": str(dom_preview_png),
                "html_preview_png_path": str(source_png),
            }
        )
        html_paths_by_index[idx] = html_path

    deck_stem = (deck_name or html_dir.parent.name)[:30]
    ppt_path = out_dir.parent / f"{deck_stem}_editable.pptx"
    DomPptxExporter().export_slides(slides, ppt_path)
    _replace_full_slide_background_picture(ppt_path, slides)
    structure_audit = _audit_exported_pptx(ppt_path, slides)
    if structure_audit["failures"]:
        ppt_path.unlink(missing_ok=True)
        raise RuntimeError("DOM editable export audit failed: " + " | ".join(structure_audit["failures"]))

    readback_audit = _render_powerpoint_readback_previews(ppt_path, preview_dir, slides)
    if not readback_audit["available"]:
        _ensure_dom_preview_fallback(slides)
        print(
            "    [audit] PowerPoint readback preview unavailable, using DOM source preview: "
            + str(readback_audit["reason"])
        )

    audit = _merge_audit_results(structure_audit, readback_audit)
    if audit["failures"]:
        ppt_path.unlink(missing_ok=True)
        raise RuntimeError("DOM editable export audit failed: " + " | ".join(audit["failures"]))

    slide_status = {}
    manifest_slides = []
    for slide_result in audit["slides"]:
        idx = int(slide_result["index"])
        slide_payload = slides[idx - 1]
        html_path = html_paths_by_index[idx]
        review_path = _write_dom_review(out_dir, slide_result, html_path)
        status_key = f"{idx:02d}"
        if slide_result["readback_available"]:
            validation_status = (
                "dom-verified-office"
                if slide_result["status"] == "pass"
                else "dom-verified-office-warn"
            )
        else:
            validation_status = "dom-verified" if slide_result["status"] == "pass" else "dom-verified-warn"
        review_status = "PASS" if slide_result["status"] == "pass" else "WARN"
        status_payload = {
            "title": slide_payload["title"],
            "page_role": slide_payload["page_role"],
            "validation_status": validation_status,
            "review_status": review_status,
            "review_rounds": 0,
            "export_ready": True,
            "html_path": slide_payload["html_path"],
            "preview_png_path": slide_result["preview_png_path"] or slide_payload["preview_png_path"],
            "html_preview_png_path": slide_payload["html_preview_png_path"],
            "preview_basis": slide_result["preview_basis"],
            "audit_basis": slide_result["audit_basis"],
            "shape_count": slide_result["shape_count"],
            "text_shape_count": slide_result["text_shape_count"],
            "picture_shape_count": slide_result["picture_shape_count"],
            "html_text_node_count": slide_result["html_text_node_count"],
            "audit_reasons": slide_result["reasons"],
            "structure_status": slide_result["structure_status"],
            "structure_reasons": slide_result["structure_reasons"],
            "readback_status": slide_result["readback_status"],
            "readback_reasons": slide_result["readback_reasons"],
            "readback_mean_pixel_delta": slide_result["readback_mean_pixel_delta"],
            "readback_dhash_distance": slide_result["readback_dhash_distance"],
            "readback_available": slide_result["readback_available"],
            "readback_error": slide_result["readback_error"],
            "preview_engine": slide_result["preview_engine"],
            "powerpoint_binary_path": slide_result["powerpoint_binary_path"],
            "powerpoint_used_ascii_workspace": slide_result["powerpoint_used_ascii_workspace"],
            "review_path": str(review_path),
        }
        slide_status[status_key] = status_payload
        manifest_slides.append(
            {
                "index": idx,
                "title": slide_payload["title"],
                "page_role": slide_payload["page_role"],
                "html_path": slide_payload["html_path"],
                "review_path": str(review_path),
                "preview_png_path": slide_result["preview_png_path"] or slide_payload["preview_png_path"],
                "html_preview_png_path": slide_payload["html_preview_png_path"],
                "export_ready": True,
                "validation_status": validation_status,
                "review_status": review_status,
                "preview_basis": slide_result["preview_basis"],
                "audit_basis": slide_result["audit_basis"],
                "shape_count": slide_result["shape_count"],
                "text_shape_count": slide_result["text_shape_count"],
                "picture_shape_count": slide_result["picture_shape_count"],
                "html_text_node_count": slide_result["html_text_node_count"],
                "audit_reasons": slide_result["reasons"],
                "structure_status": slide_result["structure_status"],
                "structure_reasons": slide_result["structure_reasons"],
                "readback_status": slide_result["readback_status"],
                "readback_reasons": slide_result["readback_reasons"],
                "readback_mean_pixel_delta": slide_result["readback_mean_pixel_delta"],
                "readback_dhash_distance": slide_result["readback_dhash_distance"],
                "readback_available": slide_result["readback_available"],
                "readback_error": slide_result["readback_error"],
                "preview_engine": slide_result["preview_engine"],
                "powerpoint_binary_path": slide_result["powerpoint_binary_path"],
                "powerpoint_used_ascii_workspace": slide_result["powerpoint_used_ascii_workspace"],
            }
        )
        if slide_result["status"] == "warn":
            print(
                f"    [audit] slide {idx:02d} has DOM export warning: "
                + "; ".join(slide_result["reasons"])
            )

    write_slide_status(out_dir, slide_status)
    (out_dir / "editable-export-manifest.json").write_text(
        json.dumps(
            {
                "version": 2,
                "generated_at": _utc_now_iso(),
                "pipeline": "dom-editable-export",
                "source_of_truth": "html",
                "html_dir": str(html_dir),
                "pptx_path": str(ppt_path),
                "slide_status_path": str(out_dir / "slide-status.json"),
                "preview_dir": str(preview_dir),
                "preview_basis": audit["preview_basis"],
                "audit_basis": audit["audit_basis"],
                "audit_summary": audit["summary"],
                "structure_audit_summary": structure_audit["summary"],
                "readback_audit_summary": audit["readback_summary"],
                "readback": {
                    "available": audit["readback_available"],
                    "engine": audit["preview_engine"],
                    "reason": audit["readback_reason"],
                    "binary_path": audit["powerpoint_binary_path"],
                    "used_ascii_workspace": audit["powerpoint_used_ascii_workspace"],
                },
                "slides": manifest_slides,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    return ppt_path
