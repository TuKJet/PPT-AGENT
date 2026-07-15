from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from pptx import Presentation

from html_pipeline.html_builder import build_pptx as build_html_pptx
from html_pipeline.html_builder import render_html_screenshot
from pptx_builder import build_pptx as build_svg_pptx
from vendor_presentation_core.export.dom_pptx_exporter import build_dom_editable_deck_from_html


class DeterministicRendererTests(unittest.TestCase):
    def test_html_renderer_does_not_rewrite_source(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            html_dir = root / "html"
            html_dir.mkdir()
            html_path = html_dir / "slide-01.html"
            source = (
                "<!doctype html><html><head><meta charset='utf-8'>"
                "<style>html,body{margin:0;width:1280px;height:720px}"
                "body{background:#123456;color:white}</style></head>"
                "<body><h1>Deterministic HTML</h1></body></html>"
            )
            html_path.write_text(source, encoding="utf-8")

            png = render_html_screenshot(html_path)
            output = root / "html.pptx"
            build_html_pptx(html_dir, output)

            self.assertTrue(png.startswith(b"\x89PNG\r\n\x1a\n"))
            self.assertEqual(html_path.read_text(encoding="utf-8"), source)
            self.assertEqual(len(Presentation(output).slides), 1)

    def test_svg_renderer_builds_one_slide_without_rewriting_source(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            svg_dir = root / "svg"
            svg_dir.mkdir()
            svg_path = svg_dir / "slide-01.svg"
            source = (
                "<svg xmlns='http://www.w3.org/2000/svg' viewBox='0 0 1280 720'>"
                "<rect width='1280' height='720' fill='#123456'/>"
                "<text x='80' y='120' fill='white'>Deterministic SVG</text>"
                "</svg>"
            )
            svg_path.write_text(source, encoding="utf-8")
            output = root / "svg.pptx"

            build_svg_pptx(svg_dir, output)

            self.assertEqual(svg_path.read_text(encoding="utf-8"), source)
            self.assertEqual(len(Presentation(output).slides), 1)

    def test_dom_editable_export_remains_available(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            html_dir = root / "html"
            html_dir.mkdir()
            (html_dir / "slide-01.html").write_text(
                "<!doctype html><html><head><meta charset='utf-8'>"
                "<style>html,body{margin:0;width:1280px;height:720px;overflow:hidden}"
                ".slide{position:relative;width:1280px;height:720px;background:#123456}"
                ".title{position:absolute;left:80px;top:80px;width:760px;height:80px;"
                "font:700 40px Arial;color:white}</style></head>"
                "<body><div class='slide'><div class='title'>Editable DOM</div></div></body></html>",
                encoding="utf-8",
            )

            output = build_dom_editable_deck_from_html(
                html_dir,
                root / "editable",
                deck_name="smoke",
            )

            presentation = Presentation(output)
            self.assertEqual(len(presentation.slides), 1)
            self.assertGreaterEqual(len(presentation.slides[0].shapes), 1)


if __name__ == "__main__":
    unittest.main()
