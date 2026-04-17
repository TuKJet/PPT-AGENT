from __future__ import annotations

import base64
import json
import math
import re
from datetime import datetime, timezone
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from PIL import Image, ImageChops, ImageFilter, ImageStat
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from html_pipeline.html_builder import render_html_screenshot, write_slide_status
from vendor_landppt.export.powerpoint_preview_renderer import (
    detect_powerpoint_render_support,
    export_powerpoint_slide_previews,
)


_STATUS_PRIORITY = {"pass": 0, "warn": 1, "fail": 2}
_RESAMPLE_LANCZOS = getattr(getattr(Image, "Resampling", Image), "LANCZOS")


def _utc_now_iso() -> str:
    return datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")


class _HtmlTextSignalParser(HTMLParser):
    _IGNORED_TAGS = {"script", "style", "title", "head", "meta", "link", "noscript"}

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._tag_stack: list[str] = []
        self.segments: list[str] = []

    def handle_starttag(self, tag: str, attrs) -> None:
        self._tag_stack.append(tag.lower())

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        for index in range(len(self._tag_stack) - 1, -1, -1):
            if self._tag_stack[index] == tag:
                del self._tag_stack[index:]
                break

    def handle_data(self, data: str) -> None:
        if any(tag in self._IGNORED_TAGS for tag in self._tag_stack):
            return
        text = re.sub(r"\s+", " ", data).strip()
        if not text:
            return
        if not re.search(r"[0-9A-Za-z\u4e00-\u9fff]", text):
            return
        self.segments.append(text[:80])


def _estimate_html_text_signal(html: str) -> dict[str, Any]:
    parser = _HtmlTextSignalParser()
    parser.feed(html)
    parser.close()
    return {
        "text_node_count": len(parser.segments),
        "samples": parser.segments[:5],
    }


def _iter_shapes(shapes) -> Any:
    for shape in shapes:
        yield shape
        if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)


def _shape_text(shape) -> str:
    raw_text = getattr(shape, "text", "") or ""
    return re.sub(r"\s+", " ", raw_text).strip()


def _merge_status(*statuses: str | None) -> str:
    merged = "pass"
    for status in statuses:
        if status is None:
            continue
        if _STATUS_PRIORITY.get(status, 0) > _STATUS_PRIORITY.get(merged, 0):
            merged = status
    return merged


def _compute_dhash(image: Image.Image, size: int = 8) -> int:
    grayscale = image.convert("L").resize((size + 1, size), _RESAMPLE_LANCZOS)
    pixels = grayscale.load()
    digest = 0
    for row in range(size):
        for col in range(size):
            digest <<= 1
            digest |= 1 if pixels[col, row] > pixels[col + 1, row] else 0
    return digest


def _hamming_distance(left: int, right: int) -> int:
    return (left ^ right).bit_count()


def _compare_preview_images(source_png: Path, preview_png: Path) -> dict[str, Any]:
    with Image.open(source_png) as source_image, Image.open(preview_png) as preview_image:
        source_rgb = source_image.convert("RGB")
        preview_rgb = preview_image.convert("RGB")
        if preview_rgb.size != source_rgb.size:
            preview_rgb = preview_rgb.resize(source_rgb.size, _RESAMPLE_LANCZOS)

        compare_size = (320, 180)
        source_compare = source_rgb.resize(compare_size, _RESAMPLE_LANCZOS).filter(
            ImageFilter.GaussianBlur(1.2)
        )
        preview_compare = preview_rgb.resize(compare_size, _RESAMPLE_LANCZOS).filter(
            ImageFilter.GaussianBlur(1.2)
        )
        diff = ImageChops.difference(source_compare, preview_compare)
        stat = ImageStat.Stat(diff)
        mean_pixel_delta = sum(stat.mean) / (len(stat.mean) * 255.0)
        dhash_distance = _hamming_distance(
            _compute_dhash(source_compare),
            _compute_dhash(preview_compare),
        )

    reasons: list[str] = []
    if mean_pixel_delta >= 0.14 or (dhash_distance >= 20 and mean_pixel_delta >= 0.05):
        reasons.append(
            "powerpoint readback diverges visually "
            f"(dhash={dhash_distance}, mean_delta={mean_pixel_delta:.4f})"
        )
        status = "fail"
    elif dhash_distance >= 15 or mean_pixel_delta >= 0.09:
        reasons.append(
            "powerpoint readback shows noticeable drift "
            f"(dhash={dhash_distance}, mean_delta={mean_pixel_delta:.4f})"
        )
        status = "warn"
    else:
        status = "pass"

    return {
        "status": status,
        "reasons": reasons,
        "mean_pixel_delta": round(mean_pixel_delta, 4),
        "dhash_distance": dhash_distance,
    }


def _ensure_dom_preview_fallback(slides: list[dict[str, Any]]) -> None:
    for slide in slides:
        source_path = Path(slide["html_preview_png_path"])
        preview_path = Path(slide["preview_png_path"])
        if preview_path.exists():
            continue
        preview_path.write_bytes(source_path.read_bytes())


def _render_powerpoint_readback_previews(
    ppt_path: Path,
    preview_dir: Path,
    slides: list[dict[str, Any]],
) -> dict[str, Any]:
    support = detect_powerpoint_render_support()
    if not support["available"]:
        return {
            "available": False,
            "engine": None,
            "reason": support["reason"],
            "binary_path": support["binary_path"],
            "used_ascii_workspace": False,
            "slides": [],
            "summary": None,
            "failures": [],
        }

    try:
        render_result = export_powerpoint_slide_previews(
            ppt_path=ppt_path,
            out_dir=preview_dir,
            slide_count=len(slides),
            prefix="editable-preview",
        )
    except Exception as exc:
        return {
            "available": False,
            "engine": support["engine"],
            "reason": str(exc),
            "binary_path": support["binary_path"],
            "used_ascii_workspace": False,
            "slides": [],
            "summary": None,
            "failures": [],
        }

    results: list[dict[str, Any]] = []
    failures: list[str] = []
    for index, slide in enumerate(slides, start=1):
        preview_path = Path(slide["preview_png_path"])
        if not preview_path.exists():
            result = {
                "index": index,
                "status": "fail",
                "reasons": ["PowerPoint readback preview missing"],
                "mean_pixel_delta": None,
                "dhash_distance": None,
                "preview_png_path": str(preview_path),
            }
        else:
            comparison = _compare_preview_images(
                Path(slide["html_preview_png_path"]),
                preview_path,
            )
            result = {
                "index": index,
                "status": comparison["status"],
                "reasons": comparison["reasons"],
                "mean_pixel_delta": comparison["mean_pixel_delta"],
                "dhash_distance": comparison["dhash_distance"],
                "preview_png_path": str(preview_path),
            }
        results.append(result)
        if result["status"] == "fail":
            failures.append(
                f"slide {index:02d}: " + "; ".join(result["reasons"] or ["PowerPoint readback audit failed"])
            )

    summary = {
        "pass": sum(1 for item in results if item["status"] == "pass"),
        "warn": sum(1 for item in results if item["status"] == "warn"),
        "fail": sum(1 for item in results if item["status"] == "fail"),
    }
    return {
        "available": True,
        "engine": render_result["engine"],
        "reason": None,
        "binary_path": render_result["binary_path"],
        "used_ascii_workspace": render_result["used_ascii_workspace"],
        "slides": results,
        "summary": summary,
        "failures": failures,
    }


def _merge_audit_results(
    structure_audit: dict[str, Any],
    readback_audit: dict[str, Any],
) -> dict[str, Any]:
    readback_by_index = {
        int(item["index"]): item for item in readback_audit.get("slides", []) if "index" in item
    }
    results: list[dict[str, Any]] = []
    failures = list(structure_audit.get("failures", []))
    readback_available = bool(readback_audit.get("available"))
    audit_basis = "pptx-structure+powerpoint-readback" if readback_available else "pptx-structure"
    preview_basis = "powerpoint-readback" if readback_available else "dom-source"

    for slide_result in structure_audit["slides"]:
        index = int(slide_result["index"])
        readback_result = readback_by_index.get(index)
        combined_reasons = list(slide_result["reasons"])
        combined_status = slide_result["status"]
        readback_status = "skipped"
        readback_reasons: list[str] = []
        readback_mean_pixel_delta = None
        readback_dhash_distance = None
        preview_png_path = None

        if readback_result:
            combined_status = _merge_status(combined_status, readback_result["status"])
            combined_reasons.extend(readback_result["reasons"])
            readback_status = readback_result["status"]
            readback_reasons = list(readback_result["reasons"])
            readback_mean_pixel_delta = readback_result.get("mean_pixel_delta")
            readback_dhash_distance = readback_result.get("dhash_distance")
            preview_png_path = readback_result.get("preview_png_path")

        merged = {
            **slide_result,
            "status": combined_status,
            "reasons": combined_reasons,
            "structure_status": slide_result["status"],
            "structure_reasons": list(slide_result["reasons"]),
            "readback_status": readback_status,
            "readback_reasons": readback_reasons,
            "readback_mean_pixel_delta": readback_mean_pixel_delta,
            "readback_dhash_distance": readback_dhash_distance,
            "preview_png_path": preview_png_path,
            "preview_basis": preview_basis,
            "audit_basis": audit_basis,
            "readback_available": readback_available,
            "readback_error": readback_audit.get("reason"),
            "preview_engine": readback_audit.get("engine"),
            "powerpoint_binary_path": readback_audit.get("binary_path"),
            "powerpoint_used_ascii_workspace": readback_audit.get("used_ascii_workspace", False),
        }
        results.append(merged)
        if merged["status"] == "fail" and not any(
            failure.startswith(f"slide {index:02d}") for failure in failures
        ):
            failures.append(
                f"slide {index:02d} {merged['title']}: " + "; ".join(merged["reasons"] or ["audit failed"])
            )

    summary = {
        "pass": sum(1 for item in results if item["status"] == "pass"),
        "warn": sum(1 for item in results if item["status"] == "warn"),
        "fail": sum(1 for item in results if item["status"] == "fail"),
    }
    return {
        "slides": results,
        "failures": failures,
        "summary": summary,
        "preview_basis": preview_basis,
        "audit_basis": audit_basis,
        "readback_available": readback_available,
        "readback_reason": readback_audit.get("reason"),
        "readback_summary": readback_audit.get("summary"),
        "preview_engine": readback_audit.get("engine"),
        "powerpoint_binary_path": readback_audit.get("binary_path"),
        "powerpoint_used_ascii_workspace": readback_audit.get("used_ascii_workspace", False),
    }


def _audit_exported_pptx(ppt_path: Path, slides: list[dict[str, Any]]) -> dict[str, Any]:
    deck = Presentation(str(ppt_path))
    results: list[dict[str, Any]] = []
    failures: list[str] = []

    if len(deck.slides) != len(slides):
        failures.append(
            f"slide count mismatch: pptx={len(deck.slides)} html={len(slides)}"
        )

    for idx, slide_payload in enumerate(slides, start=1):
        ppt_slide = deck.slides[idx - 1] if idx - 1 < len(deck.slides) else None
        text_signal = _estimate_html_text_signal(slide_payload["html"])
        html_text_nodes = text_signal["text_node_count"]
        shape_count = 0
        text_shape_count = 0
        picture_shape_count = 0
        reasons: list[str] = []

        if ppt_slide is None:
            reasons.append("missing slide in exported pptx")
            status = "fail"
        else:
            all_shapes = list(_iter_shapes(ppt_slide.shapes))
            shape_count = len(all_shapes)
            for shape in all_shapes:
                if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.PICTURE:
                    picture_shape_count += 1
                if getattr(shape, "has_text_frame", False) and _shape_text(shape):
                    text_shape_count += 1

            min_text_shapes = 0
            min_total_shapes = 0
            if html_text_nodes >= 10:
                min_text_shapes = max(3, math.ceil(html_text_nodes * 0.12))
            elif html_text_nodes >= 4:
                min_text_shapes = 1
            if html_text_nodes >= 12:
                min_total_shapes = max(4, math.ceil(html_text_nodes * 0.18))
            elif html_text_nodes >= 4:
                min_total_shapes = 2

            screenshot_like = (
                html_text_nodes >= 3
                and shape_count <= 2
                and picture_shape_count >= 1
                and text_shape_count <= 1
            )
            if screenshot_like:
                reasons.append("pptx structure looks rasterized")
            if min_text_shapes and text_shape_count < min_text_shapes:
                reasons.append(
                    f"text shapes too low ({text_shape_count} < {min_text_shapes})"
                )
            if min_total_shapes and shape_count < min_total_shapes:
                reasons.append(
                    f"shape count too low ({shape_count} < {min_total_shapes})"
                )

            if screenshot_like or (len(reasons) >= 2 and html_text_nodes >= 10):
                status = "fail"
            elif reasons:
                status = "warn"
            else:
                status = "pass"

        result = {
            "index": idx,
            "title": slide_payload.get("title", f"slide-{idx:02d}"),
            "status": status,
            "shape_count": shape_count,
            "text_shape_count": text_shape_count,
            "picture_shape_count": picture_shape_count,
            "html_text_node_count": html_text_nodes,
            "html_text_samples": text_signal["samples"],
            "reasons": reasons,
        }
        results.append(result)
        if status == "fail":
            failures.append(
                f"slide {idx:02d} {result['title']}: " + "; ".join(reasons or ["audit failed"])
            )

    summary = {
        "pass": sum(1 for item in results if item["status"] == "pass"),
        "warn": sum(1 for item in results if item["status"] == "warn"),
        "fail": sum(1 for item in results if item["status"] == "fail"),
    }
    return {"slides": results, "failures": failures, "summary": summary}


def _write_dom_review(out_dir: Path, slide_result: dict[str, Any], html_path: Path) -> Path:
    review_path = out_dir / f"review-{slide_result['index']:02d}.md"
    status_label = {
        "pass": "PASS",
        "warn": "WARN",
        "fail": "FAIL",
    }[slide_result["status"]]
    lines = [
        f"# {html_path.name}",
        "",
        f"RESULT: {status_label}",
        "",
        "## Reasons",
    ]
    lines.extend([f"- {item}" for item in slide_result["reasons"]] or ["- DOM export passed all enabled audits"])
    lines.extend(
        [
            "",
            "## Audit Layers",
            f"- structure_status: {slide_result['structure_status']}",
            f"- readback_status: {slide_result['readback_status']}",
            f"- preview_basis: {slide_result['preview_basis']}",
            f"- audit_basis: {slide_result['audit_basis']}",
            "",
            "## Metrics",
            f"- shape_count: {slide_result['shape_count']}",
            f"- text_shape_count: {slide_result['text_shape_count']}",
            f"- picture_shape_count: {slide_result['picture_shape_count']}",
            f"- html_text_node_count: {slide_result['html_text_node_count']}",
            f"- readback_mean_pixel_delta: {slide_result.get('readback_mean_pixel_delta')}",
            f"- readback_dhash_distance: {slide_result.get('readback_dhash_distance')}",
            "",
            "## HTML Text Samples",
        ]
    )
    lines.extend([f"- {item}" for item in slide_result["html_text_samples"]] or ["- no sampled text nodes"])
    if slide_result.get("readback_reasons"):
        lines.extend(["", "## Readback Notes"])
        lines.extend([f"- {item}" for item in slide_result["readback_reasons"]])
    elif slide_result.get("readback_error"):
        lines.extend(["", "## Readback Notes", f"- {slide_result['readback_error']}"])
    review_path.write_text("\n".join(lines), encoding="utf-8")
    return review_path


class DomPptxExporter:
    """Run LandPPT's browser-side DOM exporter inside Playwright."""

    def __init__(self, bundle_path: Path | None = None):
        self.bundle_path = bundle_path or (Path(__file__).resolve().parent / "dom-to-pptx.bundle.js")

    def export_slides(self, slides: list[dict[str, Any]], output_path: Path) -> Path:
        if not slides:
            raise ValueError("No slides provided for DOM export")

        shell_html = """<!DOCTYPE html><html><head><meta charset='utf-8'><title>dom export</title></head><body></body></html>"""
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch()
            page = browser.new_page(viewport={"width": 1400, "height": 900})
            page.set_content(shell_html, wait_until="domcontentloaded")
            page.add_script_tag(path=str(self.bundle_path))
            payload = {"slides": slides, "fileName": output_path.name}
            b64 = page.evaluate(
                """
                async (payload) => {
                  const host = document.createElement('div');
                  host.style.cssText = 'position:fixed;left:-9999px;top:-9999px;opacity:0;pointer-events:none;';
                  document.body.appendChild(host);

                  function stabilizeLayoutForPpt(doc) {
                    const win = doc.defaultView;
                    if (!win) return;

                    const clampPx = (value, fallback = 0) => {
                      const parsed = Number.parseFloat(String(value ?? ''));
                      return Number.isFinite(parsed) ? parsed : fallback;
                    };

                    const extractLastLinearGradient = (value) => {
                      const input = String(value || '');
                      const lower = input.toLowerCase();
                      const token = 'linear-gradient(';
                      const start = lower.lastIndexOf(token);
                      if (start < 0) return '';

                      let depth = 1;
                      let quote = '';
                      for (let index = start + token.length; index < input.length; index++) {
                        const ch = input[index];
                        if (quote) {
                          if (ch === quote && input[index - 1] !== '\\\\') {
                            quote = '';
                          }
                          continue;
                        }
                        if (ch === '"' || ch === "'") {
                          quote = ch;
                          continue;
                        }
                        if (ch === '(') {
                          depth += 1;
                          continue;
                        }
                        if (ch === ')') {
                          depth -= 1;
                          if (depth === 0) {
                            return input.slice(start, index + 1);
                          }
                        }
                      }
                      return '';
                    };

                    const extractRadialGlowSpecs = (value) => {
                      const input = String(value || '');
                      const specs = [];
                      const pattern = /radial-gradient\\(\\s*circle at\\s*([^,]+),\\s*(rgba?\\([^\\)]+\\)|#[0-9a-fA-F]+)\\s*,\\s*transparent\\s+([0-9.]+)%\\s*\\)/gi;
                      let match = null;
                      while ((match = pattern.exec(input))) {
                        specs.push({
                          position: String(match[1] || '').trim().toLowerCase(),
                          color: String(match[2] || '').trim(),
                          stop: clampPx(match[3], 30),
                        });
                      }
                      return specs.slice(0, 3);
                    };

                    const ensureRootBackdrop = () => {
                      const htmlEl = doc.documentElement;
                      const body = doc.body;
                      if (!(htmlEl instanceof win.HTMLElement) || !(body instanceof win.HTMLElement)) {
                        return;
                      }

                      htmlEl.style.width = '1280px';
                      htmlEl.style.height = '720px';
                      htmlEl.style.margin = '0';
                      htmlEl.style.overflow = 'hidden';
                      body.style.width = '1280px';
                      body.style.height = '720px';
                      body.style.minHeight = '720px';
                      body.style.margin = '0';
                      body.style.overflow = 'hidden';
                      if (!body.style.position || body.style.position === 'static') {
                        body.style.position = 'relative';
                      }

                      const htmlStyle = win.getComputedStyle(htmlEl);
                      const bodyStyle = win.getComputedStyle(body);
                      const transparentColor = 'rgba(0, 0, 0, 0)';
                      const bgSourceImage =
                        bodyStyle.backgroundImage && bodyStyle.backgroundImage !== 'none'
                          ? bodyStyle.backgroundImage
                          : htmlStyle.backgroundImage;
                      const bgSourceColor =
                        bodyStyle.backgroundColor && bodyStyle.backgroundColor !== transparentColor
                          ? bodyStyle.backgroundColor
                          : htmlStyle.backgroundColor;
                      const linearGradient = extractLastLinearGradient(bgSourceImage);
                      const radialSpecs = extractRadialGlowSpecs(bgSourceImage);
                      const hasRenderableBackdrop =
                        !!linearGradient ||
                        (!!bgSourceColor && bgSourceColor !== transparentColor) ||
                        radialSpecs.length > 0;
                      if (!hasRenderableBackdrop) {
                        return;
                      }

                      let backdrop = Array.from(body.children).find(
                        (child) =>
                          child instanceof win.HTMLElement &&
                          child.getAttribute('data-ppt-root-backdrop') === 'true'
                      );
                      if (!(backdrop instanceof win.HTMLElement)) {
                        backdrop = doc.createElement('div');
                        backdrop.setAttribute('data-ppt-root-backdrop', 'true');
                        body.insertBefore(backdrop, body.firstChild);
                      }

                      backdrop.style.position = 'absolute';
                      backdrop.style.left = '0';
                      backdrop.style.top = '0';
                      backdrop.style.width = '1280px';
                      backdrop.style.height = '720px';
                      backdrop.style.pointerEvents = 'none';
                      backdrop.style.overflow = 'hidden';
                      backdrop.style.zIndex = '0';
                      backdrop.style.borderRadius = '0';
                      backdrop.style.backgroundColor =
                        bgSourceColor && bgSourceColor !== transparentColor ? bgSourceColor : 'transparent';
                      backdrop.style.backgroundImage = linearGradient || 'none';
                      backdrop.style.backgroundRepeat = 'no-repeat';
                      backdrop.style.backgroundSize = '100% 100%';
                      backdrop.style.backgroundPosition = 'center';

                      Array.from(backdrop.children).forEach((child) => {
                        if (child instanceof win.HTMLElement && child.getAttribute('data-ppt-root-glow') === 'true') {
                          child.remove();
                        }
                      });

                      radialSpecs.forEach((spec) => {
                        const glow = doc.createElement('div');
                        const baseSize = Math.max(1280, 720);
                        const size = Math.max(
                          260,
                          Math.round(baseSize * Math.max(0.28, Math.min((spec.stop || 30) / 100 * 1.45, 0.64)))
                        );
                        const blur = Math.max(60, Math.round(size * 0.2));
                        glow.setAttribute('data-ppt-root-glow', 'true');
                        glow.style.position = 'absolute';
                        glow.style.width = `${size}px`;
                        glow.style.height = `${size}px`;
                        glow.style.borderRadius = '999px';
                        glow.style.backgroundColor = spec.color;
                        glow.style.filter = `blur(${blur}px)`;
                        glow.style.opacity = '1';
                        glow.style.pointerEvents = 'none';
                        glow.style.zIndex = '0';

                        const pos = spec.position;
                        if (pos.includes('100% 100%') || pos.includes('bottom right') || pos.includes('right bottom')) {
                          glow.style.right = `${Math.round(-size * 0.16)}px`;
                          glow.style.bottom = `${Math.round(-size * 0.16)}px`;
                        } else if (
                          pos.includes('100% 0%') ||
                          pos.includes('top right') ||
                          pos.includes('right top')
                        ) {
                          glow.style.right = `${Math.round(-size * 0.16)}px`;
                          glow.style.top = `${Math.round(-size * 0.16)}px`;
                        } else if (
                          pos.includes('0% 100%') ||
                          pos.includes('bottom left') ||
                          pos.includes('left bottom')
                        ) {
                          glow.style.left = `${Math.round(-size * 0.16)}px`;
                          glow.style.bottom = `${Math.round(-size * 0.16)}px`;
                        } else {
                          glow.style.left = `${Math.round(-size * 0.16)}px`;
                          glow.style.top = `${Math.round(-size * 0.16)}px`;
                        }

                        backdrop.appendChild(glow);
                      });

                      Array.from(body.children).forEach((child) => {
                        if (!(child instanceof win.HTMLElement) || child === backdrop) {
                          return;
                        }
                        const childStyle = win.getComputedStyle(child);
                        if (childStyle.position === 'static') {
                          child.style.position = 'relative';
                        }
                        const childZIndex = Number.parseInt(childStyle.zIndex, 10);
                        if (!Number.isFinite(childZIndex) || childZIndex < 1) {
                          child.style.zIndex = '1';
                        }
                      });

                      body.style.backgroundImage = 'none';
                      body.style.backgroundColor = 'transparent';
                      htmlEl.style.backgroundImage = 'none';
                      htmlEl.style.backgroundColor = 'transparent';
                    };

                    const stabilizeHeader = (header) => {
                      const headerStyle = win.getComputedStyle(header);
                      if (headerStyle.display !== 'flex') return;

                      const children = Array.from(header.children).filter(
                        (node) => node instanceof win.HTMLElement
                      );
                      if (children.length < 2) return;

                      const lead = children[0];
                      const tail = children[children.length - 1];
                      const headerRect = header.getBoundingClientRect();
                      const tailRect = tail.getBoundingClientRect();
                      if (!headerRect.width) return;

                      // Flex headers with `space-between` often keep the text column at max-content width.
                      // That is stable in Chromium, but PowerPoint may reflow with different font metrics.
                      const reserve = Math.max(110, tailRect.width + 36);
                      const available = Math.max(360, Math.floor(headerRect.width - reserve));

                      lead.style.flex = `0 0 ${available}px`;
                      lead.style.width = `${available}px`;
                      lead.style.maxWidth = `${available}px`;
                      lead.style.minWidth = '0';
                      if (!lead.style.paddingRight) {
                        lead.style.paddingRight = '24px';
                      }

                      tail.style.flex = '0 0 auto';
                      tail.style.marginLeft = '24px';
                      tail.style.alignSelf = 'flex-start';
                      tail.style.maxWidth = `${Math.max(132, Math.floor(headerRect.width * 0.34))}px`;
                      tail.style.minWidth = tail.style.minWidth || '92px';

                      if (tail.classList.contains('corner-cluster') || tail.querySelector('.page-tag, .meta-chip, [class*="chip"], [class*="tag"]')) {
                        tail.style.display = 'grid';
                        tail.style.justifyItems = 'end';
                        tail.style.rowGap = tail.style.rowGap || '10px';
                      }

                      const textNodes = Array.from(
                        lead.querySelectorAll('h1, h2, h3, p, .title, .subtitle')
                      ).filter((node) => node instanceof win.HTMLElement);

                      const tailNodes = Array.from(
                        tail.querySelectorAll('.page-tag, .meta-chip, [class*="chip"], [class*="tag"], [class*="badge"]')
                      ).filter((node) => node instanceof win.HTMLElement);

                      textNodes.forEach((node) => {
                        node.style.width = '100%';
                        node.style.maxWidth = '100%';
                        node.style.minWidth = '0';

                        const computed = win.getComputedStyle(node);
                        const fontSize = clampPx(computed.fontSize);
                        const lineHeight = clampPx(computed.lineHeight, fontSize * 1.08);

                        if (node.matches('h1, .title')) {
                          node.style.whiteSpace = 'normal';
                          node.style.wordBreak = 'break-word';
                          node.style.lineHeight = `${Math.max(lineHeight, fontSize * 1.1)}px`;

                          const letterSpacing = clampPx(computed.letterSpacing);
                          if (letterSpacing < -0.6) {
                            node.style.letterSpacing = '-0.4px';
                          }
                        } else if (node.matches('p, .subtitle')) {
                          node.style.lineHeight = `${Math.max(lineHeight, fontSize * 1.2)}px`;
                        }
                      });

                      tailNodes.forEach((node) => {
                        node.style.whiteSpace = 'nowrap';
                        node.style.minWidth = node.classList.contains('page-tag') ? '68px' : (node.style.minWidth || '0');
                        node.style.paddingLeft = node.style.paddingLeft || '10px';
                        node.style.paddingRight = node.style.paddingRight || '10px';
                        if (node.classList.contains('page-tag')) {
                          node.style.display = 'inline-flex';
                          node.style.alignItems = 'center';
                          node.style.justifyContent = 'center';
                        }
                      });
                    };

                    Array.from(doc.querySelectorAll('header')).forEach((header) => {
                      if (header instanceof win.HTMLElement) {
                        stabilizeHeader(header);
                      }
                    });

                    ensureRootBackdrop();
                  }

                  async function waitForIframeReady(iframe) {
                    await new Promise((resolve, reject) => {
                      const timer = setTimeout(() => resolve(), 6000);
                      iframe.onload = () => {
                        clearTimeout(timer);
                        resolve();
                      };
                      iframe.onerror = (err) => {
                        clearTimeout(timer);
                        reject(err || new Error('iframe load failed'));
                      };
                    });
                    const doc = iframe.contentDocument || iframe.contentWindow.document;
                    if (!doc || !doc.body) {
                      throw new Error('Unable to access iframe content');
                    }
                    if (doc.fonts && doc.fonts.ready) {
                      try { await Promise.race([doc.fonts.ready, new Promise(r => setTimeout(r, 1800))]); } catch (_) {}
                    }
                    const images = Array.from(doc.images || []).filter((img) => !img.complete);
                    if (images.length) {
                      await Promise.race([
                        Promise.all(images.map((img) => new Promise((resolve) => {
                          img.onload = img.onerror = () => resolve();
                        }))),
                        new Promise((resolve) => setTimeout(resolve, 2400))
                      ]);
                    }
                    stabilizeLayoutForPpt(doc);
                    await new Promise((resolve) => requestAnimationFrame(() => requestAnimationFrame(resolve)));
                    return doc.body;
                  }

                  const slideElementStream = (async function* () {
                    for (const slide of payload.slides) {
                      const iframe = document.createElement('iframe');
                      iframe.style.cssText = 'width:1280px;height:720px;border:none;position:absolute;left:0;top:0;';
                      iframe.setAttribute('sandbox', 'allow-scripts allow-same-origin');
                      host.appendChild(iframe);
                      iframe.srcdoc = slide.html;
                      const root = await waitForIframeReady(iframe);
                      try {
                        yield root;
                      } finally {
                        iframe.remove();
                      }
                    }
                  })();

                  const blob = await window.domToPptx.exportToPptx(slideElementStream, {
                    fileName: payload.fileName,
                    skipDownload: true,
                    svgAsVector: false,
                  });

                  const bytes = new Uint8Array(await blob.arrayBuffer());
                  let binary = '';
                  const chunkSize = 0x8000;
                  for (let i = 0; i < bytes.length; i += chunkSize) {
                    binary += String.fromCharCode.apply(null, bytes.subarray(i, i + chunkSize));
                  }
                  host.remove();
                  return btoa(binary);
                }
                """,
                payload,
            )
            browser.close()

        output_path.write_bytes(base64.b64decode(b64))
        return output_path


def build_dom_editable_deck_from_html(
    html_dir: Path,
    out_dir: Path,
    slide_meta: list[dict] | None = None,
    deck_name: str | None = None,
) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    preview_dir = out_dir / "previews"
    preview_dir.mkdir(exist_ok=True)
    html_files = sorted(html_dir.glob("*.html"))
    if not html_files:
        raise ValueError(f"HTML directory is empty: {html_dir}")

    meta_by_index = {}
    if slide_meta:
        meta_by_index = {item["index"]: item for item in slide_meta if "index" in item}

    slides = []
    html_paths_by_index: dict[int, Path] = {}
    for idx, html_path in enumerate(html_files, start=1):
        meta = meta_by_index.get(idx, {})
        title = meta.get("title", html_path.stem)
        page_role = meta.get("page_role", "summary")
        html = html_path.read_text(encoding="utf-8")
        source_png = preview_dir / f"html-source-{idx:02d}.png"
        source_png.write_bytes(render_html_screenshot(html_path))
        dom_preview_png = preview_dir / f"editable-preview-{idx:02d}.png"
        slides.append(
            {
                "index": idx,
                "title": title,
                "page_role": page_role,
                "html": html,
                "html_path": str(html_path),
                "preview_png_path": str(dom_preview_png),
                "html_preview_png_path": str(source_png),
            }
        )
        html_paths_by_index[idx] = html_path

    deck_stem = (deck_name or html_dir.parent.name)[:30]
    ppt_path = out_dir.parent / f"{deck_stem}_editable.pptx"
    DomPptxExporter().export_slides(slides, ppt_path)
    structure_audit = _audit_exported_pptx(ppt_path, slides)
    if structure_audit["failures"]:
        ppt_path.unlink(missing_ok=True)
        raise RuntimeError("DOM editable export audit failed: " + " | ".join(structure_audit["failures"]))

    readback_audit = _render_powerpoint_readback_previews(ppt_path, preview_dir, slides)
    if not readback_audit["available"]:
        _ensure_dom_preview_fallback(slides)
        print(
            "    [audit] PowerPoint readback preview unavailable, using DOM source preview: "
            + str(readback_audit["reason"])
        )

    audit = _merge_audit_results(structure_audit, readback_audit)
    if audit["failures"]:
        ppt_path.unlink(missing_ok=True)
        raise RuntimeError("DOM editable export audit failed: " + " | ".join(audit["failures"]))

    slide_status = {}
    manifest_slides = []
    for slide_result in audit["slides"]:
        idx = int(slide_result["index"])
        slide_payload = slides[idx - 1]
        html_path = html_paths_by_index[idx]
        review_path = _write_dom_review(out_dir, slide_result, html_path)
        status_key = f"{idx:02d}"
        if slide_result["readback_available"]:
            validation_status = (
                "dom-verified-office"
                if slide_result["status"] == "pass"
                else "dom-verified-office-warn"
            )
        else:
            validation_status = "dom-verified" if slide_result["status"] == "pass" else "dom-verified-warn"
        review_status = "PASS" if slide_result["status"] == "pass" else "WARN"
        status_payload = {
            "title": slide_payload["title"],
            "page_role": slide_payload["page_role"],
            "validation_status": validation_status,
            "review_status": review_status,
            "review_rounds": 0,
            "export_ready": True,
            "html_path": slide_payload["html_path"],
            "preview_png_path": slide_result["preview_png_path"] or slide_payload["preview_png_path"],
            "html_preview_png_path": slide_payload["html_preview_png_path"],
            "preview_basis": slide_result["preview_basis"],
            "audit_basis": slide_result["audit_basis"],
            "shape_count": slide_result["shape_count"],
            "text_shape_count": slide_result["text_shape_count"],
            "picture_shape_count": slide_result["picture_shape_count"],
            "html_text_node_count": slide_result["html_text_node_count"],
            "audit_reasons": slide_result["reasons"],
            "structure_status": slide_result["structure_status"],
            "structure_reasons": slide_result["structure_reasons"],
            "readback_status": slide_result["readback_status"],
            "readback_reasons": slide_result["readback_reasons"],
            "readback_mean_pixel_delta": slide_result["readback_mean_pixel_delta"],
            "readback_dhash_distance": slide_result["readback_dhash_distance"],
            "readback_available": slide_result["readback_available"],
            "readback_error": slide_result["readback_error"],
            "preview_engine": slide_result["preview_engine"],
            "powerpoint_binary_path": slide_result["powerpoint_binary_path"],
            "powerpoint_used_ascii_workspace": slide_result["powerpoint_used_ascii_workspace"],
            "review_path": str(review_path),
        }
        slide_status[status_key] = status_payload
        manifest_slides.append(
            {
                "index": idx,
                "title": slide_payload["title"],
                "page_role": slide_payload["page_role"],
                "html_path": slide_payload["html_path"],
                "review_path": str(review_path),
                "preview_png_path": slide_result["preview_png_path"] or slide_payload["preview_png_path"],
                "html_preview_png_path": slide_payload["html_preview_png_path"],
                "export_ready": True,
                "validation_status": validation_status,
                "review_status": review_status,
                "preview_basis": slide_result["preview_basis"],
                "audit_basis": slide_result["audit_basis"],
                "shape_count": slide_result["shape_count"],
                "text_shape_count": slide_result["text_shape_count"],
                "picture_shape_count": slide_result["picture_shape_count"],
                "html_text_node_count": slide_result["html_text_node_count"],
                "audit_reasons": slide_result["reasons"],
                "structure_status": slide_result["structure_status"],
                "structure_reasons": slide_result["structure_reasons"],
                "readback_status": slide_result["readback_status"],
                "readback_reasons": slide_result["readback_reasons"],
                "readback_mean_pixel_delta": slide_result["readback_mean_pixel_delta"],
                "readback_dhash_distance": slide_result["readback_dhash_distance"],
                "readback_available": slide_result["readback_available"],
                "readback_error": slide_result["readback_error"],
                "preview_engine": slide_result["preview_engine"],
                "powerpoint_binary_path": slide_result["powerpoint_binary_path"],
                "powerpoint_used_ascii_workspace": slide_result["powerpoint_used_ascii_workspace"],
            }
        )
        if slide_result["status"] == "warn":
            print(
                f"    [audit] slide {idx:02d} has DOM export warning: "
                + "; ".join(slide_result["reasons"])
            )

    write_slide_status(out_dir, slide_status)
    (out_dir / "editable-export-manifest.json").write_text(
        json.dumps(
            {
                "version": 2,
                "generated_at": _utc_now_iso(),
                "pipeline": "landppt-dom-export",
                "source_of_truth": "html",
                "html_dir": str(html_dir),
                "pptx_path": str(ppt_path),
                "slide_status_path": str(out_dir / "slide-status.json"),
                "preview_dir": str(preview_dir),
                "preview_basis": audit["preview_basis"],
                "audit_basis": audit["audit_basis"],
                "audit_summary": audit["summary"],
                "structure_audit_summary": structure_audit["summary"],
                "readback_audit_summary": audit["readback_summary"],
                "readback": {
                    "available": audit["readback_available"],
                    "engine": audit["preview_engine"],
                    "reason": audit["readback_reason"],
                    "binary_path": audit["powerpoint_binary_path"],
                    "used_ascii_workspace": audit["powerpoint_used_ascii_workspace"],
                },
                "slides": manifest_slides,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    return ppt_path
