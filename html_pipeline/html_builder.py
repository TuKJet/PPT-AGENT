from __future__ import annotations

import io
import json
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from playwright_runtime import launch_global_chromium, sync_playwright


SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)
VIEWPORT = {"width": 1280, "height": 720}


def _wait_for_page_assets(page) -> None:
    page.evaluate(
        """
        async () => {
          if (document.fonts && document.fonts.ready) {
            await document.fonts.ready;
          }
          const pendingImages = Array.from(document.images)
            .filter((image) => !image.complete)
            .map((image) => new Promise((resolve) => {
              image.addEventListener('load', resolve, { once: true });
              image.addEventListener('error', resolve, { once: true });
            }));
          await Promise.all(pendingImages);
          await new Promise((resolve) => requestAnimationFrame(() => requestAnimationFrame(resolve)));
        }
        """
    )


def render_html_screenshot(html_path: Path) -> bytes:
    """Render one authored HTML page without repairing or rewriting its source."""
    source_before = html_path.read_bytes()
    with sync_playwright() as playwright:
        browser = launch_global_chromium(playwright)
        try:
            page = browser.new_page(viewport=VIEWPORT)
            page.goto(html_path.resolve().as_uri(), wait_until="networkidle")
            _wait_for_page_assets(page)
            png = page.screenshot(
                type="png",
                clip={"x": 0, "y": 0, **VIEWPORT},
            )
        finally:
            browser.close()

    if html_path.read_bytes() != source_before:
        raise RuntimeError(f"HTML renderer unexpectedly changed its source: {html_path}")

    return png


def save_html_screenshot(html_path: Path, image_path: Path) -> Path:
    image_path.parent.mkdir(parents=True, exist_ok=True)
    image_path.write_bytes(render_html_screenshot(html_path))
    return image_path


def write_slide_status(out_dir: Path, slide_status: dict) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    status_path = out_dir / "slide-status.json"
    status_path.write_text(
        json.dumps({"slides": slide_status}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return status_path


def html_to_png_bytes(html_path: Path) -> bytes:
    return render_html_screenshot(html_path)


def build_pptx(html_dir: Path, output_path: Path) -> Path:
    """Build a widescreen image PPTX from Codex-authored HTML pages."""
    html_files = sorted(html_dir.glob("*.html"))
    if not html_files:
        raise ValueError(f"HTML directory is empty: {html_dir}")

    presentation = Presentation()
    presentation.slide_width = SLIDE_W
    presentation.slide_height = SLIDE_H
    blank_layout = presentation.slide_layouts[6]

    for html_path in html_files:
        png_bytes = render_html_screenshot(html_path)
        slide = presentation.slides.add_slide(blank_layout)
        slide.shapes.add_picture(
            io.BytesIO(png_bytes),
            left=0,
            top=0,
            width=SLIDE_W,
            height=SLIDE_H,
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(output_path))
    return output_path
