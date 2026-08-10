from __future__ import annotations

import argparse
import base64
import binascii
import hashlib
import io
import json
import os
import shutil
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from PIL import Image, ImageChops, ImageFilter, ImageStat

SKILL_ROOT = Path(__file__).resolve().parents[1]
BUNDLED_RUNTIME_ROOT = SKILL_ROOT / "runtime"
GLOBAL_SKILL_MODE = (BUNDLED_RUNTIME_ROOT / "filename_utils.py").is_file()
try:
    LOCAL_REPO_ROOT = Path(__file__).resolve().parents[4]
except IndexError:
    LOCAL_REPO_ROOT = SKILL_ROOT
RUNTIME_ROOT = BUNDLED_RUNTIME_ROOT if GLOBAL_SKILL_MODE else LOCAL_REPO_ROOT
WORKSPACE_ROOT = Path(
    os.getenv("PPT_AGENT_WORKSPACE", str(Path.cwd() if GLOBAL_SKILL_MODE else LOCAL_REPO_ROOT))
).resolve()

if str(RUNTIME_ROOT) not in sys.path:
    sys.path.insert(0, str(RUNTIME_ROOT))

from filename_utils import safe_filename_part, slide_filename

STATE_FILE = "workflow-state.json"
ARTIFACT_FILES = {
    "outline": "outline.json",
    "contents": "contents.json",
    "slide_plans": "slide-plans.json",
}
PREVIEW_FILES = {
    "outline": "outline-preview.md",
    "contents": "contents-preview.md",
    "slide_plans": "slide-plans-preview.md",
}
COMMON_ARTIFACT_KEYS = set(ARTIFACT_FILES)
RENDER_REVIEW_NOT_APPLICABLE = {
    "mode": "off",
    "status": "not_applicable",
}
SLIDE_PLAN_VERSION = 2
REQUIRED_DECK_STRATEGY_FIELDS = (
    "primary_audience",
    "decision_context",
    "first_questions",
    "evidence_order",
    "presentation_posture",
)
REQUIRED_DESIGN_SYSTEM_FIELDS = (
    "theme_name",
    "audience_fit",
    "palette",
    "typography",
    "component_rules",
    "chart_treatment",
    "illustration_policy",
    "design_genes",
)
REQUIRED_PALETTE_TOKENS = (
    "background",
    "surface",
    "primary",
    "accent",
    "text_primary",
    "text_muted",
)
REQUIRED_SLIDE_PLAN_FIELDS = (
    "core_message",
    "layout_structure",
    "visual_hierarchy",
    "required_elements",
    "palette_tokens",
    "style_controls",
    "audience_controls",
    "renderer_neutral_constraints",
)
PREVIEW_LABELS = {
    "primary_audience": "目标受众",
    "secondary_audience": "次要受众",
    "decision_context": "决策目标",
    "first_questions": "首要问题",
    "evidence_order": "证据顺序",
    "presentation_posture": "表达姿态",
    "theme_name": "主题名称",
    "audience_fit": "受众适配",
    "typography": "字体层级",
    "component_rules": "组件与间距",
    "chart_treatment": "图表与图解",
    "illustration_policy": "图像与插画",
    "design_genes": "设计基因",
    "title": "标题",
    "section_heading": "分区标题",
    "body": "正文",
    "metric": "核心数字",
    "cards": "卡片",
    "spacing": "间距",
    "lines": "线条",
    "footer": "页脚",
    "style": "表现方式",
    "labels": "标注",
    "risk_encoding": "风险编码",
    "density": "信息密度",
    "visual_motif": "视觉母题",
    "contrast": "对比策略",
    "technical_depth": "技术深度",
    "business_framing": "业务表达",
    "decision_orientation": "决策导向",
    "persuasion_intensity": "说服强度",
    "page_density": "页面密度",
}
PALETTE_LABELS = {
    "background": "页面背景",
    "surface": "卡片/表面",
    "primary": "主色",
    "accent": "强调色",
    "text_primary": "主要文字",
    "text_muted": "次要文字",
    "risk": "风险",
    "warning": "警示",
    "success": "成功/正向",
    "divider": "分隔线",
}
PAGE_ROLE_LABELS = {
    "cover": "封面",
    "toc": "目录",
    "content": "内容页",
    "timeline": "时间线",
    "summary": "总结页",
    "ending": "结束页",
}

IMG_SVG_REVIEW_MIN_SIMILARITY = 0.86
# Raster crops preserve fidelity-sensitive artwork. Text-free artwork stays a
# tight source crop. A complex text-bearing ornament may use one local
# ``complex_backplate`` crop after the helper removes its declared text pixels;
# the replacement wording must still be present as SVG text above the image.
# Broad screenshot bands and full-slide wrappers remain forbidden.
IMG_SVG_MAX_CROP_AREA_RATIO = 0.08
IMG_SVG_MAX_CROP_TOTAL_AREA_RATIO = 0.20
IMG_SVG_MAX_EMBEDDED_RASTER_RATIO = 0.25
IMG_SVG_CROP_TILE_GAP = 8.0
IMG_SVG_CROP_CONTENT_TYPES = {
    "icon",
    "logo",
    "badge",
    "decorative_symbol",
    "illustration",
    "photo",
    "texture",
    "complex_backplate",
}
IMG_SVG_CROP_TYPE_LIMITS = {
    "icon": (180.0, 180.0),
    "logo": (240.0, 180.0),
    "badge": (180.0, 180.0),
    "decorative_symbol": (180.0, 180.0),
    "illustration": (360.0, 260.0),
    "photo": (520.0, 420.0),
    "texture": (520.0, 420.0),
    "complex_backplate": (560.0, 360.0),
}
IMG_SVG_SOURCE_CROP_DEFAULT_TYPES = {
    "icon", "logo", "illustration", "photo", "texture", "complex_backplate"
}

IMG_SVG_COMPILED_PROMPT = """Task type: faithful visual tracing of an existing presentation slide.
This is NOT a slide redesign, restyling, simplification, or content-rewriting task.

Use the attached slide image as the sole and mandatory visual source of truth. Inspect the attached image directly in this same model turn while producing the SVG. Do not reconstruct the slide from a textual summary, slide plan, design system, previous memory, or a description of the image.

Priority order:
1. Every visible word, number, label, caption, and legend remains editable SVG <text>/<tspan>.
2. Pixel-level visual resemblance to the attached image at 1280x720.
3. Preservation of every visible element, including icons and decorations.
4. Exact wording, reading order, relative position, scale, alignment, spacing, and line breaks.
5. PowerPoint-compatible SVG rendering.
6. Editability of simple geometry.
7. SVG simplicity.

Canvas: width 1280, height 720, viewBox="0 0 1280 720".

Strict faithfulness requirements:
- Preserve all visible Chinese and English text verbatim. Do not rewrite, abbreviate, summarize, correct, or reorganize it.
- Preserve title position, font scale, line breaks, alignment, card geometry, border radius, spacing, margins, arrows, dividers, footer, shadows, strokes, fills, badges, icon backgrounds, and decorative marks.
- Sample colors from the attached image. Do not replace them with an approximate design-system palette unless visually indistinguishable.
- Do not normalize spacing, improve the composition, introduce a new visual style, or remove decorative elements.

Icon and illustration requirements:
- Every visible icon, logo, badge, illustration, and decorative symbol must remain.
- Never replace an original icon with a generic plus, checkmark, circle, arrow, user silhouette, database symbol, or other approximate icon.
- Inventory every visible icon, logo, badge, illustration, and decorative symbol before choosing how to reproduce it.
- Preserve icons, logos, illustrations, photos, and textures directly as tight source-image crops. Do not redraw or approximate them as vector artwork.
- Use one text-scrubbed complex_backplate only when the source region is compact, depends on non-trivial raster detail or coupled visual effects, integrates editable wording with that detail, can be repaired locally after text removal, and would suffer material visual drift under approximate vector tracing. This is a fidelity-and-editability decision, not a style-specific object rule.
- A complex_backplate may include its original text only in the source pixels. Declare tight text_removal_boxes and replacement_text_ids; the crop helper removes those text pixels before embedding, and the exact wording must be drawn afterward as SVG <text>/<tspan> above the <image>.
- Use vectors for editable text, lines, boxes, dividers, arrows, and other simple geometry. A simple badge or decorative symbol may use faithful_vector_trace only after direct source-vs-render comparison establishes near-pixel fidelity.
- A crop is the default for source-specific artwork. Ordinary source crops must stay tightly limited to isolated artwork and never become a screenshot of a card, panel, title band, chart area, or page region.
- Do not rasterize visible Chinese/English text, numbers, labels, captions, or legends. For isolated artwork, split artwork away from nearby text. For a compact inseparable styled component, use complex_backplate with declared text removal and editable SVG text replacement instead of approximating away its design character.
- Before declaring a crop, check its bounds against every nearby text baseline and shrink it until there is clear separation. Use the smallest faithful box (normally an isolated icon or artwork, with a few pixels of edge margin).
- For every crop, place a matching <image data-crop-id="..."> element and provide its exact source box in normalized 1280x720 coordinates in the crop manifest.

Vectorization requirements:
- Keep text as <text>/<tspan> when this preserves the original appearance.
- Rebuild cards, backgrounds, lines, dividers, arrows, and simple geometry as vector elements.
- Source-specific artwork may remain as small embedded Base64 PNG crops when it is explicitly inventoried. A complex_backplate is valid only after its declared source text is removed and replaced by editable SVG text.
- Do not use one full-page raster image, broad screenshot strips, or large raster patches that dominate the page.
- Rasterized visible text is an automatic failure even when the rendered SVG has a high pixel-similarity score.
- Do not split a text-bearing card, panel, band, or chart into adjacent smaller crops to evade crop limits.
- Build a complete visible-text inventory before cropping. Every inventory item must appear verbatim in SVG <text>/<tspan>, with its source-image box recorded so no crop can overlap it.
- Build a complete visual-element inventory. Every visible artwork item must be classified as source_crop or faithful_vector_trace. An empty crop list is valid only when every visible artwork item has a reviewed faithful-vector entry, or the source truly contains no artwork.

Compatibility requirements:
- No external URLs or linked files, <foreignObject>, script, animation, external stylesheet, or external font.
- Avoid unsupported filters only when they cause PowerPoint rendering problems; do not simplify visible styling unnecessarily.

Self-check before returning:
- Compare the SVG against the attached image from left to right and top to bottom.
- Confirm every visible icon and decorative element is present.
- Confirm every visible word and number is represented by SVG text, not pixels inside an <image>.
- Confirm no card, title, label, arrow, or footer has been moved, simplified, or redesigned.
- Confirm the result looks like the same slide, not a newly designed slide.

Output exactly one final 1280x720 SVG and one crop manifest. Do not output alternative SVG versions or explanatory prose."""


def default_img_svg_conversion() -> dict[str, Any]:
    return {
        "mode": None,
        "status": "waiting_for_img_export",
    }


def default_state() -> dict[str, Any]:
    return {
        "version": 6,
        "status": "new",
        "approvals": {},
        "artifacts": {},
        "renderer": None,
        "render_review": {
            "mode": None,
            "status": "not_applicable",
        },
        "renderers": {},
        "execution": "codex-skill",
    }


def default_renderer_state(renderer: str | None = None) -> dict[str, Any]:
    review = {"mode": None, "status": "not_applicable"}
    if renderer in {"html", "svg"}:
        review = {"mode": None, "status": "pending_choice"}
    elif renderer == "img":
        review = dict(RENDER_REVIEW_NOT_APPLICABLE)
    state = {
        "status": "new",
        "artifacts": {},
        "review": review,
    }
    if renderer == "img":
        state["svg_conversion"] = default_img_svg_conversion()
    return state


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")


def output_root() -> Path:
    return Path(os.getenv("OUTPUT_DIR", "output"))


def resolved_output_root() -> Path:
    root = output_root()
    if not root.is_absolute():
        root = WORKSPACE_ROOT / root
    return root.resolve()


def is_within(path: Path, parent: Path) -> bool:
    try:
        path.relative_to(parent)
    except ValueError:
        return False
    return True


def normalize_run_dir(raw: str | None, *, topic: str | None = None) -> Path:
    base = resolved_output_root()
    if raw is None:
        if not topic:
            raise ValueError("topic is required when run_dir is omitted")
        return (base / safe_filename_part(topic, max_length=80)).resolve()

    run_dir = Path(raw)
    if run_dir.is_absolute():
        resolved = run_dir.resolve()
    else:
        repo_relative = (WORKSPACE_ROOT / run_dir).resolve()
        resolved = repo_relative if is_within(repo_relative, base) else (base / run_dir).resolve()

    if not is_within(resolved, base):
        raise ValueError(f"run_dir must stay inside output root: {base}")
    return resolved


def run_dir_for_topic(topic: str) -> Path:
    return normalize_run_dir(None, topic=topic)


def read_json(path: Path, default: Any | None = None) -> Any:
    if not path.exists():
        if default is not None:
            return default
        raise FileNotFoundError(path)
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, data: Any) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def _normalize_state(state: dict[str, Any]) -> dict[str, Any]:
    merged = default_state()
    merged.update(state or {})
    merged["approvals"] = dict(merged.get("approvals") or {})
    artifacts = dict(merged.get("artifacts") or {})
    renderers = merged.get("renderers")
    if not isinstance(renderers, dict):
        renderers = {}

    active_renderer = merged.get("renderer")
    if active_renderer and active_renderer not in renderers:
        common_artifacts = {key: value for key, value in artifacts.items() if key in COMMON_ARTIFACT_KEYS}
        renderer_artifacts = {key: value for key, value in artifacts.items() if key not in COMMON_ARTIFACT_KEYS}
        if renderer_artifacts:
            artifacts = common_artifacts
        renderers[active_renderer] = {
            **default_renderer_state(active_renderer),
            "status": merged.get("status", "new"),
            "artifacts": renderer_artifacts,
            "review": dict(merged.get("render_review") or default_renderer_state(active_renderer)["review"]),
        }

    normalized_renderers: dict[str, Any] = {}
    for renderer, entry in renderers.items():
        normalized = default_renderer_state(renderer)
        if isinstance(entry, dict):
            normalized.update(entry)
            normalized["artifacts"] = dict(normalized.get("artifacts") or {})
            normalized["review"] = dict(normalized.get("review") or default_renderer_state(renderer)["review"])
        if renderer == "img":
            conversion = normalized.get("svg_conversion")
            normalized["svg_conversion"] = {
                **default_img_svg_conversion(),
                **(dict(conversion) if isinstance(conversion, dict) else {}),
            }
            # Version 5 required a no-op user response before an IMG run could
            # complete. Version 6 treats the exported IMG PPTX as the completed
            # deliverable and leaves SVG conversion as a later opt-in action.
            if normalized["svg_conversion"].get("status") == "pending_choice":
                normalized["svg_conversion"]["status"] = "available"
                normalized["svg_conversion"]["requires_explicit_opt_in"] = True
                normalized["svg_conversion"]["additional_model_usage_required"] = True
            if normalized.get("status") == "img_svg_choice_pending":
                normalized["status"] = "completed"
        normalized_renderers[renderer] = normalized

    merged["artifacts"] = artifacts
    merged["renderers"] = normalized_renderers
    if merged.get("status") == "img_svg_choice_pending":
        merged["status"] = "completed"
    merged["version"] = 6
    return merged


def renderer_state(state: dict[str, Any], renderer: str) -> dict[str, Any]:
    renderers = state.setdefault("renderers", {})
    entry = renderers.get(renderer)
    if not isinstance(entry, dict):
        entry = default_renderer_state(renderer)
        renderers[renderer] = entry
        return entry
    normalized = default_renderer_state(renderer)
    normalized.update(entry)
    normalized["artifacts"] = dict(normalized.get("artifacts") or {})
    normalized["review"] = dict(normalized.get("review") or default_renderer_state(renderer)["review"])
    if renderer == "img":
        conversion = normalized.get("svg_conversion")
        normalized["svg_conversion"] = {
            **default_img_svg_conversion(),
            **(dict(conversion) if isinstance(conversion, dict) else {}),
        }
    renderers[renderer] = normalized
    return normalized


def load_state(run_dir: Path) -> dict[str, Any]:
    return _normalize_state(read_json(run_dir / STATE_FILE, default=default_state()))


def save_state(run_dir: Path, state: dict[str, Any]) -> Path:
    normalized = _normalize_state(state)
    normalized["version"] = 6
    normalized["execution"] = "codex-skill"
    normalized["updated_at"] = utc_now()
    return write_json(run_dir / STATE_FILE, normalized)


def init_state(run_dir: Path, *, topic: str, audience: str, pages: str, research: str) -> dict[str, Any]:
    state = load_state(run_dir)
    state.update({
        "topic": topic,
        "audience": audience,
        "pages": pages,
        "research": research,
        "status": "initialized",
        "created_at": state.get("created_at") or utc_now(),
        "approvals": state.get("approvals") or {},
        "artifacts": state.get("artifacts") or {},
        "renderer": state.get("renderer"),
        "render_review": state.get("render_review") or {
            "mode": None,
            "status": "not_applicable",
        },
        "renderers": state.get("renderers") or {},
    })
    save_state(run_dir, state)
    return state


def ensure_render_ready(run_dir: Path, renderer: str) -> None:
    state = load_state(run_dir)
    if renderer not in {"html", "svg"}:
        return
    review = renderer_state(state, renderer).get("review") or {}
    review_status = review.get("status")
    review_mode = review.get("mode")
    if review_status == "pending_choice":
        raise RuntimeError(f"{renderer} review preference is not confirmed. Confirm review preference first.")
    if review_mode == "on" and review_status != "completed":
        raise RuntimeError(f"{renderer} review subflow is not complete. Complete the review subflow first.")


def mark_artifact(run_dir: Path, state: dict[str, Any], name: str, path: Path, preview_path: Path) -> None:
    state.setdefault("artifacts", {})[name] = {
        "path": str(path),
        "preview_path": str(preview_path),
        "status": "pending_review",
        "updated_at": utc_now(),
    }
    state.setdefault("approvals", {})[name] = False
    state["status"] = f"{name}_pending_review"
    save_state(run_dir, state)


def require_approved(run_dir: Path, artifact: str) -> None:
    state = load_state(run_dir)
    if not state.get("approvals", {}).get(artifact):
        preview = state.get("artifacts", {}).get(artifact, {}).get("preview_path")
        detail = f" Review {preview} first." if preview else ""
        raise RuntimeError(f"{artifact} is not approved.{detail}")


def _get_pages(outline: dict[str, Any]) -> list[dict[str, Any]]:
    if isinstance(outline.get("pages"), list):
        return outline["pages"]
    inner = outline.get("ppt_outline", outline)
    pages: list[dict[str, Any]] = []
    for key in ("cover", "table_of_contents"):
        value = inner.get(key)
        if isinstance(value, dict):
            pages.append(value)
    for part in inner.get("parts", []) if isinstance(inner.get("parts"), list) else []:
        if isinstance(part, dict):
            for page in part.get("pages", []):
                if isinstance(page, dict):
                    pages.append(page)
    end_page = inner.get("end_page")
    if isinstance(end_page, dict):
        pages.append(end_page)
    return pages


def _get_title(page: dict[str, Any]) -> str:
    return str(page.get("title") or page.get("page_title") or page.get("name") or "Untitled")


def write_outline_preview(outline: dict[str, Any], path: Path) -> Path:
    lines = ["# Outline Preview", ""]
    for index, page in enumerate(_get_pages(outline), start=1):
        lines.append(f"## {index:02d}. {_get_title(page)}")
        sections = page.get("sections") or page.get("content") or page.get("bullets") or []
        if isinstance(sections, list):
            for item in sections[:8]:
                lines.append(f"- {item}")
        elif sections:
            lines.append(f"- {sections}")
        lines.append("")
    path.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")
    return path


def write_contents_preview(contents: dict[str, Any], path: Path) -> Path:
    lines = ["# Contents Preview", ""]
    for title, material in contents.items():
        lines.extend([f"## {title}", "", str(material).strip(), ""])
    path.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")
    return path


def _is_non_empty(value: Any) -> bool:
    return value is not None and value != "" and value != [] and value != {}


def _require_non_empty_fields(value: Any, path: str, fields: tuple[str, ...]) -> dict[str, Any]:
    if not isinstance(value, dict):
        raise ValueError(f"{path} must be an object")
    missing = [field for field in fields if not _is_non_empty(value.get(field))]
    if missing:
        raise ValueError(f"{path} is missing required fields: {', '.join(missing)}")
    return value


def _require_non_empty_string(value: Any, path: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{path} must be a non-empty string")
    return value


def _require_non_empty_string_list(value: Any, path: str) -> list[str]:
    if not isinstance(value, list) or not value or not all(
        isinstance(item, str) and item.strip() for item in value
    ):
        raise ValueError(f"{path} must be a non-empty string list")
    return value


def validate_slide_plans(data: dict[str, Any]) -> None:
    if not isinstance(data, dict):
        raise ValueError("slide-plans.json must contain a JSON object")
    try:
        version = int(data.get("version", 1))
    except (TypeError, ValueError) as exc:
        raise ValueError("slide-plans.json version must be an integer") from exc
    if version != SLIDE_PLAN_VERSION:
        raise ValueError(
            "slide-plans.json must use version 2 with structured deck_strategy, "
            "design_system, and per-slide plan fields; legacy version 1/free-form plans "
            "must be regenerated before preview or approval"
        )

    deck_strategy = _require_non_empty_fields(
        data.get("deck_strategy"),
        "slide-plans.json deck_strategy",
        REQUIRED_DECK_STRATEGY_FIELDS,
    )
    for field in ("primary_audience", "decision_context", "presentation_posture"):
        _require_non_empty_string(deck_strategy[field], f"slide-plans.json deck_strategy.{field}")
    _require_non_empty_string_list(
        deck_strategy["first_questions"],
        "slide-plans.json deck_strategy.first_questions",
    )
    _require_non_empty_string_list(
        deck_strategy["evidence_order"],
        "slide-plans.json deck_strategy.evidence_order",
    )

    design_system = _require_non_empty_fields(
        data.get("design_system"),
        "slide-plans.json design_system",
        REQUIRED_DESIGN_SYSTEM_FIELDS,
    )
    for field in ("theme_name", "audience_fit", "illustration_policy"):
        _require_non_empty_string(design_system[field], f"slide-plans.json design_system.{field}")
    for field in ("typography", "component_rules", "chart_treatment"):
        if not isinstance(design_system[field], dict):
            raise ValueError(f"slide-plans.json design_system.{field} must be an object")
    _require_non_empty_string_list(
        design_system["design_genes"],
        "slide-plans.json design_system.design_genes",
    )
    palette = design_system.get("palette")
    if not isinstance(palette, dict):
        raise ValueError("slide-plans.json design_system requires a palette object")
    missing = sorted(token for token in REQUIRED_PALETTE_TOKENS if not _is_non_empty(palette.get(token)))
    if missing:
        raise ValueError(
            "slide-plans.json design_system.palette is missing required tokens: "
            + ", ".join(missing)
        )
    invalid_palette_values = sorted(
        token for token, color in palette.items() if not isinstance(color, str) or not color.strip()
    )
    if invalid_palette_values:
        raise ValueError(
            "slide-plans.json design_system.palette tokens must map to non-empty color strings: "
            + ", ".join(invalid_palette_values)
        )

    slides = data.get("slides")
    if not isinstance(slides, list) or not slides:
        raise ValueError("slide-plans.json slides must be a non-empty list")

    seen_indexes: set[int] = set()
    for position, slide in enumerate(slides, start=1):
        slide_path = f"slide-plans.json slides[{position - 1}]"
        slide = _require_non_empty_fields(
            slide,
            slide_path,
            ("index", "title", "material", "page_role", "plan"),
        )
        try:
            index = int(slide["index"])
        except (TypeError, ValueError) as exc:
            raise ValueError(f"{slide_path}.index must be an integer") from exc
        if index < 1 or index in seen_indexes:
            raise ValueError(f"{slide_path}.index must be a unique positive integer")
        seen_indexes.add(index)
        for field in ("title", "material", "page_role"):
            _require_non_empty_string(slide[field], f"{slide_path}.{field}")

        plan = _require_non_empty_fields(
            slide["plan"],
            f"{slide_path}.plan",
            REQUIRED_SLIDE_PLAN_FIELDS,
        )
        for field in ("core_message", "layout_structure"):
            _require_non_empty_string(plan[field], f"{slide_path}.plan.{field}")
        for field in (
            "visual_hierarchy",
            "required_elements",
            "palette_tokens",
            "renderer_neutral_constraints",
        ):
            _require_non_empty_string_list(plan[field], f"{slide_path}.plan.{field}")
        for field in ("style_controls", "audience_controls"):
            if not isinstance(plan[field], dict):
                raise ValueError(f"{slide_path}.plan.{field} must be an object")
        palette_tokens = plan["palette_tokens"]
        unknown_tokens = sorted(set(palette_tokens) - set(palette))
        if unknown_tokens:
            raise ValueError(
                f"{slide_path}.plan.palette_tokens references unknown design_system.palette tokens: "
                + ", ".join(unknown_tokens)
            )


def _preview_label(key: str) -> str:
    return PREVIEW_LABELS.get(key, key.replace("_", " ").strip().title())


def _append_preview_mapping(lines: list[str], value: dict[str, Any], *, indent: int = 0) -> None:
    prefix = " " * indent
    for key, item in value.items():
        label = _preview_label(str(key))
        if isinstance(item, dict):
            lines.append(f"{prefix}- **{label}**：")
            _append_preview_mapping(lines, item, indent=indent + 4)
        elif isinstance(item, list):
            lines.append(f"{prefix}- **{label}**：")
            for child in item:
                lines.append(f"{prefix}    - {str(child).strip()}")
        else:
            lines.append(f"{prefix}- **{label}**：{str(item).strip()}")


def _append_numbered_preview_list(lines: list[str], values: list[Any]) -> None:
    for index, value in enumerate(values, start=1):
        lines.append(f"{index}. {str(value).strip()}")


def _append_bulleted_preview_list(lines: list[str], values: list[Any]) -> None:
    for value in values:
        lines.append(f"- {str(value).strip()}")


def _append_deck_strategy_preview(lines: list[str], strategy: dict[str, Any]) -> None:
    lines.extend([
        "## 全局叙事策略",
        "",
        f"- **目标受众**：{strategy['primary_audience']}",
        f"- **决策目标**：{strategy['decision_context']}",
        f"- **表达姿态**：{strategy['presentation_posture']}",
        "",
        "### 受众首先会问",
        "",
    ])
    _append_numbered_preview_list(lines, strategy["first_questions"])
    lines.extend(["", "### 证据展开顺序", ""])
    _append_numbered_preview_list(lines, strategy["evidence_order"])
    extras = {
        key: value
        for key, value in strategy.items()
        if key not in REQUIRED_DECK_STRATEGY_FIELDS
    }
    if extras:
        lines.extend(["", "### 其他叙事约束", ""])
        _append_preview_mapping(lines, extras)
    lines.append("")


def _append_palette_preview(lines: list[str], palette: dict[str, Any]) -> None:
    lines.extend([
        "### 配色方案",
        "",
        "| 用途 | Token | 色值 |",
        "| --- | --- | --- |",
    ])
    for token, color in palette.items():
        lines.append(f"| {PALETTE_LABELS.get(token, _preview_label(token))} | `{token}` | `{color}` |")
    lines.append("")


def _append_design_system_preview(lines: list[str], design_system: dict[str, Any]) -> None:
    lines.extend([
        "## 全局设计规范",
        "",
        f"- **主题名称**：{design_system['theme_name']}",
        f"- **受众适配**：{design_system['audience_fit']}",
        f"- **图像与插画**：{design_system['illustration_policy']}",
        "",
    ])
    _append_palette_preview(lines, design_system["palette"])
    for key, title in (
        ("typography", "字体层级"),
        ("component_rules", "组件与间距"),
        ("chart_treatment", "图表与图解"),
    ):
        lines.extend([f"### {title}", ""])
        _append_preview_mapping(lines, design_system[key])
        lines.append("")
    lines.extend(["### 设计基因", ""])
    lines.append(" · ".join(str(item).strip() for item in design_system["design_genes"]))
    extras = {
        key: value
        for key, value in design_system.items()
        if key not in REQUIRED_DESIGN_SYSTEM_FIELDS
    }
    if extras:
        lines.extend(["", "### 其他设计约束", ""])
        _append_preview_mapping(lines, extras)
    lines.append("")


def _append_slide_plan_details(
    lines: list[str],
    plan: dict[str, Any],
    palette: dict[str, Any],
) -> None:
    lines.extend([
        "> **核心信息**",
        ">",
        f"> {plan['core_message']}",
        "",
        "### 页面布局",
        "",
        str(plan["layout_structure"]).strip(),
        "",
        "### 视觉层级",
        "",
    ])
    _append_numbered_preview_list(lines, plan["visual_hierarchy"])
    lines.extend(["", "### 页面元素", ""])
    _append_bulleted_preview_list(lines, plan["required_elements"])
    lines.extend(["", "### 本页配色", ""])
    colors = [
        f"{PALETTE_LABELS.get(token, _preview_label(token))} `{palette[token]}`"
        for token in plan["palette_tokens"]
    ]
    lines.append(" · ".join(colors))
    lines.extend(["", "### 样式控制", ""])
    _append_preview_mapping(lines, plan["style_controls"])
    lines.extend([
        "",
        "_受众控制与渲染通用约束已通过校验，并保留在 `slide-plans.json` 中供渲染阶段使用。_",
        "",
    ])


def write_plans_preview(data: dict[str, Any], path: Path) -> Path:
    validate_slide_plans(data)
    lines = ["# 幻灯片规划预览", ""]
    _append_deck_strategy_preview(lines, data["deck_strategy"])
    _append_design_system_preview(lines, data["design_system"])
    palette = data["design_system"]["palette"]
    for job in data.get("slides", []):
        role = str(job.get("page_role", "content"))
        lines.extend([
            "---",
            "",
            f"## 第 {int(job['index']):02d} 页｜{job['title']}",
            "",
            f"- **页面类型**：{PAGE_ROLE_LABELS.get(role, role)}",
            "",
        ])
        _append_slide_plan_details(lines, job["plan"], palette)
    path.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")
    return path


def preview_artifact(run_dir: Path, artifact: str) -> Path:
    data_path = run_dir / ARTIFACT_FILES[artifact]
    preview_path = run_dir / PREVIEW_FILES[artifact]
    data = read_json(data_path)
    if artifact == "outline":
        write_outline_preview(data, preview_path)
    elif artifact == "contents":
        write_contents_preview(data, preview_path)
    elif artifact == "slide_plans":
        write_plans_preview(data, preview_path)
    state = load_state(run_dir)
    mark_artifact(run_dir, state, artifact, data_path, preview_path)
    return preview_path


def slide_jobs(run_dir: Path) -> list[dict[str, Any]]:
    data = slide_plan_document(run_dir)
    return list(data.get("slides") or [])


def slide_plan_document(run_dir: Path) -> dict[str, Any]:
    data = read_json(run_dir / "slide-plans.json", default={"slides": []})
    validate_slide_plans(data)
    return data


def renderer_pptx_name(run_dir: Path, renderer: str) -> str:
    topic = safe_filename_part(infer_topic(run_dir), max_length=30)
    return f"{topic}-{renderer}.pptx"


def slide_status_path(run_dir: Path, renderer: str | None = None) -> Path:
    return run_dir / ("slide-status.json" if renderer is None else f"slide-status-{renderer}.json")


def write_slide_status(run_dir: Path, slides: dict[str, Any], renderer: str | None = None) -> Path:
    payload: dict[str, Any] = {"slides": slides}
    if renderer:
        payload["renderer"] = renderer
        write_json(slide_status_path(run_dir, renderer), payload)
    return write_json(slide_status_path(run_dir), payload)


def infer_topic(run_dir: Path) -> str:
    return str(load_state(run_dir).get("topic") or run_dir.name)


def mark_completed(run_dir: Path, renderer: str, artifacts: dict[str, Path]) -> None:
    state = load_state(run_dir)
    state["renderer"] = renderer
    entry = renderer_state(state, renderer)
    entry["status"] = "completed"
    entry["completed_at"] = utc_now()
    entry["artifacts"] = {
        **entry.get("artifacts", {}),
        **{
            name: {"path": str(path), "status": "completed"}
            for name, path in artifacts.items()
        },
    }
    state["render_review"] = dict(entry.get("review") or {})
    state["status"] = "completed"
    save_state(run_dir, state)


def render_jobs_root(run_dir: Path, renderer: str) -> Path:
    return run_dir / "render-jobs" / renderer


def render_target_path(run_dir: Path, renderer: str, index: int, title: str) -> Path:
    extensions = {
        "html": "html",
        "svg": "svg",
        "img": "png",
    }
    return run_dir / renderer / slide_filename(index, title, extensions[renderer])


def prepare_render_jobs(run_dir: Path, renderer: str) -> Path:
    require_approved(run_dir, "slide_plans")
    if renderer not in {"html", "svg", "img"}:
        raise ValueError("renderer must be html, svg, or img")

    state = load_state(run_dir)
    plan_document = slide_plan_document(run_dir)
    slides = list(plan_document.get("slides") or [])
    jobs_dir = render_jobs_root(run_dir, renderer)
    jobs_dir.mkdir(parents=True, exist_ok=True)
    target_dir = run_dir / renderer
    target_dir.mkdir(parents=True, exist_ok=True)

    shared_context_path = jobs_dir / "shared-context.json"
    shared_context = {
        "version": 2 if int(plan_document.get("version", 1)) >= 2 else 1,
        "topic": infer_topic(run_dir),
        "audience": state.get("audience"),
        "deck_strategy": plan_document.get("deck_strategy") or {},
        "design_system": plan_document.get("design_system") or {},
        "renderer": renderer,
        "slide_count": len(slides),
        "slides": [
            {
                "index": int(job.get("index", idx)),
                "title": job.get("title", f"Slide {idx}"),
                "page_role": job.get("page_role", "content"),
                "material": job.get("material", ""),
                "plan": job.get("plan", ""),
            }
            for idx, job in enumerate(slides, start=1)
        ],
    }
    write_json(shared_context_path, shared_context)

    manifest_slides: list[dict[str, Any]] = []
    for idx, job in enumerate(slides, start=1):
        slide_index = int(job.get("index", idx))
        title = str(job.get("title") or f"Slide {slide_index}")
        job_path = jobs_dir / f"slide-{slide_index:02d}.json"
        target_path = render_target_path(run_dir, renderer, slide_index, title)
        payload = {
            "version": shared_context["version"],
            "topic": shared_context["topic"],
            "audience": shared_context["audience"],
            "deck_strategy": shared_context["deck_strategy"],
            "design_system": shared_context["design_system"],
            "renderer": renderer,
            "index": slide_index,
            "title": title,
            "page_role": job.get("page_role", "content"),
            "material": job.get("material", ""),
            "plan": job.get("plan", ""),
            "total_pages": len(slides),
            "target_path": str(target_path),
            "shared_context_path": str(shared_context_path),
        }
        write_json(job_path, payload)
        manifest_slides.append({
            "index": slide_index,
            "title": title,
            "job_path": str(job_path),
            "target_path": str(target_path),
        })

    manifest_path = jobs_dir / "manifest.json"
    write_json(manifest_path, {
        "version": shared_context["version"],
        "renderer": renderer,
        "topic": shared_context["topic"],
        "audience": shared_context["audience"],
        "slide_count": len(slides),
        "shared_context_path": str(shared_context_path),
        "slides": manifest_slides,
    })

    entry = renderer_state(state, renderer)
    entry["jobs"] = {
        "manifest_path": str(manifest_path),
        "shared_context_path": str(shared_context_path),
        "status": "prepared",
        "updated_at": utc_now(),
    }
    state["renderer"] = renderer
    state["render_review"] = dict(entry.get("review") or {})
    state["status"] = "render_jobs_prepared"
    save_state(run_dir, state)
    return manifest_path


def img_source_images(run_dir: Path) -> list[Path]:
    img_dir = run_dir / "img"
    return sorted(
        [path for path in img_dir.iterdir() if path.suffix.lower() in {".png", ".jpg", ".jpeg"}]
    ) if img_dir.exists() else []


def require_complete_img_sources(run_dir: Path) -> list[Path]:
    images = img_source_images(run_dir)
    if not images:
        raise ValueError(f"IMG directory is empty: {run_dir / 'img'}")
    expected = len(slide_jobs(run_dir))
    if expected and len(images) != expected:
        raise ValueError(
            f"IMG page count mismatch: expected {expected} pages from slide-plans.json, found {len(images)}"
        )
    return images


def mark_img_exported_complete(run_dir: Path, pptx_path: Path, images: list[Path]) -> None:
    state = load_state(run_dir)
    state["renderer"] = "img"
    entry = renderer_state(state, "img")
    entry["artifacts"] = {
        **entry.get("artifacts", {}),
        "pptx": {"path": str(pptx_path), "status": "completed"},
    }
    entry["svg_conversion"] = {
        "mode": None,
        "status": "available",
        "source_pptx_path": str(pptx_path),
        "source_images": [str(path) for path in images],
        "available_after_img_export": True,
        "requires_explicit_opt_in": True,
        "additional_model_usage_required": True,
        "updated_at": utc_now(),
    }
    entry["status"] = "completed"
    entry["completed_at"] = utc_now()
    state["render_review"] = dict(entry.get("review") or {})
    state["status"] = "completed"
    save_state(run_dir, state)


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def file_sha256(path: Path) -> str:
    return sha256_bytes(path.read_bytes())


def text_sha256(value: str) -> str:
    return sha256_bytes(value.encode("utf-8"))


def img_svg_conversion_evidence_template(
    source_image_path: Path,
    prompt_sha256: str,
) -> dict[str, Any]:
    return {
        "version": 1,
        "source_image_path": str(source_image_path),
        "source_image_sha256": file_sha256(source_image_path),
        "prompt_sha256": prompt_sha256,
        "input_mode": "image_and_prompt_same_model_turn",
        "source_image_attached": True,
        "visual_source_of_truth": "source_image_only",
        "model_or_agent": "REQUIRED: record the vision-capable model or Codex agent",
        "generated_at": "REQUIRED: ISO-8601 timestamp",
    }


def img_svg_crop_manifest_template(
    source_image_path: Path,
    svg_path: Path,
) -> dict[str, Any]:
    return {
        "version": 3,
        "source_image_path": str(source_image_path),
        "svg_path": str(svg_path),
        "canvas": {"width": 1280, "height": 720},
        "icon_strategy": "REQUIRED: source_crops | mixed | faithful_vector_trace | no_icons_visible",
        "crop_policy": {
            "max_area_ratio": IMG_SVG_MAX_CROP_AREA_RATIO,
            "max_total_area_ratio": IMG_SVG_MAX_CROP_TOTAL_AREA_RATIO,
            "max_embedded_raster_ratio": IMG_SVG_MAX_EMBEDDED_RASTER_RATIO,
            "type_limits": {
                key: {"max_width": value[0], "max_height": value[1]}
                for key, value in IMG_SVG_CROP_TYPE_LIMITS.items()
            },
            "text_must_remain_vector": True,
            "text_scrubbed_complex_backplates_allowed": True,
            "adjacent_crop_tiles_forbidden": True,
        },
        "visible_text_inventory": {
            "complete": False,
            "items": [],
            "no_visible_text_reason": "REQUIRED only when the source image contains no visible text",
        },
        "visual_element_inventory": {
            "complete": False,
            "items": [],
            "no_visible_artwork_reason": "REQUIRED only when the source image contains no visible icons, logos, illustrations, badges, or decorative symbols",
        },
        "crops": [],
        "no_crops_reason": "REQUIRED when crops is empty",
    }


def render_img_svg_review(
    source_image_path: Path,
    svg_path: Path,
    preview_path: Path,
) -> dict[str, float]:
    from pptx_builder import svg_to_png_bytes

    preview_bytes = svg_to_png_bytes(svg_path)
    preview_path.parent.mkdir(parents=True, exist_ok=True)
    preview_path.write_bytes(preview_bytes)

    with Image.open(source_image_path) as source, Image.open(io.BytesIO(preview_bytes)) as rendered:
        source_rgb = source.convert("RGB").resize((1280, 720), Image.Resampling.LANCZOS)
        rendered_rgb = rendered.convert("RGB").resize((1280, 720), Image.Resampling.LANCZOS)

        pixel_difference = ImageChops.difference(source_rgb, rendered_rgb)
        pixel_mae = sum(ImageStat.Stat(pixel_difference).mean) / 3.0
        pixel_similarity = max(0.0, 1.0 - pixel_mae / 255.0)

        source_edges = source_rgb.filter(ImageFilter.FIND_EDGES)
        rendered_edges = rendered_rgb.filter(ImageFilter.FIND_EDGES)
        edge_difference = ImageChops.difference(source_edges, rendered_edges)
        edge_mae = sum(ImageStat.Stat(edge_difference).mean) / 3.0
        edge_similarity = max(0.0, 1.0 - edge_mae / 255.0)

    combined_similarity = pixel_similarity * 0.65 + edge_similarity * 0.35
    return {
        "pixel_similarity": round(pixel_similarity, 6),
        "edge_similarity": round(edge_similarity, 6),
        "combined_similarity": round(combined_similarity, 6),
        "recommended_minimum": IMG_SVG_REVIEW_MIN_SIMILARITY,
    }


def write_img_svg_review_artifact(job: dict[str, Any]) -> dict[str, Any]:
    source_image_path = Path(job["source_image_path"])
    svg_path = Path(job["target_path"])
    preview_path = Path(job["rendered_preview_path"])
    review_path = Path(job["fidelity_review_path"])
    source_hash = file_sha256(source_image_path)
    svg_hash = file_sha256(svg_path)
    prompt_hash = str(job["compiled_prompt_sha256"])
    metrics = render_img_svg_review(source_image_path, svg_path, preview_path)

    previous = read_json(review_path, {}) if review_path.exists() else {}
    unchanged = (
        previous.get("source_image_sha256") == source_hash
        and previous.get("svg_sha256") == svg_hash
        and previous.get("prompt_sha256") == prompt_hash
    )
    review = {
        "version": 1,
        "slide_index": int(job["index"]),
        "source_image_path": str(source_image_path),
        "source_image_sha256": source_hash,
        "svg_path": str(svg_path),
        "svg_sha256": svg_hash,
        "rendered_preview_path": str(preview_path),
        "prompt_sha256": prompt_hash,
        "automatic_metrics": metrics,
        "status": previous.get("status", "pending_visual_review") if unchanged else "pending_visual_review",
        "source_image_inspected": bool(previous.get("source_image_inspected")) if unchanged else False,
        "rendered_svg_inspected": bool(previous.get("rendered_svg_inspected")) if unchanged else False,
        "layout_preserved": bool(previous.get("layout_preserved")) if unchanged else False,
        "icons_preserved": bool(previous.get("icons_preserved")) if unchanged else False,
        "source_specific_artwork_preserved": bool(previous.get("source_specific_artwork_preserved")) if unchanged else False,
        "all_visible_text_editable": bool(previous.get("all_visible_text_editable")) if unchanged else False,
        "no_redesign": bool(previous.get("no_redesign")) if unchanged else False,
        "reviewer": str(previous.get("reviewer") or "") if unchanged else "",
        "notes": str(previous.get("notes") or "") if unchanged else "",
        "low_similarity_override_reason": str(previous.get("low_similarity_override_reason") or "") if unchanged else "",
        "updated_at": utc_now(),
    }
    write_json(review_path, review)
    return review


def validate_img_svg_conversion_evidence(job: dict[str, Any]) -> dict[str, Any]:
    path = Path(job["conversion_evidence_path"])
    if not path.is_file():
        raise ValueError(f"IMG-to-SVG conversion evidence is missing: {path}")
    evidence = read_json(path)
    source_path = Path(job["source_image_path"])
    if Path(str(evidence.get("source_image_path") or "")).resolve() != source_path.resolve():
        raise ValueError(f"IMG-to-SVG evidence source path mismatch: {path}")
    if evidence.get("source_image_sha256") != file_sha256(source_path):
        raise ValueError(f"IMG-to-SVG evidence source hash mismatch: {path}")
    if evidence.get("prompt_sha256") != job.get("compiled_prompt_sha256"):
        raise ValueError(f"IMG-to-SVG evidence prompt hash mismatch: {path}")
    if evidence.get("input_mode") != "image_and_prompt_same_model_turn":
        raise ValueError(f"IMG-to-SVG evidence must record image_and_prompt_same_model_turn: {path}")
    if evidence.get("source_image_attached") is not True:
        raise ValueError(f"IMG-to-SVG evidence must confirm the source image was attached: {path}")
    if evidence.get("visual_source_of_truth") != "source_image_only":
        raise ValueError(f"IMG-to-SVG evidence must use source_image_only as visual truth: {path}")
    if not str(evidence.get("model_or_agent") or "").strip() or str(evidence.get("model_or_agent")).startswith("REQUIRED"):
        raise ValueError(f"IMG-to-SVG evidence must record the model or agent: {path}")
    if not str(evidence.get("generated_at") or "").strip() or str(evidence.get("generated_at")).startswith("REQUIRED"):
        raise ValueError(f"IMG-to-SVG evidence must record generated_at: {path}")
    return evidence


def _rects_intersect(first: tuple[float, float, float, float], second: tuple[float, float, float, float]) -> bool:
    fx, fy, fw, fh = first
    sx, sy, sw, sh = second
    return min(fx + fw, sx + sw) > max(fx, sx) and min(fy + fh, sy + sh) > max(fy, sy)


def _normalize_img_svg_text(value: str) -> str:
    return "".join(str(value).split())


def _svg_visible_text_items(svg_path: Path) -> list[str]:
    try:
        root = ElementTree.fromstring(svg_path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, ElementTree.ParseError):
        return []
    fragments: list[str] = []
    for element in root.iter():
        if element.tag.rsplit("}", 1)[-1].lower() == "text":
            normalized = _normalize_img_svg_text("".join(element.itertext()))
            if normalized:
                fragments.append(normalized)
    return fragments


def _svg_crop_ids(svg_path: Path) -> list[str]:
    try:
        root = ElementTree.fromstring(svg_path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, ElementTree.ParseError):
        return []
    return [
        str(element.attrib.get("data-crop-id") or "").strip()
        for element in root.iter()
        if element.tag.rsplit("}", 1)[-1].lower() == "image"
    ]


def _svg_text_fragments_after_crop(svg_path: Path, crop_id: str) -> list[str]:
    """Return SVG text that is painted after one raster backplate.

    SVG document order controls z-order. A text-scrubbed backplate is useful
    only when its replacement text is emitted later and therefore remains
    visible/editable above the embedded image.
    """
    try:
        root = ElementTree.fromstring(svg_path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, ElementTree.ParseError):
        return []
    elements = list(root.iter())
    crop_index = next(
        (
            index
            for index, element in enumerate(elements)
            if element.tag.rsplit("}", 1)[-1].lower() == "image"
            and str(element.attrib.get("data-crop-id") or "").strip() == crop_id
        ),
        None,
    )
    if crop_index is None:
        return []
    fragments: list[str] = []
    for element in elements[crop_index + 1:]:
        if element.tag.rsplit("}", 1)[-1].lower() == "text":
            normalized = _normalize_img_svg_text("".join(element.itertext()))
            if normalized:
                fragments.append(normalized)
    return fragments


def _rects_form_crop_tiles(
    first: tuple[float, float, float, float],
    second: tuple[float, float, float, float],
) -> bool:
    fx, fy, fw, fh = first
    sx, sy, sw, sh = second
    horizontal_gap = max(sx - (fx + fw), fx - (sx + sw), 0.0)
    vertical_gap = max(sy - (fy + fh), fy - (sy + sh), 0.0)
    vertical_overlap = max(0.0, min(fy + fh, sy + sh) - max(fy, sy))
    horizontal_overlap = max(0.0, min(fx + fw, sx + sw) - max(fx, sx))
    vertical_alignment = vertical_overlap / max(1.0, min(fh, sh))
    horizontal_alignment = horizontal_overlap / max(1.0, min(fw, sw))
    return (
        horizontal_gap <= IMG_SVG_CROP_TILE_GAP and vertical_alignment >= 0.65
    ) or (
        vertical_gap <= IMG_SVG_CROP_TILE_GAP and horizontal_alignment >= 0.65
    )


def _approximate_svg_text_bounds(svg_path: Path) -> list[tuple[float, float, float, float]]:
    """Approximate visible <text> bounds to catch crops that swallow vector text.

    SVG text has no portable layout API. This conservative estimate is only a
    safety net; the manifest's explicit contains_text=false declaration and
    the model prompt remain the primary crop policy.
    """
    try:
        root = ElementTree.fromstring(svg_path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, ElementTree.ParseError):
        return []

    bounds: list[tuple[float, float, float, float]] = []
    for element in root.iter():
        if element.tag.rsplit("}", 1)[-1].lower() != "text":
            continue
        if element.attrib.get("transform"):
            # Transformed text is uncommon in generated slides; skip it rather
            # than inventing a coordinate transform and risk false positives.
            continue
        try:
            x = float((element.attrib.get("x") or "0").split()[0])
            y = float((element.attrib.get("y") or "0").split()[0])
            font_size = float(str(element.attrib.get("font-size") or "16").replace("px", ""))
        except (TypeError, ValueError):
            continue
        font_size = max(6.0, min(font_size, 160.0))
        lines: list[str] = []
        if element.text and element.text.strip():
            lines.append(element.text.strip())
        for child in element:
            if child.tag.rsplit("}", 1)[-1].lower() == "tspan":
                lines.append("".join(child.itertext()).strip())
        lines = [line for line in lines if line]
        if not lines:
            continue
        max_chars = max(len(line) for line in lines)
        # CJK glyphs are close to 1em; Latin text is narrower. The slightly
        # generous estimate is intentional so a crop cannot touch a label.
        width = max(12.0, max_chars * font_size * 0.95)
        text_anchor = str(element.attrib.get("text-anchor") or "start").strip().lower()
        if text_anchor == "middle":
            x -= width / 2.0
        elif text_anchor == "end":
            x -= width
        line_height = font_size * 1.25
        top = y - font_size
        height = max(line_height, len(lines) * line_height)
        bounds.append((x, top, width, height))
    return bounds


def validate_img_svg_crop_manifest(job: dict[str, Any], svg_stats: dict[str, Any]) -> dict[str, Any]:
    path = Path(job["crop_manifest_path"])
    if not path.is_file():
        raise ValueError(f"IMG-to-SVG crop manifest is missing: {path}")
    manifest = read_json(path)
    source_path = Path(job["source_image_path"])
    svg_path = Path(job["target_path"])
    if Path(str(manifest.get("source_image_path") or "")).resolve() != source_path.resolve():
        raise ValueError(f"IMG-to-SVG crop manifest source path mismatch: {path}")
    if Path(str(manifest.get("svg_path") or "")).resolve() != svg_path.resolve():
        raise ValueError(f"IMG-to-SVG crop manifest SVG path mismatch: {path}")
    canvas = manifest.get("canvas") or {}
    if canvas != {"width": 1280, "height": 720}:
        raise ValueError(f"IMG-to-SVG crop manifest canvas must be 1280x720: {path}")
    if manifest.get("version") != 3:
        raise ValueError(
            f"IMG-to-SVG crop manifest must use version 3 with text and visual inventories: {path}"
        )

    inventory = manifest.get("visible_text_inventory")
    if not isinstance(inventory, dict) or inventory.get("complete") is not True:
        raise ValueError(f"IMG-to-SVG crop manifest needs a complete visible_text_inventory: {path}")
    inventory_items = inventory.get("items")
    if not isinstance(inventory_items, list):
        raise ValueError(f"IMG-to-SVG visible_text_inventory items must be a list: {path}")

    text_boxes: list[tuple[float, float, float, float]] = []
    text_counts: Counter[str] = Counter()
    inventory_ids: set[str] = set()
    inventory_by_id: dict[str, tuple[str, tuple[float, float, float, float]]] = {}
    for index, item in enumerate(inventory_items):
        if not isinstance(item, dict):
            raise ValueError(f"IMG-to-SVG text inventory entry {index} must be an object: {path}")
        text_id = str(item.get("id") or "").strip()
        if not text_id or text_id in inventory_ids:
            raise ValueError(f"IMG-to-SVG text inventory entry {index} needs a unique id: {path}")
        inventory_ids.add(text_id)
        normalized_text = _normalize_img_svg_text(str(item.get("text") or ""))
        if not normalized_text:
            raise ValueError(f"IMG-to-SVG text inventory {text_id} needs exact visible text: {path}")
        raw_box = item.get("source_box")
        if not isinstance(raw_box, list) or len(raw_box) != 4:
            raise ValueError(
                f"IMG-to-SVG text inventory {text_id} source_box must be [x, y, width, height]: {path}"
            )
        try:
            x, y, width, height = (float(value) for value in raw_box)
        except (TypeError, ValueError) as exc:
            raise ValueError(f"IMG-to-SVG text inventory {text_id} source_box must be numeric: {path}") from exc
        if x < 0 or y < 0 or width <= 0 or height <= 0 or x + width > 1280 or y + height > 720:
            raise ValueError(f"IMG-to-SVG text inventory {text_id} must stay inside 1280x720: {path}")
        text_rect = (x, y, width, height)
        text_boxes.append(text_rect)
        inventory_by_id[text_id] = (normalized_text, text_rect)
        text_counts[normalized_text] += 1

    if not inventory_items:
        reason = str(inventory.get("no_visible_text_reason") or "").strip()
        if len(reason) < 20:
            raise ValueError(
                f"IMG-to-SVG empty visible_text_inventory needs a specific no_visible_text_reason: {path}"
            )

    svg_text_items = _svg_visible_text_items(svg_path)
    missing_text = [
        text for text, count in text_counts.items()
        if sum(fragment.count(text) for fragment in svg_text_items) < count
    ]
    if missing_text:
        preview = ", ".join(missing_text[:5])
        raise ValueError(
            f"IMG-to-SVG visible text is missing from SVG <text>/<tspan> ({preview}); "
            f"rasterized text is not editable: {path}"
        )

    crops = manifest.get("crops")
    if not isinstance(crops, list):
        raise ValueError(f"IMG-to-SVG crop manifest crops must be a list: {path}")
    icon_strategy = str(manifest.get("icon_strategy") or "")
    allowed_strategies = {"source_crops", "mixed", "faithful_vector_trace", "no_icons_visible"}
    if icon_strategy not in allowed_strategies:
        raise ValueError(f"IMG-to-SVG crop manifest has invalid icon_strategy: {path}")

    visual_inventory = manifest.get("visual_element_inventory")
    if not isinstance(visual_inventory, dict) or visual_inventory.get("complete") is not True:
        raise ValueError(f"IMG-to-SVG crop manifest needs a complete visual_element_inventory: {path}")
    visual_items = visual_inventory.get("items")
    if not isinstance(visual_items, list):
        raise ValueError(f"IMG-to-SVG visual_element_inventory items must be a list: {path}")

    visual_ids: set[str] = set()
    source_crop_visuals: dict[str, dict[str, Any]] = {}
    traced_visuals: dict[str, dict[str, Any]] = {}
    for index, item in enumerate(visual_items):
        if not isinstance(item, dict):
            raise ValueError(f"IMG-to-SVG visual inventory entry {index} must be an object: {path}")
        visual_id = str(item.get("id") or "").strip()
        if not visual_id or visual_id in visual_ids:
            raise ValueError(f"IMG-to-SVG visual inventory entry {index} needs a unique id: {path}")
        visual_ids.add(visual_id)
        content_type = str(item.get("content_type") or "").strip().lower()
        if content_type not in IMG_SVG_CROP_CONTENT_TYPES:
            allowed = ", ".join(sorted(IMG_SVG_CROP_CONTENT_TYPES))
            raise ValueError(
                f"IMG-to-SVG visual inventory {visual_id} must declare content_type ({allowed}): {path}"
            )
        raw_box = item.get("source_box")
        if not isinstance(raw_box, list) or len(raw_box) != 4:
            raise ValueError(
                f"IMG-to-SVG visual inventory {visual_id} source_box must be [x, y, width, height]: {path}"
            )
        try:
            x, y, width, height = (float(value) for value in raw_box)
        except (TypeError, ValueError) as exc:
            raise ValueError(
                f"IMG-to-SVG visual inventory {visual_id} source_box must be numeric: {path}"
            ) from exc
        if x < 0 or y < 0 or width <= 0 or height <= 0 or x + width > 1280 or y + height > 720:
            raise ValueError(f"IMG-to-SVG visual inventory {visual_id} must stay inside 1280x720: {path}")
        strategy = str(item.get("strategy") or "").strip()
        if strategy not in {"source_crop", "faithful_vector_trace"}:
            raise ValueError(
                f"IMG-to-SVG visual inventory {visual_id} needs source_crop or faithful_vector_trace strategy: {path}"
            )
        if content_type in IMG_SVG_SOURCE_CROP_DEFAULT_TYPES and strategy != "source_crop":
            raise ValueError(
                f"IMG-to-SVG visual inventory {visual_id} with content_type {content_type} "
                f"must use source_crop; keep only text, lines, boxes, and simple geometry vector: {path}"
            )
        if item.get("fidelity_reviewed") is not True:
            raise ValueError(
                f"IMG-to-SVG visual inventory {visual_id} must confirm fidelity_reviewed=true: {path}"
            )
        notes = str(item.get("notes") or "").strip()
        if len(notes) < 20:
            raise ValueError(
                f"IMG-to-SVG visual inventory {visual_id} needs concrete source-vs-render fidelity notes: {path}"
            )
        if strategy == "source_crop":
            source_crop_visuals[visual_id] = item
        else:
            traced_visuals[visual_id] = item

    if not visual_items:
        reason = str(visual_inventory.get("no_visible_artwork_reason") or "").strip()
        if len(reason) < 20:
            raise ValueError(
                f"IMG-to-SVG empty visual_element_inventory needs a specific no_visible_artwork_reason: {path}"
            )
        if icon_strategy != "no_icons_visible":
            raise ValueError(
                f"IMG-to-SVG empty visual_element_inventory requires no_icons_visible strategy: {path}"
            )
    elif source_crop_visuals and traced_visuals and icon_strategy != "mixed":
        raise ValueError(f"IMG-to-SVG mixed visual strategies require mixed icon_strategy: {path}")
    elif source_crop_visuals and not traced_visuals and icon_strategy != "source_crops":
        raise ValueError(f"IMG-to-SVG source-crop visuals require source_crops icon_strategy: {path}")
    elif traced_visuals and not source_crop_visuals and icon_strategy != "faithful_vector_trace":
        raise ValueError(
            f"IMG-to-SVG traced visuals require faithful_vector_trace icon_strategy: {path}"
        )

    declared_crops_by_id = {
        str(item.get("id") or "").strip(): item
        for item in crops
        if isinstance(item, dict) and str(item.get("id") or "").strip()
    }
    if set(declared_crops_by_id) != set(source_crop_visuals):
        raise ValueError(
            f"IMG-to-SVG source_crop visual inventory ids must exactly match crop ids: {path}"
        )
    for crop_id, visual_item in source_crop_visuals.items():
        crop_item = declared_crops_by_id[crop_id]
        if list(visual_item.get("source_box") or []) != list(crop_item.get("source_box") or []):
            raise ValueError(
                f"IMG-to-SVG visual inventory and crop source_box mismatch for {crop_id}: {path}"
            )
        if str(visual_item.get("content_type") or "").strip().lower() != str(
            crop_item.get("content_type") or ""
        ).strip().lower():
            raise ValueError(
                f"IMG-to-SVG visual inventory and crop content_type mismatch for {crop_id}: {path}"
            )
    if crops:
        if icon_strategy not in {"source_crops", "mixed"}:
            raise ValueError(f"IMG-to-SVG crops require source_crops or mixed icon_strategy: {path}")
        svg_crop_ids = _svg_crop_ids(svg_path)
        declared_crop_ids = [str(item.get("id") or "").strip() for item in crops if isinstance(item, dict)]
        if any(not crop_id for crop_id in svg_crop_ids) or Counter(svg_crop_ids) != Counter(declared_crop_ids):
            raise ValueError(
                f"IMG-to-SVG SVG <image data-crop-id> values must exactly match the crop manifest: {path}"
            )
        text_bounds = _approximate_svg_text_bounds(Path(job["target_path"]))
        total_area = 1280.0 * 720.0
        crop_rects: list[tuple[str, tuple[float, float, float, float]]] = []
        crop_area = 0.0
        seen_crop_ids: set[str] = set()
        for index, item in enumerate(crops):
            if not isinstance(item, dict):
                raise ValueError(f"IMG-to-SVG crop entry {index} must be an object: {path}")
            crop_id = str(item.get("id") or "").strip()
            if not crop_id or crop_id in seen_crop_ids:
                raise ValueError(f"IMG-to-SVG crop entry {index} needs a unique stable id: {path}")
            seen_crop_ids.add(crop_id)
            content_type = str(item.get("content_type") or "").strip().lower()
            if content_type not in IMG_SVG_CROP_CONTENT_TYPES:
                allowed = ", ".join(sorted(IMG_SVG_CROP_CONTENT_TYPES))
                raise ValueError(
                    f"IMG-to-SVG crop {crop_id} must declare content_type ({allowed}): {path}"
                )
            if item.get("contains_text") is not False:
                raise ValueError(
                    f"IMG-to-SVG crop {crop_id} must declare contains_text=false; keep all text as SVG: {path}"
                )
            raw_box = item.get("source_box")
            if not isinstance(raw_box, list) or len(raw_box) != 4:
                raise ValueError(f"IMG-to-SVG crop {crop_id} source_box must be [x, y, width, height]: {path}")
            try:
                x, y, width, height = (float(value) for value in raw_box)
            except (TypeError, ValueError) as exc:
                raise ValueError(f"IMG-to-SVG crop {crop_id} source_box must be numeric: {path}") from exc
            if x < 0 or y < 0 or width <= 0 or height <= 0 or x + width > 1280 or y + height > 720:
                raise ValueError(f"IMG-to-SVG crop {crop_id} source_box must stay inside 1280x720: {path}")
            max_width, max_height = IMG_SVG_CROP_TYPE_LIMITS[content_type]
            if width > max_width or height > max_height:
                raise ValueError(
                    f"IMG-to-SVG crop {crop_id} is too broad for {content_type} "
                    f"({width:g}x{height:g}, max {max_width:g}x{max_height:g}); "
                    f"crop only the incompatible artwork: {path}"
                )
            if width * height > total_area * IMG_SVG_MAX_CROP_AREA_RATIO:
                raise ValueError(
                    f"IMG-to-SVG crop {crop_id} covers more than {IMG_SVG_MAX_CROP_AREA_RATIO:.0%} of the slide: {path}"
                )
            crop_rect = (x, y, width, height)
            crop_area += width * height
            crop_rects.append((crop_id, crop_rect))
            overlapping_inventory_ids = {
                text_id
                for text_id, (_, text_rect) in inventory_by_id.items()
                if _rects_intersect(crop_rect, text_rect)
            }
            is_complex_backplate = content_type == "complex_backplate"
            if is_complex_backplate:
                if item.get("source_contains_text") is not True:
                    raise ValueError(
                        f"IMG-to-SVG complex_backplate {crop_id} must declare source_contains_text=true: {path}"
                    )
                replacement_ids = item.get("replacement_text_ids")
                if not isinstance(replacement_ids, list) or not replacement_ids:
                    raise ValueError(
                        f"IMG-to-SVG complex_backplate {crop_id} needs replacement_text_ids: {path}"
                    )
                replacement_id_set = {str(value).strip() for value in replacement_ids if str(value).strip()}
                if len(replacement_id_set) != len(replacement_ids):
                    raise ValueError(
                        f"IMG-to-SVG complex_backplate {crop_id} replacement_text_ids must be unique and non-empty: {path}"
                    )
                missing_inventory_ids = replacement_id_set - set(inventory_by_id)
                if missing_inventory_ids:
                    raise ValueError(
                        f"IMG-to-SVG complex_backplate {crop_id} references unknown replacement text ids "
                        f"{sorted(missing_inventory_ids)}: {path}"
                    )
                if replacement_id_set != overlapping_inventory_ids:
                    raise ValueError(
                        f"IMG-to-SVG complex_backplate {crop_id} replacement_text_ids must exactly match "
                        f"the visible text inside the backplate ({sorted(overlapping_inventory_ids)}): {path}"
                    )
                removal_mode = str(item.get("text_removal_mode") or "").strip()
                if removal_mode not in {"light_neutral", "dark_neutral", "all"}:
                    raise ValueError(
                        f"IMG-to-SVG complex_backplate {crop_id} text_removal_mode must be "
                        f"light_neutral, dark_neutral, or all: {path}"
                    )
                dilation = item.get("text_removal_dilation", 2)
                if not isinstance(dilation, int) or dilation < 0 or dilation > 4:
                    raise ValueError(
                        f"IMG-to-SVG complex_backplate {crop_id} text_removal_dilation must be an integer 0-4: {path}"
                    )
                removal_boxes = item.get("text_removal_boxes")
                if not isinstance(removal_boxes, list) or not removal_boxes:
                    raise ValueError(
                        f"IMG-to-SVG complex_backplate {crop_id} needs tight text_removal_boxes: {path}"
                    )
                parsed_removal_boxes: list[tuple[float, float, float, float]] = []
                for removal_box in removal_boxes:
                    if not isinstance(removal_box, list) or len(removal_box) != 4:
                        raise ValueError(
                            f"IMG-to-SVG complex_backplate {crop_id} has invalid text_removal_boxes: {path}"
                        )
                    try:
                        removal_rect = tuple(float(value) for value in removal_box)
                    except (TypeError, ValueError) as exc:
                        raise ValueError(
                            f"IMG-to-SVG complex_backplate {crop_id} text_removal_boxes must be numeric: {path}"
                        ) from exc
                    rx, ry, rw, rh = removal_rect
                    if (
                        rw <= 0 or rh <= 0 or rx < x or ry < y
                        or rx + rw > x + width or ry + rh > y + height
                    ):
                        raise ValueError(
                            f"IMG-to-SVG complex_backplate {crop_id} text_removal_boxes must stay inside the crop: {path}"
                        )
                    parsed_removal_boxes.append(removal_rect)
                for replacement_id in replacement_id_set:
                    replacement_text, replacement_rect = inventory_by_id[replacement_id]
                    if not any(
                        _rects_intersect(removal_rect, replacement_rect)
                        for removal_rect in parsed_removal_boxes
                    ):
                        raise ValueError(
                            f"IMG-to-SVG complex_backplate {crop_id} has no removal box for "
                            f"replacement text {replacement_id}: {path}"
                        )
                    later_text = _svg_text_fragments_after_crop(svg_path, crop_id)
                    if not any(replacement_text in fragment for fragment in later_text):
                        raise ValueError(
                            f"IMG-to-SVG complex_backplate {crop_id} replacement text {replacement_id} "
                            f"must be emitted as SVG text after the <image>: {path}"
                        )
            else:
                if item.get("source_contains_text") is True:
                    raise ValueError(
                        f"IMG-to-SVG crop {crop_id} may declare source text only as complex_backplate: {path}"
                    )
                if overlapping_inventory_ids:
                    raise ValueError(
                        f"IMG-to-SVG crop {crop_id} overlaps the visible-text inventory; "
                        f"keep every word editable or use a text-scrubbed complex_backplate: {path}"
                    )
                if any(_rects_intersect(crop_rect, text_rect) for text_rect in text_bounds):
                    raise ValueError(
                        f"IMG-to-SVG crop {crop_id} overlaps vector text; shrink the crop or use a "
                        f"text-scrubbed complex_backplate: {path}"
                    )
            exclusion_boxes = item.get("text_exclusion_boxes") or []
            if not isinstance(exclusion_boxes, list):
                raise ValueError(f"IMG-to-SVG crop {crop_id} text_exclusion_boxes must be a list: {path}")
            for exclusion in exclusion_boxes:
                if not isinstance(exclusion, list) or len(exclusion) != 4:
                    raise ValueError(f"IMG-to-SVG crop {crop_id} has invalid text_exclusion_boxes: {path}")
                try:
                    exclusion_rect = tuple(float(value) for value in exclusion)
                except (TypeError, ValueError) as exc:
                    raise ValueError(f"IMG-to-SVG crop {crop_id} text_exclusion_boxes must be numeric: {path}") from exc
                if _rects_intersect(crop_rect, exclusion_rect):
                    raise ValueError(
                        f"IMG-to-SVG crop {crop_id} intersects a declared text exclusion box; keep text outside crops: {path}"
                    )
        if crop_area > total_area * IMG_SVG_MAX_CROP_TOTAL_AREA_RATIO:
            raise ValueError(
                f"IMG-to-SVG crop manifest covers more than {IMG_SVG_MAX_CROP_TOTAL_AREA_RATIO:.0%} "
                f"of the slide in aggregate; reconstruct cards and text as vectors: {path}"
            )
        for first_index, (first_id, first_rect) in enumerate(crop_rects):
            for second_id, second_rect in crop_rects[first_index + 1:]:
                if _rects_form_crop_tiles(first_rect, second_rect):
                    raise ValueError(
                        f"IMG-to-SVG crops {first_id} and {second_id} form adjacent/overlapping tiles; "
                        f"do not split a text-bearing region to evade crop limits: {path}"
                    )
    else:
        if svg_stats.get("embedded_image_count", 0):
            raise ValueError(
                f"IMG-to-SVG SVG contains embedded images but the crop manifest is empty; "
                f"declare every <image data-crop-id> crop: {path}"
            )
        reason = str(manifest.get("no_crops_reason") or "").strip()
        if len(reason) < 20:
            raise ValueError(f"IMG-to-SVG empty crop manifest needs a specific no_crops_reason: {path}")
        if icon_strategy not in {"faithful_vector_trace", "no_icons_visible"}:
            raise ValueError(f"IMG-to-SVG empty crop manifest needs a trace/no-icons strategy: {path}")
    return manifest


def validate_img_svg_fidelity_review(job: dict[str, Any], review: dict[str, Any]) -> dict[str, Any]:
    path = Path(job["fidelity_review_path"])
    if review.get("status") != "pass":
        raise RuntimeError(f"IMG-to-SVG visual review is pending: {path}")
    required_true = (
        "source_image_inspected",
        "rendered_svg_inspected",
        "layout_preserved",
        "icons_preserved",
        "source_specific_artwork_preserved",
        "all_visible_text_editable",
        "no_redesign",
    )
    missing = [key for key in required_true if review.get(key) is not True]
    if missing:
        raise ValueError(f"IMG-to-SVG visual review is missing confirmations {missing}: {path}")
    if not str(review.get("reviewer") or "").strip():
        raise ValueError(f"IMG-to-SVG visual review must record a reviewer: {path}")
    if len(str(review.get("notes") or "").strip()) < 12:
        raise ValueError(f"IMG-to-SVG visual review needs concrete notes: {path}")
    similarity = float((review.get("automatic_metrics") or {}).get("combined_similarity", 0.0))
    if similarity < IMG_SVG_REVIEW_MIN_SIMILARITY:
        override = str(review.get("low_similarity_override_reason") or "").strip()
        if len(override) < 20:
            raise ValueError(
                f"IMG-to-SVG similarity {similarity:.3f} is below {IMG_SVG_REVIEW_MIN_SIMILARITY:.2f}; "
                f"revise the slide or record a concrete low_similarity_override_reason: {path}"
            )
    return review


def prepare_img_svg_jobs(run_dir: Path, state: dict[str, Any]) -> Path:
    images = require_complete_img_sources(run_dir)
    jobs = slide_jobs(run_dir)
    jobs_dir = render_jobs_root(run_dir, "img-svg")
    target_dir = run_dir / "img-svg"
    reviews_dir = jobs_dir / "reviews"
    jobs_dir.mkdir(parents=True, exist_ok=True)
    target_dir.mkdir(parents=True, exist_ok=True)

    # A fresh opt-in must regenerate every SVG from the current IMG pages.
    for stale in target_dir.glob("*.svg"):
        stale.unlink()
    for stale in jobs_dir.glob("*.json"):
        stale.unlink()
    if reviews_dir.exists():
        shutil.rmtree(reviews_dir)
    reviews_dir.mkdir(parents=True, exist_ok=True)

    prompt_sha256 = text_sha256(IMG_SVG_COMPILED_PROMPT)

    shared_context_path = jobs_dir / "shared-context.json"
    write_json(shared_context_path, {
        "version": 2,
        "topic": infer_topic(run_dir),
        "audience": state.get("audience"),
        "renderer": "img-svg",
        "source_renderer": "img",
        "slide_count": len(images),
        "task_type": "faithful_visual_tracing_not_redesign",
        "priority_order": [
            "editable_visible_text",
            "pixel_level_visual_resemblance",
            "preserve_every_visible_element_and_icon",
            "exact_copy_geometry_and_reading_order",
            "powerpoint_compatibility",
            "editability",
            "svg_simplicity",
        ],
        "model_input_rule": (
            "Attach each source_image_path and the compiled_prompt in the same vision-model turn. "
            "The source image is the sole visual truth; this is tracing, not redesign."
        ),
        "output_rule": (
            "Write one final 1280x720 hybrid SVG and one version-3 crop manifest per source image. "
            "Inventory text and artwork; rebuild stable geometry as vectors and use exact tight source "
            "crops by default for icons, logos, illustrations, photos, and textures."
        ),
        "compiled_prompt": IMG_SVG_COMPILED_PROMPT,
        "compiled_prompt_sha256": prompt_sha256,
        "completion_rule": (
            "complete-img-svg renders each SVG for comparison and requires conversion evidence, a crop "
            "manifest, and an explicit source-vs-render visual review pass before export."
        ),
        "crop_helper_path": str(
            Path(__file__).resolve().parent / "embed_img_crops.py"
        ),
    })

    manifest_slides: list[dict[str, Any]] = []
    for index, image_path in enumerate(images, start=1):
        plan = jobs[index - 1] if index - 1 < len(jobs) else {}
        target_path = target_dir / f"{image_path.stem}.svg"
        job_path = jobs_dir / f"slide-{index:02d}.json"
        crop_manifest_path = jobs_dir / f"slide-{index:02d}-crops.json"
        conversion_evidence_path = jobs_dir / f"slide-{index:02d}-conversion-evidence.json"
        fidelity_review_path = reviews_dir / f"slide-{index:02d}-fidelity-review.json"
        rendered_preview_path = reviews_dir / f"slide-{index:02d}-rendered.png"
        payload = {
            "version": 2,
            "renderer": "img-svg",
            "source_renderer": "img",
            "index": int(plan.get("index", index)),
            "title": plan.get("title", image_path.stem),
            "page_role": plan.get("page_role", "content"),
            "source_image_path": str(image_path),
            "target_path": str(target_path),
            "crop_manifest_path": str(crop_manifest_path),
            "conversion_evidence_path": str(conversion_evidence_path),
            "fidelity_review_path": str(fidelity_review_path),
            "rendered_preview_path": str(rendered_preview_path),
            "shared_context_path": str(shared_context_path),
            "prompt_contract_path": str(
                Path(__file__).resolve().parents[1] / "references" / "prompt-contracts.md"
            ),
            "prompt_contract_section": "IMG-to-SVG Model Conversion Contract",
            "compiled_prompt": IMG_SVG_COMPILED_PROMPT,
            "compiled_prompt_sha256": prompt_sha256,
            "required_model_input": {
                "mode": "image_and_prompt_same_model_turn",
                "image_path": str(image_path),
                "prompt_field": "compiled_prompt",
                "visual_source_of_truth": "source_image_only",
            },
            "crop_helper_path": str(
                Path(__file__).resolve().parent / "embed_img_crops.py"
            ),
        }
        write_json(job_path, payload)
        write_json(
            crop_manifest_path,
            img_svg_crop_manifest_template(image_path, target_path),
        )
        write_json(
            conversion_evidence_path,
            img_svg_conversion_evidence_template(image_path, prompt_sha256),
        )
        manifest_slides.append({
            "index": payload["index"],
            "title": payload["title"],
            "source_image_path": str(image_path),
            "job_path": str(job_path),
            "target_path": str(target_path),
            "crop_manifest_path": str(crop_manifest_path),
            "conversion_evidence_path": str(conversion_evidence_path),
            "fidelity_review_path": str(fidelity_review_path),
            "rendered_preview_path": str(rendered_preview_path),
            "compiled_prompt_sha256": prompt_sha256,
        })

    manifest_path = jobs_dir / "manifest.json"
    write_json(manifest_path, {
        "version": 2,
        "renderer": "img-svg",
        "source_renderer": "img",
        "topic": infer_topic(run_dir),
        "slide_count": len(images),
        "shared_context_path": str(shared_context_path),
        "slides": manifest_slides,
    })

    entry = renderer_state(state, "img")
    conversion = dict(entry.get("svg_conversion") or {})
    entry["svg_conversion"] = {
        **conversion,
        "mode": "on",
        "status": "pending_generation",
        "manifest_path": str(manifest_path),
        "shared_context_path": str(shared_context_path),
        "target_dir": str(target_dir),
        "updated_at": utc_now(),
    }
    entry["status"] = "img_svg_generation_pending"
    state["status"] = "img_svg_generation_pending"
    save_state(run_dir, state)
    return manifest_path


def validate_ppt_compatible_svg(path: Path) -> dict[str, Any]:
    try:
        root = ElementTree.fromstring(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, ElementTree.ParseError) as exc:
        raise ValueError(f"Invalid SVG XML: {path}: {exc}") from exc

    root_name = root.tag.rsplit("}", 1)[-1].lower()
    if root_name != "svg":
        raise ValueError(f"SVG root element is required: {path}")

    view_box = (root.attrib.get("viewBox") or root.attrib.get("viewbox") or "").replace(",", " ").split()
    try:
        normalized_view_box = [float(value) for value in view_box]
    except ValueError as exc:
        raise ValueError(f"SVG viewBox must be numeric: {path}") from exc
    if normalized_view_box != [0.0, 0.0, 1280.0, 720.0]:
        raise ValueError(f"SVG viewBox must be exactly '0 0 1280 720': {path}")

    forbidden = {"foreignobject", "script"}
    raster_area = 0.0
    vector_element_count = 0
    embedded_image_count = 0
    for element in root.iter():
        local_name = element.tag.rsplit("}", 1)[-1].lower()
        if local_name in forbidden:
            raise ValueError(f"SVG contains forbidden <{local_name}> element: {path}")
        if local_name in {
            "text", "tspan", "path", "rect", "circle", "ellipse", "line",
            "polyline", "polygon",
        }:
            vector_element_count += 1
        if local_name == "image":
            embedded_image_count += 1
            href = ""
            for attribute, value in element.attrib.items():
                if attribute.rsplit("}", 1)[-1].lower() == "href":
                    href = value or ""
                    break
            allowed_prefixes = (
                "data:image/png;base64,",
                "data:image/jpeg;base64,",
                "data:image/jpg;base64,",
            )
            if not href.startswith(allowed_prefixes):
                raise ValueError(
                    f"SVG <image> must use an embedded Base64 PNG/JPEG data URI: {path}"
                )
            try:
                encoded = href.split(",", 1)[1]
                raster_bytes = base64.b64decode(encoded, validate=True)
                with Image.open(io.BytesIO(raster_bytes)) as raster:
                    raster.verify()
            except (IndexError, binascii.Error, OSError, ValueError) as exc:
                raise ValueError(f"SVG contains invalid embedded raster data: {path}") from exc

            try:
                x = float(element.attrib.get("x", "0"))
                y = float(element.attrib.get("y", "0"))
                width = float(element.attrib.get("width", "0"))
                height = float(element.attrib.get("height", "0"))
            except ValueError as exc:
                raise ValueError(f"SVG <image> geometry must be numeric: {path}") from exc
            if width <= 0 or height <= 0:
                raise ValueError(f"SVG <image> width and height must be positive: {path}")
            if x <= 1 and y <= 1 and width >= 1278 and height >= 718:
                raise ValueError(
                    f"SVG must not wrap the complete source slide in one full-page <image>: {path}"
                )
            raster_area += width * height
        for attribute, value in element.attrib.items():
            if (
                attribute.rsplit("}", 1)[-1].lower() == "href"
                and value
                and not value.startswith(("#", "data:image/png;base64,", "data:image/jpeg;base64,", "data:image/jpg;base64,"))
            ):
                raise ValueError(f"SVG contains an external href: {path}")

    source = path.read_text(encoding="utf-8").lower()
    if any(marker in source for marker in ("url(http://", "url(https://", "@import")):
        raise ValueError(f"SVG contains external image or stylesheet data: {path}")
    if vector_element_count == 0:
        raise ValueError(f"SVG must contain vector text or shape elements: {path}")
    if raster_area > 1280 * 720 * IMG_SVG_MAX_EMBEDDED_RASTER_RATIO:
        raise ValueError(
            f"SVG embedded raster regions cover more than {IMG_SVG_MAX_EMBEDDED_RASTER_RATIO:.0%} "
            f"of the slide; reconstruct cards and text as vectors: {path}"
        )
    return {
        "vector_element_count": vector_element_count,
        "embedded_image_count": embedded_image_count,
        "embedded_raster_area": raster_area,
        "embedded_raster_ratio": round(raster_area / (1280 * 720), 6),
    }


def complete_img_svg_generation(run_dir: Path) -> list[Path]:
    state = load_state(run_dir)
    if state.get("renderer") != "img":
        raise RuntimeError("IMG-to-SVG completion is only available for the img renderer")
    entry = renderer_state(state, "img")
    conversion = dict(entry.get("svg_conversion") or {})
    if conversion.get("mode") != "on":
        raise RuntimeError("IMG-to-SVG conversion is not enabled")
    if conversion.get("status") != "pending_generation":
        raise RuntimeError("IMG-to-SVG generation is not pending")

    manifest = read_json(Path(conversion["manifest_path"]))
    expected_paths = [Path(item["target_path"]) for item in manifest.get("slides", [])]
    actual_paths = sorted((run_dir / "img-svg").glob("*.svg"))
    if {path.resolve() for path in actual_paths} != {path.resolve() for path in expected_paths}:
        raise ValueError(
            f"IMG-to-SVG output mismatch: expected {len(expected_paths)} exact page files, found {len(actual_paths)}"
        )
    issues: list[str] = []
    review_paths: list[str] = []
    for item in manifest.get("slides", []):
        job = read_json(Path(item["job_path"]))
        svg_path = Path(job["target_path"])
        svg_stats = validate_ppt_compatible_svg(svg_path)
        try:
            validate_img_svg_conversion_evidence(job)
        except (ValueError, FileNotFoundError) as exc:
            issues.append(str(exc))
        try:
            validate_img_svg_crop_manifest(job, svg_stats)
        except (ValueError, FileNotFoundError) as exc:
            issues.append(str(exc))

        review = write_img_svg_review_artifact(job)
        review_paths.append(str(job["fidelity_review_path"]))
        try:
            validate_img_svg_fidelity_review(job, review)
        except (ValueError, RuntimeError) as exc:
            issues.append(str(exc))

    if issues:
        detail = "\n- ".join(issues)
        raise RuntimeError(
            "IMG-to-SVG fidelity gate is not satisfied. Inspect each source image beside its rendered "
            "preview, repair the SVG when needed, complete the evidence/crop manifests, mark the fidelity "
            f"review pass, and rerun complete-img-svg:\n- {detail}"
        )

    entry["svg_conversion"] = {
        **conversion,
        "status": "completed",
        "svg_count": len(actual_paths),
        "svg_paths": [str(path) for path in actual_paths],
        "fidelity_review_paths": review_paths,
        "fidelity_gate": "passed",
        "completed_at": utc_now(),
    }
    entry["status"] = "img_svg_export_ready"
    state["status"] = "img_svg_export_ready"
    save_state(run_dir, state)
    return actual_paths


def export_svg(run_dir: Path) -> Path:
    require_approved(run_dir, "slide_plans")
    from pptx_builder import build_pptx

    svg_dir = run_dir / "svg"
    svg_files = sorted(svg_dir.glob("*.svg")) if svg_dir.exists() else []
    if not svg_files:
        raise ValueError(
            f"SVG directory is empty: {svg_dir}. "
            "This branch needs Codex-authored SVG source files before export."
        )

    pptx_path = run_dir / renderer_pptx_name(run_dir, "svg")
    build_pptx(svg_dir, pptx_path)
    jobs = slide_jobs(run_dir)
    write_slide_status(run_dir, {
        f"{int(job.get('index', i)):02d}": {
            "title": job.get("title", f"Slide {i}"),
            "page_role": job.get("page_role", "content"),
            "validation_status": "codex_generated",
            "export_ready": True,
        }
        for i, job in enumerate(jobs, start=1)
    }, renderer="svg")
    mark_completed(run_dir, "svg", {"pptx": pptx_path})
    return pptx_path


def export_html(run_dir: Path, editable_engine: str | None = None) -> Path:
    require_approved(run_dir, "slide_plans")
    from html_pipeline.html_builder import build_pptx

    topic = infer_topic(run_dir)
    html_dir = run_dir / "html"
    html_files = sorted(html_dir.glob("*.html")) if html_dir.exists() else []
    if not html_files:
        raise ValueError(
            f"HTML directory is empty: {html_dir}. "
            "This branch needs Codex-authored HTML source files before export. "
            "If you used clean-render, note that it only clears derived outputs; "
            "HTML pages still must exist in this branch."
        )

    image_pptx = run_dir / renderer_pptx_name(run_dir, "html")
    build_pptx(html_dir, image_pptx)

    jobs = slide_jobs(run_dir)
    slide_meta = []
    for i, html_path in enumerate(html_files, start=1):
        job = jobs[i - 1] if i - 1 < len(jobs) else {}
        slide_meta.append({
            "index": int(job.get("index", i)),
            "title": job.get("title") or html_path.stem,
            "page_role": job.get("page_role", "content"),
            "html_path": str(html_path),
        })

    status = {
        f"{item['index']:02d}": {
            "title": item["title"],
            "page_role": item["page_role"],
            "validation_status": "codex_generated",
            "export_ready": True,
            "html_path": item["html_path"],
        }
        for item in slide_meta
    }
    write_slide_status(run_dir, status, renderer="html")

    artifacts = {"image_pptx": image_pptx}
    editable_dir = run_dir / "editable"
    try:
        from vendor_presentation_core.export.dom_pptx_exporter import build_dom_editable_deck_from_html

        editable_pptx = build_dom_editable_deck_from_html(
            html_dir=html_dir,
            out_dir=editable_dir,
            slide_meta=slide_meta,
            deck_name=topic,
        )
        artifacts["editable_pptx"] = editable_pptx
    except Exception as exc:
        print(f"[warning] editable export skipped: {exc}", flush=True)

    manifest = {
        "version": 1,
        "source": "codex-skill",
        "topic": topic,
        "renderer": "html",
        "slides": slide_meta,
        "image_pptx_path": str(image_pptx),
        "editable_pptx_path": str(artifacts.get("editable_pptx", "")),
        "editable_engine": editable_engine or "dom_export",
        "updated_at": utc_now(),
    }
    manifest_path = write_json(run_dir / "editable-ppt-chain.json", manifest)
    artifacts["chain_manifest"] = manifest_path
    mark_completed(run_dir, "html", artifacts)
    return image_pptx


def export_img(run_dir: Path) -> Path:
    require_approved(run_dir, "slide_plans")
    from pptx import Presentation
    from pptx.util import Inches

    images = require_complete_img_sources(run_dir)

    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for image in images:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(image), 0, 0, width=prs.slide_width, height=prs.slide_height)

    pptx_path = run_dir / renderer_pptx_name(run_dir, "img")
    prs.save(str(pptx_path))
    write_slide_status(run_dir, {
        f"{i:02d}": {
            "title": image.stem,
            "page_role": "image",
            "validation_status": "codex_imagegen",
            "export_ready": True,
            "image_path": str(image),
        }
        for i, image in enumerate(images, start=1)
    }, renderer="img")
    mark_img_exported_complete(run_dir, pptx_path, images)
    return pptx_path


def export_img_svg(run_dir: Path) -> Path:
    require_approved(run_dir, "slide_plans")
    state = load_state(run_dir)
    if state.get("renderer") != "img":
        raise RuntimeError("IMG-to-SVG export is only available for the img renderer")
    entry = renderer_state(state, "img")
    conversion = dict(entry.get("svg_conversion") or {})
    if conversion.get("mode") != "on" or conversion.get("status") != "completed":
        raise RuntimeError("IMG-to-SVG pages are not complete. Finish the model conversion before export.")

    from pptx_builder import build_native_svg_pptx

    svg_dir = run_dir / "img-svg"
    svg_paths = sorted(svg_dir.glob("*.svg"))
    for svg_path in svg_paths:
        validate_ppt_compatible_svg(svg_path)

    pptx_path = run_dir / renderer_pptx_name(run_dir, "img-svg")
    build_native_svg_pptx(svg_dir, pptx_path)

    manifest = read_json(Path(conversion["manifest_path"]))
    slides = manifest.get("slides", [])
    write_slide_status(run_dir, {
        f"{int(item.get('index', i)):02d}": {
            "title": item.get("title", f"Slide {i}"),
            "page_role": "vectorized_image",
            "validation_status": "codex_img_to_hybrid_svg",
            "export_ready": True,
            "source_image_path": item.get("source_image_path"),
            "svg_path": item.get("target_path"),
        }
        for i, item in enumerate(slides, start=1)
    }, renderer="img-svg")

    chain_manifest = write_json(run_dir / "img-svg-chain.json", {
        "version": 1,
        "source_renderer": "img",
        "conversion": "vision_model_img_to_hybrid_svg",
        "source_pptx_path": conversion.get("source_pptx_path"),
        "svg_dir": str(svg_dir),
        "svg_pptx_path": str(pptx_path),
        "slides": slides,
        "updated_at": utc_now(),
    })

    state = load_state(run_dir)
    entry = renderer_state(state, "img")
    conversion = dict(entry.get("svg_conversion") or {})
    entry["svg_conversion"] = {
        **conversion,
        "status": "exported",
        "pptx_path": str(pptx_path),
        "exported_at": utc_now(),
    }
    save_state(run_dir, state)
    mark_completed(run_dir, "img", {
        "svg_pptx": pptx_path,
        "svg_chain_manifest": chain_manifest,
    })
    return pptx_path


def clean_render_outputs(run_dir: Path) -> None:
    root = run_dir.resolve()
    # Keep renderer source dirs. Codex authors html/svg/img directly in this branch,
    # so default cleanup should only remove derived outputs and review artifacts.
    for name in ("reviews", "editable"):
        target = (run_dir / name).resolve()
        if target.exists():
            if root not in target.parents:
                raise RuntimeError(f"refusing to delete outside run dir: {target}")
            shutil.rmtree(target)
            print(f"deleted={target}", flush=True)
    for pattern in ("*.pptx", "slide-status*.json", "editable-ppt-chain.json", "img-svg-chain.json"):
        for target in run_dir.glob(pattern):
            resolved = target.resolve()
            if root != resolved.parent:
                raise RuntimeError(f"refusing to delete outside run dir: {resolved}")
            target.unlink()
            print(f"deleted={resolved}", flush=True)


def cmd_init(args: argparse.Namespace) -> None:
    run_dir = normalize_run_dir(args.run_dir, topic=args.topic)
    run_dir.mkdir(parents=True, exist_ok=True)
    init_state(run_dir, topic=args.topic, audience=args.audience, pages=args.pages, research=args.research or "")
    print(f"run_dir={run_dir}", flush=True)


def cmd_save_artifact(args: argparse.Namespace) -> None:
    run_dir = normalize_run_dir(args.run_dir)
    run_dir.mkdir(parents=True, exist_ok=True)
    source = Path(args.file)
    data = read_json(source)
    target = write_json(run_dir / ARTIFACT_FILES[args.artifact], data)
    preview = preview_artifact(run_dir, args.artifact)
    print(f"{args.artifact}={target}", flush=True)
    print(f"preview={preview}", flush=True)


def cmd_preview(args: argparse.Namespace) -> None:
    preview = preview_artifact(normalize_run_dir(args.run_dir), args.artifact)
    print(f"preview={preview}", flush=True)


def cmd_approve(args: argparse.Namespace) -> None:
    run_dir = normalize_run_dir(args.run_dir)
    state = load_state(run_dir)
    if args.artifact not in state.get("artifacts", {}):
        raise ValueError(f"unknown artifact: {args.artifact}")
    if args.artifact == "slide_plans":
        slide_plan_document(run_dir)
    state.setdefault("approvals", {})[args.artifact] = True
    state["artifacts"][args.artifact]["status"] = "approved"
    state["artifacts"][args.artifact]["approved_at"] = utc_now()
    state["status"] = f"{args.artifact}_approved"
    save_state(run_dir, state)
    print(f"approved={args.artifact}", flush=True)
    if args.artifact == "slide_plans" and not state.get("renderer"):
        print("next=choose-renderer", flush=True)
        print("recommended_renderer=img", flush=True)


def cmd_choose_renderer(args: argparse.Namespace) -> None:
    run_dir = normalize_run_dir(args.run_dir)
    require_approved(run_dir, "slide_plans")
    state = load_state(run_dir)
    state["renderer"] = args.renderer
    entry = renderer_state(state, args.renderer)
    if args.renderer in {"html", "svg"}:
        entry["review"] = {
            "mode": None,
            "status": "pending_choice",
        }
        entry["status"] = "review_choice_pending"
        state["render_review"] = dict(entry["review"])
        state["status"] = "review_choice_pending"
    else:
        entry["review"] = {
            **RENDER_REVIEW_NOT_APPLICABLE,
            "confirmed_at": utc_now(),
        }
        entry["svg_conversion"] = default_img_svg_conversion()
        entry["status"] = "render_ready"
        state["render_review"] = dict(entry["review"])
        state["status"] = "render_ready"
    save_state(run_dir, state)
    print(f"renderer={args.renderer}", flush=True)
    if args.renderer in {"html", "svg"}:
        print("next=choose-review", flush=True)


def cmd_choose_review(args: argparse.Namespace) -> None:
    run_dir = normalize_run_dir(args.run_dir)
    state = load_state(run_dir)
    renderer = state.get("renderer")
    if renderer not in {"html", "svg"}:
        raise RuntimeError("review choice is only available for html or svg renderer")
    entry = renderer_state(state, renderer)
    entry["review"] = {
        "mode": args.mode,
        "status": "pending" if args.mode == "on" else "skipped",
        "confirmed_at": utc_now(),
    }
    entry["status"] = "render_review_pending" if args.mode == "on" else "render_ready"
    state["render_review"] = dict(entry["review"])
    state["status"] = entry["status"]
    save_state(run_dir, state)
    print(f"review_mode={args.mode}", flush=True)
    if args.mode == "on":
        print("next=complete-review", flush=True)


def cmd_complete_review(args: argparse.Namespace) -> None:
    run_dir = normalize_run_dir(args.run_dir)
    state = load_state(run_dir)
    renderer = state.get("renderer")
    if renderer not in {"html", "svg"}:
        raise RuntimeError("render review completion is only available for html or svg renderer")
    entry = renderer_state(state, renderer)
    review = entry.get("review") or {}
    if review.get("mode") != "on":
        raise RuntimeError("render review is not enabled")
    entry["review"] = {
        **review,
        "status": "completed",
        "completed_at": utc_now(),
    }
    entry["status"] = "render_ready"
    state["render_review"] = dict(entry["review"])
    state["status"] = "render_ready"
    save_state(run_dir, state)
    print(f"review_completed={renderer}", flush=True)


def cmd_choose_img_svg(args: argparse.Namespace) -> None:
    run_dir = normalize_run_dir(args.run_dir)
    state = load_state(run_dir)
    if state.get("renderer") != "img":
        raise RuntimeError("IMG-to-SVG conversion is only available for the img renderer")
    entry = renderer_state(state, "img")
    conversion = dict(entry.get("svg_conversion") or {})
    if conversion.get("status") not in {"available", "skipped"}:
        raise RuntimeError(
            "IMG-to-SVG conversion is only available after the IMG PPTX has been exported"
        )
    source_pptx = Path(str(conversion.get("source_pptx_path") or ""))
    if not source_pptx.is_file():
        raise RuntimeError("The exported IMG PPTX is missing; re-export it before starting IMG-to-SVG conversion")

    if args.mode == "on":
        conversion.update({
            "mode": "on",
            "status": "pending_generation",
            "confirmed_at": utc_now(),
        })
        entry["svg_conversion"] = conversion
        save_state(run_dir, state)
        manifest = prepare_img_svg_jobs(run_dir, load_state(run_dir))
        print("img_svg_mode=on", flush=True)
        print(f"img_svg_jobs={manifest}", flush=True)
        print("next=complete-img-svg", flush=True)
        return

    entry["svg_conversion"] = {
        **conversion,
        "mode": "off",
        "status": "skipped",
        "confirmed_at": utc_now(),
    }
    entry["status"] = "completed"
    entry["completed_at"] = utc_now()
    state["status"] = "completed"
    save_state(run_dir, state)
    print("img_svg_mode=off", flush=True)
    print("completed=img", flush=True)


def cmd_complete_img_svg(args: argparse.Namespace) -> None:
    run_dir = normalize_run_dir(args.run_dir)
    svg_paths = complete_img_svg_generation(run_dir)
    print(f"img_svg_completed={len(svg_paths)}", flush=True)
    print("next=export-img-svg", flush=True)


def cmd_export_img_svg(args: argparse.Namespace) -> None:
    run_dir = normalize_run_dir(args.run_dir)
    out = export_img_svg(run_dir)
    print(f"completed={out}", flush=True)
    print("post_export_guidance=office-convert-svg-to-shape", flush=True)
    print("office_editability=vector-shapes-not-semantic-text-or-charts", flush=True)


def cmd_prepare_render_jobs(args: argparse.Namespace) -> None:
    run_dir = normalize_run_dir(args.run_dir)
    renderer = args.renderer or load_state(run_dir).get("renderer")
    if renderer not in {"html", "svg", "img"}:
        raise RuntimeError("renderer must be html, svg, or img")
    manifest = prepare_render_jobs(run_dir, renderer)
    print(f"render_jobs={manifest}", flush=True)


def cmd_export(args: argparse.Namespace) -> None:
    run_dir = normalize_run_dir(args.run_dir)
    state = load_state(run_dir)
    renderer = args.renderer or state.get("renderer")
    if renderer:
        ensure_render_ready(run_dir, renderer)
    if renderer == "html":
        out = export_html(run_dir, editable_engine=args.editable_engine)
    elif renderer == "svg":
        out = export_svg(run_dir)
    elif renderer == "img":
        out = export_img(run_dir)
    else:
        raise ValueError("renderer must be html, svg, or img")
    if renderer == "img":
        print(f"img_pptx={out}", flush=True)
        print(f"completed={out}", flush=True)
        print("workflow_complete=yes", flush=True)
        print("user_reply_required=no", flush=True)
        print("offer_img_svg=yes", flush=True)
        print("img_svg_conversion=available_on_explicit_request", flush=True)
        print("img_svg_editability=powerpoint-convert-to-shape", flush=True)
        print("img_svg_additional_model_usage=yes", flush=True)
        print("optional_next=choose-img-svg --mode on", flush=True)
    else:
        print(f"completed={out}", flush=True)


def cmd_status(args: argparse.Namespace) -> None:
    run_dir = normalize_run_dir(args.run_dir)
    state = load_state(run_dir)
    print(f"run_dir={run_dir}", flush=True)
    print(f"status={state.get('status')}", flush=True)
    print(f"renderer={state.get('renderer')}", flush=True)
    review = state.get("render_review") or {}
    print(f"review_mode={review.get('mode')}", flush=True)
    print(f"review_status={review.get('status')}", flush=True)
    next_steps = {
        "img_svg_generation_pending": "complete-img-svg",
        "img_svg_export_ready": "export-img-svg",
    }
    if state.get("status") in next_steps:
        print(f"next={next_steps[state['status']]}", flush=True)
    for name, entry in sorted((state.get("artifacts") or {}).items()):
        print(f"{name}: {entry.get('status')} {entry.get('path')}", flush=True)
    for renderer_name, entry in sorted((state.get("renderers") or {}).items()):
        renderer_review = entry.get("review") or {}
        print(
            f"renderer[{renderer_name}]: status={entry.get('status')} "
            f"review={renderer_review.get('mode')}/{renderer_review.get('status')}",
            flush=True,
        )
        if renderer_name == "img":
            conversion = entry.get("svg_conversion") or {}
            print(
                f"renderer[img].svg_conversion: mode={conversion.get('mode')} "
                f"status={conversion.get('status')}",
                flush=True,
            )
        for artifact_name, artifact_entry in sorted((entry.get("artifacts") or {}).items()):
            print(f"renderer[{renderer_name}].{artifact_name}: {artifact_entry.get('status')} {artifact_entry.get('path')}", flush=True)
    slides = read_json(run_dir / "slide-status.json", default={}).get("slides") or {}
    if slides:
        ready = sum(1 for item in slides.values() if item.get("export_ready"))
        print(f"slides_ready={ready}/{len(slides)}", flush=True)
    patterns = {"html": "*.html", "svg": "*.svg", "img": "*.*", "img-svg": "*.svg"}
    for name, pattern in patterns.items():
        count = len(list((run_dir / name).glob(pattern))) if (run_dir / name).exists() else 0
        if count:
            print(f"{name}_files={count}", flush=True)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Codex-skill PPT workflow artifact helper")
    sub = parser.add_subparsers(dest="command", required=True)

    init = sub.add_parser("init")
    init.add_argument("--topic", required=True)
    init.add_argument("--audience", default="通用受众")
    init.add_argument("--pages", default="12-15页")
    init.add_argument("--research", default="")
    init.add_argument("--run-dir", default=None)
    init.set_defaults(func=cmd_init)

    save = sub.add_parser("save-artifact")
    save.add_argument("--run-dir", required=True)
    save.add_argument("--artifact", choices=sorted(ARTIFACT_FILES), required=True)
    save.add_argument("--file", required=True)
    save.set_defaults(func=cmd_save_artifact)

    preview = sub.add_parser("preview")
    preview.add_argument("--run-dir", required=True)
    preview.add_argument("--artifact", choices=sorted(ARTIFACT_FILES), required=True)
    preview.set_defaults(func=cmd_preview)

    approve = sub.add_parser("approve")
    approve.add_argument("--run-dir", required=True)
    approve.add_argument("--artifact", choices=sorted(ARTIFACT_FILES), required=True)
    approve.set_defaults(func=cmd_approve)

    choose = sub.add_parser("choose-renderer")
    choose.add_argument("--run-dir", required=True)
    choose.add_argument("--renderer", choices=["html", "svg", "img"], required=True)
    choose.set_defaults(func=cmd_choose_renderer)

    choose_review = sub.add_parser("choose-review")
    choose_review.add_argument("--run-dir", required=True)
    choose_review.add_argument("--mode", choices=["off", "on"], required=True)
    choose_review.set_defaults(func=cmd_choose_review)

    choose_img_svg = sub.add_parser("choose-img-svg")
    choose_img_svg.add_argument("--run-dir", required=True)
    choose_img_svg.add_argument("--mode", choices=["off", "on"], required=True)
    choose_img_svg.set_defaults(func=cmd_choose_img_svg)

    jobs = sub.add_parser("prepare-render-jobs")
    jobs.add_argument("--run-dir", required=True)
    jobs.add_argument("--renderer", choices=["html", "svg", "img"], default=None)
    jobs.set_defaults(func=cmd_prepare_render_jobs)

    complete_review = sub.add_parser("complete-review")
    complete_review.add_argument("--run-dir", required=True)
    complete_review.set_defaults(func=cmd_complete_review)

    complete_img_svg = sub.add_parser("complete-img-svg")
    complete_img_svg.add_argument("--run-dir", required=True)
    complete_img_svg.set_defaults(func=cmd_complete_img_svg)

    export_img_svg_parser = sub.add_parser("export-img-svg")
    export_img_svg_parser.add_argument("--run-dir", required=True)
    export_img_svg_parser.set_defaults(func=cmd_export_img_svg)

    export = sub.add_parser("export")
    export.add_argument("--run-dir", required=True)
    export.add_argument("--renderer", choices=["html", "svg", "img"], default=None)
    export.add_argument("--editable-engine", default=None)
    export.set_defaults(func=cmd_export)

    clean = sub.add_parser("clean-render")
    clean.add_argument("--run-dir", required=True)
    clean.set_defaults(func=lambda args: clean_render_outputs(normalize_run_dir(args.run_dir)))

    status = sub.add_parser("status")
    status.add_argument("--run-dir", required=True)
    status.set_defaults(func=cmd_status)
    return parser


def main() -> None:
    args = build_parser().parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
