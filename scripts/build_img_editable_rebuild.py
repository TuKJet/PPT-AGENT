from __future__ import annotations

import json
import re
from pathlib import Path

from PIL import Image, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


SLIDE_W = 13.333
SLIDE_H = 7.5
FONT = "Microsoft YaHei"

NAVY = RGBColor(11, 31, 54)
BLUE = RGBColor(37, 99, 235)
TEAL = RGBColor(20, 184, 166)
CYAN = RGBColor(6, 182, 212)
GREEN = RGBColor(34, 197, 94)
ORANGE = RGBColor(245, 158, 11)
PURPLE = RGBColor(124, 58, 237)
SLATE = RGBColor(71, 85, 105)
LIGHT = RGBColor(248, 250, 252)
CARD = RGBColor(255, 255, 255)
BORDER = RGBColor(226, 232, 240)
PALE_BLUE = RGBColor(239, 246, 255)
PALE_TEAL = RGBColor(240, 253, 250)
PALE_ORANGE = RGBColor(255, 247, 237)
PALE_PURPLE = RGBColor(245, 243, 255)

ACCENTS = [BLUE, TEAL, ORANGE, PURPLE, CYAN, GREEN]
PALES = [PALE_BLUE, PALE_TEAL, PALE_ORANGE, PALE_PURPLE]


def rgb_tuple(c: RGBColor) -> tuple[int, int, int]:
    return int(c[0]), int(c[1]), int(c[2])


def run_dir() -> Path:
    candidates = [
        p
        for p in Path("output").iterdir()
        if p.is_dir() and "AI_Agent" in p.name and (p / "img").exists()
    ]
    if not candidates:
        raise FileNotFoundError("No AI_Agent run directory with img/ found.")
    return max(candidates, key=lambda p: (p / "img").stat().st_mtime)


def clean_bullets(text: str) -> list[str]:
    lines: list[str] = []
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        line = re.sub(r"^[\-•\d\.\s]+", "", line).strip()
        if line:
            lines.append(line)
    return lines


def fit_font(text: str, base: int, min_size: int = 9) -> int:
    length = len(text)
    if length > 95:
        return max(min_size, base - 5)
    if length > 70:
        return max(min_size, base - 3)
    if length > 42:
        return max(min_size, base - 2)
    return base


def set_run(run, size: int, color: RGBColor = SLATE, bold: bool = False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold


def set_shape_fill(shape, fill_color: RGBColor, transparency: int = 0):
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    shape.fill.transparency = transparency


def set_line(shape, color: RGBColor = BORDER, width: float = 1.0, transparency: int = 0):
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    shape.line.transparency = transparency


def textbox(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    text: str,
    size: int = 16,
    color: RGBColor = SLATE,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.01)
    tf.margin_bottom = Inches(0.01)
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = Pt(0)
    run = p.add_run()
    run.text = text
    set_run(run, size, color, bold)
    return shape


def bullet_box(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    bullets: list[str],
    size: int = 11,
    color: RGBColor = SLATE,
    max_items: int | None = None,
):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.05)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.02)
    tf.margin_bottom = Inches(0.02)
    items = bullets[:max_items] if max_items else bullets
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = FONT
        p.font.size = Pt(fit_font(item, size, 8))
        p.font.color.rgb = color
        p.space_after = Pt(5)
    return shape


def rounded_rect(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    fill: RGBColor = CARD,
    line: RGBColor = BORDER,
    radius_shape=MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
    transparency: int = 0,
):
    shape = slide.shapes.add_shape(radius_shape, Inches(x), Inches(y), Inches(w), Inches(h))
    set_shape_fill(shape, fill, transparency)
    set_line(shape, line, 0.8)
    return shape


def add_card(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    title: str,
    bullets: list[str],
    accent: RGBColor = BLUE,
    pale: RGBColor = PALE_BLUE,
    number: str | None = None,
    max_items: int = 3,
):
    rounded_rect(slide, x, y, w, h, CARD, BORDER)
    rounded_rect(slide, x + 0.13, y + 0.13, 0.07, h - 0.26, accent, accent)
    if number:
        badge = rounded_rect(slide, x + 0.28, y + 0.18, 0.42, 0.32, pale, accent)
        badge.text_frame.clear()
        p = badge.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.vertical_anchor = MSO_ANCHOR.MIDDLE
        r = p.add_run()
        r.text = number
        set_run(r, 10, accent, True)
        title_x = x + 0.82
        title_w = w - 0.98
    else:
        title_x = x + 0.32
        title_w = w - 0.48
    textbox(slide, title_x, y + 0.17, title_w, 0.34, title, 12, NAVY, True)
    bullet_box(slide, x + 0.32, y + 0.62, w - 0.52, h - 0.75, bullets, 9, SLATE, max_items)


def add_header(slide, title: str, kicker: str, page: int, accent: RGBColor = BLUE):
    textbox(slide, 0.55, 0.27, 0.9, 0.22, f"{page:02d}", 8, accent, True)
    rounded_rect(slide, 0.55, 0.52, 0.36, 0.05, accent, accent)
    textbox(slide, 0.55, 0.68, 9.8, 0.45, title, 20, NAVY, True)
    if kicker:
        textbox(slide, 0.57, 1.08, 10.6, 0.3, kicker, 9, SLATE, False)
    textbox(slide, 10.8, 0.32, 1.95, 0.25, "AI AGENT ASSOCIATION", 7, SLATE, False, PP_ALIGN.RIGHT)


def add_footer(slide, page: int, accent: RGBColor = BLUE):
    rounded_rect(slide, 0.55, 6.92, 12.25, 0.34, NAVY, NAVY)
    textbox(slide, 0.78, 6.99, 8.4, 0.18, "管理层汇报 · IMG视觉重建可编辑稿", 7, RGBColor(226, 232, 240), False)
    textbox(slide, 11.65, 6.99, 0.9, 0.18, f"{page:02d}/14", 7, RGBColor(226, 232, 240), False, PP_ALIGN.RIGHT)
    rounded_rect(slide, 10.65, 6.995, 0.78, 0.16, accent, accent)


def add_background(slide, image_path: Path):
    slide.shapes.add_picture(str(image_path), Inches(0), Inches(0), width=Inches(SLIDE_W), height=Inches(SLIDE_H))


def make_soft_backgrounds(run: Path) -> list[Path]:
    out_dir = run / "img_editable_rebuild_assets"
    out_dir.mkdir(exist_ok=True)
    bg_paths: list[Path] = []
    for src in sorted((run / "img").glob("slide-*.png")):
        out = out_dir / f"{src.stem}-soft-bg.jpg"
        im = Image.open(src).convert("RGB").resize((1920, 1080), Image.LANCZOS)
        im = im.filter(ImageFilter.GaussianBlur(6))
        white = Image.new("RGB", im.size, "white")
        im = Image.blend(white, im, 0.13)
        im.save(out, quality=90)
        bg_paths.append(out)
    return bg_paths


def draw_cover(slide, title: str, bullets: list[str]):
    textbox(slide, 0.72, 0.58, 1.8, 0.22, "AI AGENT GOVERNANCE", 8, BLUE, True)
    textbox(slide, 0.72, 1.05, 6.3, 0.75, title, 28, NAVY, True)
    textbox(slide, 0.75, 1.86, 7.0, 0.46, "以组织化机制推动AI知识沉淀、人才培养、研发提效与业务创新", 14, SLATE, False)
    add_card(slide, 0.72, 2.65, 5.9, 1.65, "本次建议", bullets[:3], BLUE, PALE_BLUE, "01", 3)
    add_card(slide, 0.72, 4.55, 5.9, 1.35, "资源诉求", bullets[3:], TEAL, PALE_TEAL, "02", 2)
    # Editable ring-like motif.
    for i, (size, color) in enumerate([(3.05, BLUE), (2.35, TEAL), (1.65, RGBColor(255, 255, 255))]):
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(8.25 + i * 0.35), Inches(1.38 + i * 0.35), Inches(size), Inches(size))
        set_shape_fill(shape, color if i < 2 else CARD, 18 if i < 2 else 0)
        set_line(shape, color, 2)
    textbox(slide, 9.05, 2.1, 2.2, 0.42, "2026", 24, BLUE, True, PP_ALIGN.CENTER)
    textbox(slide, 8.75, 2.62, 2.9, 0.48, "公司级组织化推进", 13, NAVY, True, PP_ALIGN.CENTER)
    for idx, label in enumerate(["统一治理", "标准复用", "规模落地"]):
        add_card(slide, 7.65 + idx * 1.65, 5.25, 1.35, 0.78, label, [], ACCENTS[idx], PALES[idx], None, 0)


def draw_agenda(slide, title: str, bullets: list[str]):
    add_header(slide, title, "围绕建设必要性、运行机制、落地路径与投入产出判断展开", 2, BLUE)
    labels = ["为什么建", "怎么运作", "怎么落地", "投入产出"]
    for i, label in enumerate(labels):
        x = 0.72 + i * 3.05
        add_card(slide, x, 1.85, 2.63, 3.2, label, [bullets[i]], ACCENTS[i], PALES[i], f"{i+1}", 1)
        if i < 3:
            textbox(slide, x + 2.68, 3.15, 0.35, 0.25, "→", 18, ACCENTS[i], True, PP_ALIGN.CENTER)
    rounded_rect(slide, 1.55, 5.65, 10.2, 0.62, NAVY, NAVY)
    textbox(slide, 1.9, 5.82, 9.5, 0.24, "背景必要性 → 运行机制 → 落地计划 → 价值决策", 12, RGBColor(226, 232, 240), True, PP_ALIGN.CENTER)
    add_footer(slide, 2, BLUE)


def draw_three_focus(slide, title: str, bullets: list[str], page: int):
    add_header(slide, title, "把建设目的、推进方案和资源诉求放到同一张决策视图中", page, TEAL)
    add_card(slide, 0.7, 1.55, 3.95, 1.55, "规模化窗口", bullets[:2], BLUE, PALE_BLUE, "88%", 2)
    add_card(slide, 4.95, 1.55, 3.3, 1.55, "协会牵引", bullets[2:3], TEAL, PALE_TEAL, "PoC", 1)
    add_card(slide, 8.55, 1.55, 3.75, 1.55, "治理前置", bullets[3:4], ORANGE, PALE_ORANGE, "11%", 1)
    add_card(slide, 0.7, 3.55, 5.45, 1.95, "管理层需对齐", bullets[1:4], PURPLE, PALE_PURPLE, "3", 3)
    add_card(slide, 6.45, 3.55, 5.85, 1.95, "本次决策事项", bullets[-1:], TEAL, PALE_TEAL, "授权", 1)
    add_footer(slide, page, TEAL)


def draw_cards_grid(slide, title: str, bullets: list[str], page: int, labels: list[str], accent: RGBColor):
    add_header(slide, title, "", page, accent)
    for i, label in enumerate(labels):
        r, c = divmod(i, 2)
        add_card(slide, 0.75 + c * 5.9, 1.55 + r * 2.25, 5.45, 1.78, label, [bullets[i % len(bullets)]], ACCENTS[i % len(ACCENTS)], PALES[i % len(PALES)], f"{i+1:02d}", 1)
    add_card(slide, 0.75, 5.92, 11.8, 0.75, "关键结论", [bullets[-1]], accent, PALE_TEAL, None, 1)
    add_footer(slide, page, accent)


def draw_hub(slide, title: str, bullets: list[str], page: int):
    add_header(slide, title, "协会不替代部门职责，而是公司级AI能力建设与场景孵化接口", page, PURPLE)
    hub = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, Inches(5.15), Inches(2.45), Inches(2.4), Inches(2.4))
    set_shape_fill(hub, PURPLE, 8)
    set_line(hub, PURPLE, 2)
    textbox(slide, 5.45, 3.05, 1.8, 0.52, "AI协会平台", 18, RGBColor(255, 255, 255), True, PP_ALIGN.CENTER)
    positions = [(0.75, 1.55), (8.25, 1.55), (0.75, 4.55), (8.25, 4.55)]
    labels = ["承接战略", "连接资源", "项目孵化", "知识沉淀"]
    for i, (x, y) in enumerate(positions):
        add_card(slide, x, y, 3.75, 1.45, labels[i], [bullets[i]], ACCENTS[i], PALES[i], None, 1)
    add_footer(slide, page, PURPLE)


def draw_org(slide, title: str, bullets: list[str], page: int):
    add_header(slide, title, "形成决策、运营、专项和治理协同机制", page, BLUE)
    add_card(slide, 1.1, 1.45, 11.1, 0.95, "指导委员会", [bullets[0]], BLUE, PALE_BLUE, "决策", 1)
    add_card(slide, 2.25, 2.75, 8.8, 0.95, "协会负责人 / 秘书处", [bullets[1]], TEAL, PALE_TEAL, "运营", 1)
    labels = ["场景专项组", "培训专项组", "孵化专项组", "治理岗位"]
    for i, label in enumerate(labels):
        add_card(slide, 0.75 + i * 3.0, 4.25, 2.65, 1.45, label, [bullets[min(i + 2, len(bullets) - 1)]], ACCENTS[i], PALES[i % len(PALES)], None, 1)
    add_footer(slide, page, BLUE)


def draw_tiers(slide, title: str, bullets: list[str], page: int):
    add_header(slide, title, "按贡献深度分层管理，让资源和权限跟随贡献动态流动", page, TEAL)
    tiers = ["核心成员", "培养成员", "专项成员", "普通社群"]
    heights = [0.85, 0.95, 1.05, 1.15]
    widths = [3.1, 4.35, 5.6, 6.85]
    for i, tier in enumerate(tiers):
        x = 1.0 + i * 0.63
        y = 1.55 + i * 1.05
        add_card(slide, x, y, widths[i], heights[i], tier, [bullets[i]], ACCENTS[i], PALES[i % len(PALES)], None, 1)
    add_card(slide, 8.2, 1.75, 3.9, 3.8, "动态管理", bullets[-2:], BLUE, PALE_BLUE, "半年", 2)
    add_footer(slide, page, TEAL)


def draw_roadmap(slide, title: str, bullets: list[str], page: int):
    add_header(slide, title, "12个月分阶段推进，覆盖人群、活动频率和阶段产出同步明确", page, BLUE)
    phases = [
        ("1-2月", "启动建制", "15-20人核心小组", bullets[0]),
        ("3-4月", "训练扩面", "50-80名种子用户", bullets[1]),
        ("5-8月", "场景孵化", "5-8个PoC原型", bullets[2]),
        ("9-12月", "复用推广", "2-3个方法复用", bullets[3]),
    ]
    for i, (months, name, scope, body) in enumerate(phases):
        x = 0.68 + i * 3.08
        add_card(slide, x, 1.55, 2.72, 3.35, name, [scope, body], ACCENTS[i], PALES[i % len(PALES)], months, 2)
    add_card(slide, 0.75, 5.32, 11.75, 0.92, "活动频率", [bullets[4]], TEAL, PALE_TEAL, "全年", 1)
    add_footer(slide, page, BLUE)


def draw_value(slide, title: str, bullets: list[str], page: int):
    add_header(slide, title, "从知识资产、人才梯队、生产力提升和组织活性四个维度衡量协会价值", page, GREEN)
    labels = ["知识沉淀", "人才培养", "生产力衔接", "组织活性"]
    for i, label in enumerate(labels):
        x = 0.8 + (i % 2) * 5.75
        y = 1.55 + (i // 2) * 2.25
        add_card(slide, x, y, 5.25, 1.72, label, [bullets[i]], ACCENTS[i], PALES[i % len(PALES)], f"{i+1}", 1)
    add_card(slide, 2.0, 6.0, 9.3, 0.55, "价值锚点", [bullets[-1]], GREEN, PALE_TEAL, None, 1)
    add_footer(slide, page, GREEN)


def draw_resources(slide, title: str, bullets: list[str], page: int):
    add_header(slide, title, "建议先批准12个月试运行，用季度复盘控制投入、价值和风险", page, ORANGE)
    add_card(slide, 0.82, 1.55, 4.1, 2.0, "建议决策", bullets[:2], ORANGE, PALE_ORANGE, "12个月", 2)
    add_card(slide, 5.18, 1.55, 3.25, 2.0, "资源配置", bullets[2:3], BLUE, PALE_BLUE, "资源", 1)
    add_card(slide, 8.68, 1.55, 3.55, 2.0, "联合把关", bullets[3:4], TEAL, PALE_TEAL, "治理", 1)
    add_card(slide, 0.82, 4.05, 11.4, 1.55, "验收口径", bullets[4:], PURPLE, PALE_PURPLE, "指标", 1)
    add_footer(slide, page, ORANGE)


def build_deck():
    root = run_dir()
    contents = json.loads((root / "contents.json").read_text(encoding="utf-8"))
    slides = [(title, clean_bullets(body)) for title, body in contents.items()]
    bg_paths = make_soft_backgrounds(root)

    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    blank = prs.slide_layouts[6]

    for idx, (title, bullets) in enumerate(slides, 1):
        slide = prs.slides.add_slide(blank)
        add_background(slide, bg_paths[idx - 1])
        if idx == 1:
            draw_cover(slide, title, bullets)
        elif idx == 2:
            draw_agenda(slide, title, bullets)
        elif idx == 3:
            draw_three_focus(slide, title, bullets, idx)
        elif idx == 4:
            draw_cards_grid(slide, title, bullets, idx, ["外部趋势", "内部需求", "当前痛点", "治理前置"], BLUE)
        elif idx == 5:
            draw_hub(slide, title, bullets, idx)
        elif idx == 6:
            draw_org(slide, title, bullets, idx)
        elif idx == 7:
            draw_tiers(slide, title, bullets, idx)
        elif idx == 8:
            draw_cards_grid(slide, title, bullets, idx, ["准入机制", "权限额度", "风险边界", "台账留痕"], TEAL)
        elif idx == 9:
            draw_cards_grid(slide, title, bullets, idx, ["评估依据", "数据审计", "培养保留", "降级释放"], ORANGE)
        elif idx == 10:
            draw_cards_grid(slide, title, bullets, idx, ["需求池", "高复用场景", "小型验证", "成功率对比"], BLUE)
        elif idx == 11:
            draw_cards_grid(slide, title, bullets, idx, ["周节奏", "月度机制", "季度复盘", "半年调整"], CYAN)
        elif idx == 12:
            draw_roadmap(slide, title, bullets, idx)
        elif idx == 13:
            draw_value(slide, title, bullets, idx)
        elif idx == 14:
            draw_resources(slide, title, bullets, idx)
        else:
            draw_cards_grid(slide, title, bullets, idx, ["要点一", "要点二", "要点三", "要点四"], BLUE)

    out = root / "AI_Agent_association_management_img_editable_rebuild.pptx"
    prs.save(out)
    return out


if __name__ == "__main__":
    print(build_deck())
